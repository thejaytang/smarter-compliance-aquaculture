"""Owning relational source state, with a one-time, immutable workbook import.

The workbook loaded by ``workbook`` is an in-memory compatibility view over
database records. The derived Excel file is never read back into this store.
"""
from contextlib import contextmanager
from datetime import date, datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from tempfile import NamedTemporaryFile
from copy import copy
from functools import lru_cache
from pathlib import Path
import sqlite3
from system1.sqlite_support import connect as connect_sqlite

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.table import Table
from openpyxl.utils.datetime import from_excel
from openpyxl.styles.numbers import is_date_format
from openpyxl.formula.translate import Translator

SCHEMA = 'system1-governance/1'
DERIVED = {'selection_status', 'current_snapshot_date', 'needs_human_action'}


def encode(record):
    def value(item):
        if isinstance(item, datetime): return ['datetime', item.isoformat()]
        if isinstance(item, date): return ['date', item.isoformat()]
        if item is None or isinstance(item, (str, int, float, bool)): return ['value', item]
        raise ValueError('Unsupported source-state value: '+type(item).__name__)
    return json.dumps({k: value(v) for k, v in record.items()}, sort_keys=True, ensure_ascii=False)


def decode(raw):
    def value(item):
        kind, data = item
        if kind == 'datetime': return datetime.fromisoformat(data)
        if kind == 'date': return date.fromisoformat(data)
        if kind == 'value': return data
        raise ValueError('Unsupported stored value type')
    return {k: value(v) for k, v in json.loads(raw).items()}


def digest(data):
    return sha256(data).hexdigest()


def publish_bytes(path, raw):
    """Publish an immutable companion exclusively; a matching prior retry is safe."""
    if path.exists():
        if path.read_bytes() != raw: raise ValueError('Immutable archive identity conflict')
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=path.parent, prefix='.'+path.name, delete=False) as handle:
        temporary = Path(handle.name)
        handle.write(raw); handle.flush(); os.fsync(handle.fileno())
    try:
        try: os.link(temporary, path)
        except FileExistsError:
            if path.read_bytes() != raw: raise ValueError('Immutable archive identity conflict')
    finally: temporary.unlink(missing_ok=True)


def original_inventory(root):
    result = []
    for item in sorted(root.rglob('*')):
        if item.is_file():
            if root not in item.resolve().parents: raise ValueError('Source artifact leaves managed root')
            raw = item.read_bytes()
            result.append({'path':str(item.relative_to(root)), 'sha256':digest(raw), 'bytes':len(raw)})
    return result


def authority_version(config):
    """Logical committed versions include SQLite WAL state, not file timestamps."""
    store = GovernanceStore(config['governance_db'])
    with store.connect() as db:
        state = db.execute('SELECT revision,import_sha256 FROM state').fetchone()
    assessment = config.get('assessment_db', config['log_root']/'source-assessments.sqlite')
    if config.get('assessment_db'):
        from .source_assessment import Assessments
        Assessments(assessment)
    records = None
    if assessment.is_file():
        with connect_sqlite(assessment.as_uri()+'?mode=ro',uri=True) as db:
            db.execute('BEGIN')
            records = {name: sorted(list(db.execute('SELECT * FROM '+name))) for name in
                       ('assessments','assessment_events','assessment_policy','assessment_holds')}
    return {'revision':state[0], 'import_sha256':state[1],
            'assessment_sha256':digest(json.dumps(records,sort_keys=True,default=str).encode())}


def business_rows(wb, config, check_cached=None):
    import human_operations as h
    import source_updater as u
    from .formula_cache import WorkbookCalculator
    calc = WorkbookCalculator(wb)
    result = {}
    for kind, sheet, headers, start, identity in (
        ('sources', config['sheet_name'], u.workbook_headers(wb[config['sheet_name']], config['header_row']), config['data_start_row'], 'source_id'),
        ('operations', h.human_sheet_name(config), h.operation_headers(wb[h.human_sheet_name(config)]), h.HUMAN_DATA_START_ROW, 'operation_id'),
    ):
        rows = []
        seen = set()
        for row in range(start, wb[sheet].max_row+1):
            rec = {key: wb[sheet].cell(row, col).value for key, col in headers.items()}
            if not rec.get(identity): continue
            if rec[identity] in seen: raise ValueError('Duplicate '+identity+': '+str(rec[identity]))
            seen.add(rec[identity])
            for key, col in headers.items():
                cell = wb[sheet].cell(row, col)
                if cell.data_type == 'e': raise ValueError('Error in business cell: '+sheet+'!'+cell.coordinate)
                if cell.data_type != 'f': continue
                if kind != 'sources' or key not in DERIVED:
                    raise ValueError('Unmapped business formula: '+sheet+'!'+cell.coordinate)
                v = calc.cell(sheet, cell.coordinate)
                if isinstance(v, (int, float)) and not isinstance(v, bool) and is_date_format(cell.number_format):
                    v = from_excel(v, wb.epoch)
                # Excel writes empty-string formula results as a blank cached cell.
                if v == '': v = None
                if check_cached is not None and v != check_cached[sheet][cell.coordinate].value:
                    raise ValueError('Cached business value differs: '+sheet+'!'+cell.coordinate)
                rec[key] = v
            rows.append({'id': rec[identity], 'position': row, 'record': rec})
        result[kind] = rows
    return result


@lru_cache(maxsize=8)
def business_layout(raw, source_sheet, header_row, data_start_row, operation_sheet):
    """Cache immutable layout metadata, never mutable business values or decisions."""
    import human_operations as h
    import source_updater as u
    wb = load_workbook(BytesIO(raw))
    try:
        layout = {'epoch':wb.epoch, 'sheets':{}}
        for kind, name, header, start, headers in (
            ('sources', source_sheet, header_row, data_start_row, u.workbook_headers(wb[source_sheet],header_row)),
            ('operations', operation_sheet, h.HUMAN_HEADER_ROW, h.HUMAN_DATA_START_ROW, h.operation_headers(wb[operation_sheet])),
        ):
            sheet=wb[name];formulas={}
            if kind=='sources':
                for key in DERIVED & headers.keys():
                    for row in range(start,sheet.max_row+1):
                        cell=sheet.cell(row,headers[key])
                        if cell.data_type=='f':
                            formulas[key]=(cell.value,cell.coordinate)
                            break
            layout['sheets'][kind] = {
                'name':name,'header':header,'headers':headers,
                'labels':{key:sheet.cell(header,col).value for key,col in headers.items()},
                'formats':{key:sheet.cell(start,col).number_format for key,col in headers.items()},
                'formulas':formulas,'tables':[(table.name,table.ref) for table in sheet.tables.values()],
            }
        return layout
    finally: wb.close()


class GovernanceStore:
    def __init__(self, path):
        self.path = Path(path).resolve()

    @contextmanager
    def connect(self):
        with connect_sqlite(self.path.as_uri()+'?mode=rw', uri=True, timeout=10) as db:
            db.row_factory = sqlite3.Row
            db.execute('PRAGMA foreign_keys=ON')
            db.execute('PRAGMA synchronous=FULL')
            if db.execute('SELECT schema FROM state').fetchone()[0] != SCHEMA:
                raise ValueError('Unsupported governance database schema')
            yield db

    @classmethod
    def migrate(cls, config, path):
        """One-time import; a matching immutable archive can resume an interrupted attempt."""
        path = Path(path).resolve()
        if path.exists(): raise ValueError('Governance database already exists; import is one-time only.')
        raw = config['workbook'].read_bytes()
        source_hash = digest(raw)
        formula = load_workbook(BytesIO(raw), data_only=False)
        cached = load_workbook(BytesIO(raw), data_only=True)
        try: rows = business_rows(formula, config, check_cached=cached)
        finally: formula.close(); cached.close()
        path.parent.mkdir(parents=True, exist_ok=True)
        archive = path.parent / (path.stem+'-migration-input.xlsx')
        # An interrupted attempt may leave this immutable, byte-identical archive.
        publish_bytes(archive, raw)
        source_root = config['source_root'].resolve()
        artifacts = original_inventory(source_root)
        with NamedTemporaryFile(dir=path.parent, prefix='.'+path.name, delete=False) as handle:
            temporary = Path(handle.name)
        try:
            return cls._import_candidate(config, path, temporary, archive, rows, artifacts, source_hash)
        finally:
            temporary.unlink(missing_ok=True)

    @classmethod
    def _import_candidate(cls, config, path, temporary, archive, rows, artifacts, source_hash):
        with connect_sqlite(temporary) as db:
            db.execute('PRAGMA foreign_keys=ON')
            db.executescript('''
                CREATE TABLE state(schema TEXT NOT NULL, revision INTEGER NOT NULL, import_sha256 TEXT NOT NULL, archive TEXT NOT NULL, imported_at TEXT NOT NULL);
                CREATE TABLE sources(id TEXT PRIMARY KEY, position INTEGER UNIQUE NOT NULL, revision INTEGER NOT NULL, selection_status TEXT, data TEXT NOT NULL);
                CREATE TABLE operations(id TEXT PRIMARY KEY, position INTEGER UNIQUE NOT NULL, revision INTEGER NOT NULL, source_id TEXT, program_status TEXT, data TEXT NOT NULL);
                CREATE INDEX operations_by_source_status ON operations(source_id,program_status);
                CREATE TABLE history(sequence INTEGER PRIMARY KEY AUTOINCREMENT, revision INTEGER NOT NULL, entity TEXT NOT NULL, entity_id TEXT NOT NULL, actor TEXT NOT NULL, at TEXT NOT NULL, before_data TEXT, after_data TEXT NOT NULL);
                CREATE TRIGGER immutable_history_update BEFORE UPDATE ON history BEGIN SELECT RAISE(ABORT,'History is append-only'); END;
                CREATE TRIGGER immutable_history_delete BEFORE DELETE ON history BEGIN SELECT RAISE(ABORT,'History is append-only'); END;
                CREATE TABLE artifacts(path TEXT PRIMARY KEY, sha256 TEXT NOT NULL, bytes INTEGER NOT NULL);
                CREATE TRIGGER immutable_artifact_update BEFORE UPDATE ON artifacts BEGIN SELECT RAISE(ABORT,'Original artifact identity is immutable'); END;
                CREATE TRIGGER immutable_artifact_delete BEFORE DELETE ON artifacts BEGIN SELECT RAISE(ABORT,'Original artifact identity is immutable'); END;
                CREATE TABLE source_versions(source_id TEXT NOT NULL REFERENCES sources(id), snapshot_id TEXT NOT NULL, path TEXT NOT NULL, sha256 TEXT, first_revision INTEGER NOT NULL, PRIMARY KEY(source_id,snapshot_id,path));
                CREATE TRIGGER immutable_version_update BEFORE UPDATE ON source_versions BEGIN SELECT RAISE(ABORT,'Source versions are immutable'); END;
                CREATE TRIGGER immutable_version_delete BEFORE DELETE ON source_versions BEGIN SELECT RAISE(ABORT,'Source versions are immutable'); END;
            ''')
            now = datetime.now(timezone.utc).isoformat()
            db.execute('INSERT INTO state VALUES(?,?,?,?,?)', (SCHEMA, 1, source_hash, archive.name, now))
            for kind, entries in rows.items():
                for row in entries:
                    rec = row['record']; encoded = encode(rec)
                    cls._write_record(db, kind, row, 1, encoded)
                    db.execute('INSERT INTO history(revision,entity,entity_id,actor,at,after_data) VALUES(?,?,?,?,?,?)', (1, kind, row['id'], 'system1.migration', now, encoded))
            db.executemany('INSERT INTO artifacts VALUES(:path,:sha256,:bytes)', artifacts)
            cls._versions(db, rows['sources'], 1, config)
        store = cls(temporary)
        with store.connect() as db:
            if db.execute('PRAGMA quick_check').fetchone()[0] != 'ok': raise ValueError('Migration database failed integrity check')
        if digest(config['workbook'].read_bytes()) != source_hash or original_inventory(config['source_root'].resolve()) != artifacts:
            raise ValueError('Source state changed during migration; candidate was not activated')
        with temporary.open('r+b') as handle: os.fsync(handle.fileno())
        # Same-filesystem, exclusive publication: no observer can open a partial database.
        os.link(temporary, path)
        return cls(path)

    @staticmethod
    def _write_record(db, kind, row, revision, encoded):
        rec = row['record']
        if kind == 'sources':
            db.execute('INSERT INTO sources VALUES(?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET position=excluded.position,revision=excluded.revision,selection_status=excluded.selection_status,data=excluded.data',
                       (row['id'], row['position'], revision, rec.get('selection_status'), encoded))
        else:
            db.execute('INSERT INTO operations VALUES(?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET position=excluded.position,revision=excluded.revision,source_id=excluded.source_id,program_status=excluded.program_status,data=excluded.data',
                       (row['id'], row['position'], revision, rec.get('source_id'), rec.get('program_status'), encoded))

    @staticmethod
    def _versions(db, sources, revision, config=None):
        for row in sources:
            rec = row['record']
            # Planned identifiers on missing/paywalled sources are not acquired originals.
            if not rec.get('snapshot_id') or not rec.get('stored_filename') or rec.get('snapshot_status') != 'STORED' or not rec.get('content_hash'): continue
            path = str(Path(str(rec.get('folder_code') or '')) / rec['stored_filename'])
            identity = (row['id'], rec['snapshot_id'], path)
            old = db.execute('SELECT sha256 FROM source_versions WHERE source_id=? AND snapshot_id=? AND path=?', identity).fetchone()
            if old and old[0] != rec.get('content_hash'): raise ValueError('Existing source-version hash cannot be rewritten')
            if old is None and config is not None and rec.get('snapshot_status') == 'STORED':
                root = config['source_root'].resolve(); original = (root/path).resolve()
                if root not in original.parents or not original.is_file() or digest(original.read_bytes()) != rec.get('content_hash'):
                    raise ValueError('New stored source version must match its managed original')
                known = db.execute('SELECT sha256 FROM artifacts WHERE path=?', (path,)).fetchone()
                if known and known[0] != rec['content_hash']: raise ValueError('Original artifact identity cannot be overwritten')
                db.execute('INSERT OR IGNORE INTO artifacts VALUES(?,?,?)', (path, rec['content_hash'], original.stat().st_size))
            db.execute('INSERT OR IGNORE INTO source_versions VALUES(?,?,?,?,?)', (*identity, rec.get('content_hash'), revision))

    def snapshot(self):
        with self.connect() as db:
            db.execute('BEGIN')
            result = {'state': dict(db.execute('SELECT * FROM state').fetchone())}
            for kind in ('sources', 'operations'):
                result[kind] = [dict(id=r['id'], position=r['position'], revision=r['revision'], record=decode(r['data'])) for r in db.execute('SELECT * FROM '+kind+' ORDER BY position')]
            return result

    def revision(self):
        with self.connect() as db: return db.execute('SELECT revision FROM state').fetchone()[0]

    def business_workbook(self, config, data_only=False):
        """Fresh transactional cell adapter without Excel presentation reconstruction.

        This adapter retains the existing, validated domain operations. Only immutable
        header/formula layout is cached; every call reads a new consistent SQL snapshot.
        The full presentation template is used separately by the one-way exporter.
        """
        import human_operations as h
        snap=self.snapshot()
        raw=(self.path.parent/snap['state']['archive']).read_bytes()
        if digest(raw)!=snap['state']['import_sha256']:raise ValueError('Immutable migration template hash mismatch')
        layout=business_layout(raw,config['sheet_name'],config['header_row'],config['data_start_row'],h.human_sheet_name(config))
        wb=Workbook();wb.remove(wb.active);wb.epoch=layout['epoch']
        for kind, info in layout['sheets'].items():
            sheet=wb.create_sheet(info['name'])
            for key,col in info['headers'].items():sheet.cell(info['header'],col).value=info['labels'][key]
            for name,ref in info['tables']:sheet.add_table(Table(displayName=name,ref=ref))
            for entry in snap[kind]:
                for key,value in entry['record'].items():
                    cell=sheet.cell(entry['position'],info['headers'][key]);cell.number_format=info['formats'][key]
                    if not data_only and key in info['formulas']:
                        formula,origin=info['formulas'][key]
                        cell.value=Translator(formula,origin=origin).translate_formula(cell.coordinate)
                    else:cell.value=value
        wb._governance_revision=snap['state']['revision'];wb._governance_path=str(self.path)
        return wb

    def workbook(self, config, data_only=False):
        import human_operations as h
        import source_updater as u
        snap = self.snapshot()
        raw = (self.path.parent / snap['state']['archive']).read_bytes()
        if digest(raw) != snap['state']['import_sha256']: raise ValueError('Immutable migration template hash mismatch')
        wb = load_workbook(BytesIO(raw), data_only=data_only)
        for kind, sheet, headers in (
            ('sources', config['sheet_name'], u.workbook_headers(wb[config['sheet_name']], config['header_row'])),
            ('operations', h.human_sheet_name(config), h.operation_headers(wb[h.human_sheet_name(config)])),
        ):
            formulas = {}
            if kind == 'sources' and not data_only:
                for key in DERIVED & headers.keys():
                    for cells in wb[sheet].iter_rows(min_row=config['data_start_row'], min_col=headers[key], max_col=headers[key]):
                        cell = cells[0]
                        if cell.data_type == 'f':
                            formulas[key] = (cell.value, cell.coordinate)
                            break
            first_row = config['data_start_row'] if kind == 'sources' else h.HUMAN_DATA_START_ROW
            styles = {key:copy(wb[sheet].cell(first_row,col)._style) for key,col in headers.items()}
            template_last = wb[sheet].max_row
            for entry in snap[kind]:
                for key, value in entry['record'].items():
                    cell = wb[sheet].cell(entry['position'], headers[key])
                    if entry['position'] > template_last: cell._style = copy(styles[key])
                    if key in formulas:
                        formula, origin = formulas[key]
                        cell.value = Translator(formula, origin=origin).translate_formula(cell.coordinate)
                    else:
                        cell.value = value
        wb._governance_revision = snap['state']['revision']
        wb._governance_path = str(self.path)
        return wb

    def save(self, wb, config, expected_revision, actor='system1.local'):
        if getattr(wb, '_governance_path', None) != str(self.path): raise ValueError('Only database-bound state can be saved; Excel readback is forbidden')
        if getattr(wb, '_governance_revision', None) != expected_revision: raise ValueError('STALE: state changed after this view was loaded')
        rows = business_rows(wb, config)
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            current = db.execute('SELECT revision FROM state').fetchone()[0]
            if current != expected_revision: raise ValueError('STALE: governance revision changed')
            changes = []
            for kind, entries in rows.items():
                old = {r['id']: r for r in db.execute('SELECT * FROM '+kind)}
                if set(old) - {r['id'] for r in entries}: raise ValueError('Source records and operation history cannot be deleted')
                for row in entries:
                    encoded = encode(row['record']); prior = old.get(row['id'])
                    if prior is not None and row['position'] != prior['position']:
                        raise ValueError('Business identity positions cannot be reordered in the compatibility view')
                    if prior is None or encoded != prior['data'] or row['position'] != prior['position']:
                        changes.append((kind, row, encoded, prior['data'] if prior else None))
            if not changes: return current
            revision = current+1
            now = datetime.now(timezone.utc).isoformat()
            for kind, row, encoded, before in changes:
                self._write_record(db, kind, row, revision, encoded)
                db.execute('INSERT INTO history(revision,entity,entity_id,actor,at,before_data,after_data) VALUES(?,?,?,?,?,?,?)',
                           (revision, kind, row['id'], actor, now, before, encoded))
            self._versions(db, rows['sources'], revision, config)
            db.execute('UPDATE state SET revision=?', (revision,))
        wb._governance_revision = revision
        return revision

    def backup(self, destination):
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.exists(): raise FileExistsError(destination)
        with NamedTemporaryFile(dir=destination.parent, prefix='.'+destination.name, delete=False) as handle:
            temporary = Path(handle.name)
        try:
            with self.connect() as src, connect_sqlite(temporary) as dest: src.backup(dest)
            state = GovernanceStore(temporary).snapshot()['state']
            raw = (self.path.parent/state['archive']).read_bytes()
            if digest(raw) != state['import_sha256']: raise ValueError('Backup template hash mismatch')
            publish_bytes(destination.parent/state['archive'], raw)
            with temporary.open('r+b') as handle: os.fsync(handle.fileno())
            os.link(temporary, destination)
        finally: temporary.unlink(missing_ok=True)
        return destination
