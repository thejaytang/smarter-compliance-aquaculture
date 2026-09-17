"""Derived, replaceable Excel register. Review authority remains in Workflow/Canonical."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone
from hashlib import sha256
from ..platform_support import lock_file
from ..sqlite_support import operational_root
import json
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
import os
from pathlib import Path
import re
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SCHEMA = 'system2-requirement-workbook/1'
RENDER_VERSION = 21
NAME = 'System2_Requirement_Register.xlsx'
HEADERS = ['Original number', 'Requirement / content text', 'Classification', 'Content review',
           'Requirement review', 'Delivery status', 'Context / applicability', 'Source position',
           'Title', 'Criteria', 'Level', 'Notes / footnotes', 'Original text', 'Content confidence',
           'Requirement confidence', 'Decision origin', 'Review reasons', 'Unit ID', 'Unit version',
           'Document ID', 'Document revision', 'Generation', 'Snapshot ID', 'Source SHA256',
           'Policy revision', 'Canonical references', 'Structure / dependencies', 'Text part', 'Unit type', 'Last human decision', 'Parent content', 'Heading level', 'Reading order']
BLUE, TEAL, INK, PALE = '173B54', 'DDF1E9', '243747', 'F1F5F9'
HEADERS += ['Generated subitem label', 'Requirement parent ID', 'Counted Requirements',
            'Provisional subdivision policy', 'Exact accepted text span', 'Unassigned text decision']


def _text(value):
    if value is None:
        return ''
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else str(value)


def _cell(ws, row, col, value):
    # Explicit text prevents source expressions (=, +, -, @) becoming executable formulae.
    cell = ws.cell(row, col)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
    else:
        text = _text(value)
        text = re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]', lambda m: f'\\u{ord(m[0]):04x}', text)
        if len(text.encode('utf-16-le')) // 2 > 32767:
            raise ValueError('excel_cell_limit_exceeded')
        cell.value = text
        cell.data_type = 's'
    # NamedStyle copies the workbook-bound style tuple per cell. Reusing one
    # immutable definition avoids constructing/registering ~1.2m style objects
    # in the representative register; later number/fill changes stay local.
    cell.style = 'RegisterBody'
    return cell


def _chunks(value):
    # 15,000 Python characters also fit Excel's 32,767 UTF-16 code-unit limit.
    if not isinstance(value, str):
        return [value]
    return [value[i:i + 15000] for i in range(0, len(value), 15000)] or ['']


def _table(ws, row, headers, values, name):
    for col, header in enumerate(headers, 1):
        c = _cell(ws, row, col, header)
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.font = Font(name='Aptos', size=11, color='FFFFFF', bold=True)
    ws.row_dimensions[row].height = 34
    for n, record in enumerate(values, row + 1):
        for col, value in enumerate(record, 1):
            _cell(ws, n, col, value)
        import math
        lines=max((sum(max(1,math.ceil(len(line)/width)) for line in str(record[c] or '').split('\n')) for c,width in ((1,72),(8,32),(6,48),(29,24)) if c<len(record)),default=1)
        ws.row_dimensions[n].height = min(260,max(32,lines*15+12)) if len(headers)>10 else 48
    if values:
        table = Table(displayName=name, ref=f'A{row}:{get_column_letter(len(headers))}{row + len(values)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = f'B{row + 1}'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.print_title_rows = f'1:{row}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _banner(ws, title, subtitle, width=8):
    for row, text in [(1, title), (2, subtitle)]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        c = _cell(ws, row, 1, text)
        c.font = Font(name='Aptos', size=21 if row == 1 else 11, bold=row == 1, color=INK)
        ws.row_dimensions[row].height = 36 if row == 1 else 38


def build_workbook(records, documents, policy, registry_hash, event_cursor, fingerprint, registry_kind='workbook', sampling=None):
    from ..review.requirement_subdivision import count as requirement_count
    book = Workbook()
    book.add_named_style(NamedStyle(name='RegisterBody', font=Font(name='Aptos', size=11, color=INK),
                                   alignment=Alignment(vertical='top', wrap_text=True)))
    index = book.active
    index.title = 'Source Index'
    grouped = defaultdict(list)
    sources = {r['source_id']: r for r in records if r.get('selection_status') == 'INCLUDE'}
    for doc in documents:
        sid = doc['source']['source_id']
        grouped[sid].append(doc)
        sources.setdefault(sid, next((r for r in records if r['source_id'] == sid), doc['source']))
    generated = datetime.now(timezone.utc).isoformat(timespec='seconds')
    _banner(index, 'System2 · Requirement register',
            f'{len(sources)} sources · One worksheet per source · Snapshot {generated}')
    _cell(index, 4, 1, 'How to use')
    index.merge_cells('B4:H4')
    _cell(index, 4, 2, 'Click a source ID. Filter Delivery status = Available for accepted Requirements. Review and correct in the browser.')
    index.row_dimensions[4].height = 40
    _cell(index, 5, 1, 'Snapshot scope')
    index.merge_cells('B5:H5')
    _cell(index, 5, 2, 'Current INCLUDE sources plus retained processed sources. Not started means no parsing job. Incomplete sources retain pending scope.')
    index.row_dimensions[5].height = 38
    rows = []
    used = {'source index', 'read me'}
    for count, (sid, source) in enumerate(sorted(sources.items()), 1):
        name = re.sub(r'[\[\]:*?/\\]', '_', sid).strip("'")[:31] or 'Source'
        if name.lower() in used:
            name = name[:22] + '_' + sha256(sid.encode()).hexdigest()[:8]
        used.add(name.lower())
        ws = book.create_sheet(name)
        docs = grouped[sid]
        live = [d for d in docs if d.get('eligible')]
        units = sum(len(d['units']) for d in live)
        pending = sum(u.get('content_status') not in {'human_accepted','machine_accepted'} or
                      u.get('requirement_status') not in {'human_accepted','machine_accepted'} for d in live for u in d['units'])
        available = sum(requirement_count(d['published']) for d in live)
        complete = bool(live) and all(d.get('complete', False) for d in live)
        state = ', '.join(sorted({d['state'] for d in live})) if live else ('Not started' if source.get('selection_status') == 'INCLUDE' else 'Withdrawn')
        if docs and not live and source.get('selection_status') == 'INCLUDE':
            state = 'Current version not started / earlier result inactive'
        title = source.get('source_title') or sid
        _banner(ws, f'{sid} · {title}',
                f'{state} · {"Source complete" if complete else "Source incomplete"} · {available} available · {pending} pending units')
        for row, label, value in [(4, 'Selection / format', f"{source.get('selection_status', 'Unknown')} / {source.get('file_format', '')}"),
                                  (5, 'Source version', f"{source.get('snapshot_id', '')} | {source.get('content_hash', '')}"),
                                  (6, 'Source / evidence', source.get('official_url') or source.get('relative_path', '')),
                                  (7, 'Readback policy', f"Policy {policy['revision']} | Generated {generated} | Full cell text: formula bar or expand row")]:
            _cell(ws, row, 1, label)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            _cell(ws, row, 2, value)
            ws.row_dimensions[row].height = 28
        _cell(ws, 8, 1, 'Back to index').hyperlink = "#'Source Index'!A1"
        ws.merge_cells('B8:H8')
        _cell(ws, 8, 2, 'Program-managed view. Use the browser for decisions; downloaded copies do not update automatically.')
        values = []
        for doc in docs:
            by_id={u['id']:u for u in doc['units']}
            display=[]
            for source_unit in doc['units']:
                if source_unit.get('table_rows') and not source_unit.get('edits',{}).get('body'):
                    for row in source_unit['table_rows']:
                        display.append(dict(source_unit,original=dict(source_unit['original'],fields=dict(source_unit['original']['fields'],body=row['text'])),table_row=row['row']))
                else:display.append(source_unit)
            for unit in display:
                original = unit['original']
                from ..review.effective import resolve
                effective = resolve(unit, by_id)
                fields = effective['fields']
                available_row = unit['id'] in doc['published'] and doc.get('eligible') and not unit.get('evidence_only') and not unit.get('superseded_by')
                delivery = 'Available' if available_row else ('Inactive version / source' if not doc.get('eligible') else
                    'Superseded' if unit.get('superseded_by') else 'Evidence only' if unit.get('evidence_only') else 'Not delivered')
                classification = 'Coverage check' if unit['kind'] == 'coverage' else unit.get('classification', 'undetermined')
                refs = effective.get('references', [])
                related = effective.get('related_content', [])
                linked_refs = [r for item in related for r in item.get('references', [])]
                linked_context = []
                for item in related:
                    if item.get('role') not in {'context','reference'}:continue
                    linked_pages=sorted({r['page_index']+1 for r in item.get('references',[]) if r.get('page_index') is not None})
                    location='pages '+', '.join(map(str,linked_pages)) if linked_pages else 'bound original region'
                    text='\n'.join(_text(item.get('fields',{}).get(k)) for k in ('identifier','title','body','criteria','level','applicability','notes','context') if item.get('fields',{}).get(k))
                    linked_context.append('Linked '+item['role']+' · '+location+'\n'+text)
                canonical_hashes=list(dict.fromkeys(r['canonical_sha256'] for r in refs+linked_refs if r.get('canonical_sha256')))
                if not canonical_hashes:
                    canonical_hashes=list(dict.fromkeys(c['sha256'] for c in doc.get('canonical',[]) if c.get('sha256')))
                hierarchy=effective.get('hierarchy',{});parent=by_id.get(hierarchy.get('parent_id'),{})
                from ..domains.requirements.review_hierarchy import label as hierarchy_label
                parent_label=hierarchy_label(parent) if parent else ('Parsed source container' if str(hierarchy.get('parent_id','')).startswith('container:') else 'Unassigned')
                pages=sorted({r['page_index']+1 for r in refs if r.get('page_index') is not None})
                locations=list(dict.fromkeys(str(r.get('locator','')) for r in refs if r.get('locator')))
                positions=('Pages '+', '.join(map(str,pages)) if pages else (unit.get('chapter') or fields.get('context') or 'Original order'))+' | '+str(len(refs))+' evidence references'
                if locations and len(locations[0])<100:positions+=' | '+locations[0]
                linked_pages=sorted({r['page_index']+1 for r in linked_refs if r.get('page_index') is not None})
                if linked_pages:positions+=' | Linked original pages '+', '.join(map(str,linked_pages))
                reasons = {k: unit.get(k) for k in ('blockers','content_reasons','requirement_reasons') if unit.get(k)}
                if doc.get('issues') or doc.get('error'):
                    reasons['source_issues'] = [doc.get('issues'), doc.get('error')]
                record = [fields.get('identifier'), fields.get('body') or fields.get('title'), classification,
                    unit.get('content_status'), unit.get('requirement_status'), delivery,
                    '\n\n'.join([_text(fields[k]) for k in ('context','applicability') if fields.get(k)]+linked_context), positions,
                    fields.get('title'), fields.get('criteria'), fields.get('level'), fields.get('notes'),
                    original.get('fields', {}).get('body'), unit.get('content_confidence'), unit.get('requirement_confidence'),
                    _text({'content': 'human' if unit.get('content_human') else 'machine assessment',
                           'requirement': 'human' if unit.get('requirement_human') else 'machine assessment'}),
                    _text(reasons), unit['id'], unit.get('version'), doc['id'], doc['revision'], doc.get('generation', 0),
                    doc['source'].get('snapshot_id'), doc['source'].get('content_hash'), policy['revision'],
                    _text({'unit_id':unit['id'],'canonical_sha256':canonical_hashes,'origin':unit.get('provenance',{}).get('origin','parser')}),
                    _text({'kind':unit['kind'],'chapter':unit.get('chapter'),'member_count':len(unit.get('members',[])),'dependency_count':len(unit.get('dependencies',[])),'superseded_by':unit.get('superseded_by',[]),'table_assembly_id':effective.get('table_assembly_id') or (unit['id'] if unit.get('table_assembly') else None),'column_headers':effective.get('table_assembly',{}).get('column_headers',[]),
                           'related_content':[{k:item.get(k) for k in ('role','target_unit_id','version','references')} for item in related]})]
                subdivision = unit.get('requirement_subdivision')
                parent_id = doc['source']['source_id'] + ':' + unit['id']
                count_parent = not subdivision or subdivision['policy']['count_basis'] == 'parent'
                record += ['', effective['kind'], next((h.get('actor','')+' · '+h.get('action','')+' · '+h.get('note','') for h in reversed(doc.get('history',[])) if h.get('unit_id')==unit['id']), ''), parent_label if hierarchy else '', hierarchy.get('heading_level'), effective.get('reading_order_key',''),
                           '', parent_id if subdivision else '', int(available_row and count_parent),
                           {'schema':subdivision['schema'],'status':subdivision['status'],'revision':subdivision['revision'],**subdivision['policy']} if subdivision else '',
                           '', subdivision.get('remainder_reason','') if subdivision else '']
                if available_row and not count_parent:
                    record[5] = 'Parent context; subitems available'
                output_records = [record]
                for i, child in enumerate(subdivision['subitems'] if subdivision else [], 1):
                    child_record = list(record)
                    child_record[0], child_record[1] = child['original_number'], child['text']
                    child_record[5] = ('Included in parent; not counted' if count_parent else 'Available') if available_row else 'Subitems awaiting review'
                    child_record[12] = child['text']
                    child_record[17] = f"{unit['id']}:b{subdivision['revision']}:{i}"
                    child_record[28] = 'requirement_subitem'
                    child_record[33], child_record[35] = child['generated_label'], int(available_row and not count_parent)
                    child_record[37] = {k:child[k] for k in ('field','start','end','text')}
                    output_records.append(child_record)
                for output_record in output_records:
                    chunks = [_chunks(_text(v) if isinstance(v, (dict, list)) else v) for v in output_record]
                    parts = max(map(len, chunks))
                    for part in range(parts):
                        excel_row = [c[part] if len(c)>1 and part<len(c) else '' if len(c)>1 else c[0] for c in chunks]
                        excel_row[27] = (f'Table row {unit["table_row"]} · ' if 'table_row' in unit else '') + f'{part+1}/{parts}'
                        if part: excel_row[35] = 0
                        values.append(excel_row)
        if len(values) > 1048565:
            raise ValueError('source_exceeds_excel_row_limit:' + sid)
        _table(ws, 10, HEADERS, values, f'Source{count}')
        if not values:
            ws.merge_cells('A11:H11')
            _cell(ws, 11, 1, 'No extracted items. Start this source in the browser; this workbook does not start parsing.')
            ws.row_dimensions[11].height = 50
        widths = [22, 75, 23, 24, 24, 26, 50, 45, 35, 45, 18, 45, 65, 18, 18, 30, 45]
        for col in range(1, len(HEADERS) + 1):
            ws.column_dimensions[get_column_letter(col)].width = widths[col-1] if col <= len(widths) else 26
            if 18 <= col <= 28:
                ws.column_dimensions[get_column_letter(col)].hidden = True
        for row in ws.iter_rows(min_row=11, max_row=10 + len(values)):
            for col in (14, 15):
                row[col-1].number_format = '0.0%'
            if row[5].value == 'Available':
                row[5].fill = PatternFill('solid', fgColor=TEAL)
        ws.sheet_properties.tabColor = '318565' if complete else '6085A0'
        rows.append([sid, title, source.get('file_format', ''), state, 'Complete' if complete else 'Incomplete', units, pending, available])
    _table(index, 8, ['Source ID', 'Source title', 'Format', 'Processing state', 'Completeness', 'Content units', 'Pending units', 'Available Requirements'], rows, 'SourceIndex')
    for row, ws in enumerate(book.worksheets[1:], 9):
        index.cell(row, 1).hyperlink = f"#'{ws.title.replace(chr(39), chr(39)*2)}'!A1"
        index.cell(row, 1).font = Font(name='Aptos', size=11, color='146BB3', underline='single')
    for col, width in enumerate([22, 66, 14, 36, 20, 19, 19, 22], 1):
        index.column_dimensions[get_column_letter(col)].width = width
    guide = book.create_sheet('Read Me')
    _banner(guide, 'Reading this register', 'A source-based Excel carrier for System2 review results', 3)
    info = [('Authority', 'Original Canonical and versioned review decisions remain authoritative. Excel is a program-managed register, not a decision submission channel.'),
            ('Refresh', 'The running workbench synchronizes this file after state changes. The browser downloads the last generated snapshot and displays its event version. Close and reopen Excel to see changes; downloaded copies are snapshots.'),
            ('Evidence only', 'Retained overlap evidence cannot be delivered as an independent Requirement. Its content and review history remain visible.'),
            ('Available', 'Only currently published Requirements appear as Available. Candidates, context, coverage checks, superseded and inactive versions remain distinguishable.'),
            ('Provisional B subitems', 'Full parent content is retained. Generated labels are not original numbers. Sum Counted Requirements; do not count parent and children together. Each saved subdivision identifies its temporary parent/subitem counting rule, exact text spans and unassigned-text decision. Peer confirmation remains pending.'),
            ('Completeness', 'An available item does not imply its source is complete. See pending units, source state and the coverage-check rows for unfinished scope.'),
            ('Confidence', 'Blank scores are unknown. Percentages guide routing and are not document accuracy. Human confirmation does not rewrite machine scores.'),
            ('Long content', 'Expand rows or use the formula bar for long cells. Content above 15,000 characters continues across Text part rows; count by Document ID + Unit ID, not Excel rows.'),
            ('Provenance', 'Unhide columns R:AB for unit IDs, versions, snapshot/hash, Canonical references, relationships and text-part numbers. Prior generations and full decision history remain in the browser/database.'),
            ('Corrections', 'Browser-applied corrections appear in current text; Original text retains the source body. Excel changes are not imported and are replaced on refresh.'),
            ('Schema', SCHEMA), ('Policy', _text(policy)), ('Registry identity kind', registry_kind), ('Registry SHA256', registry_hash),
            ('Event cursor', event_cursor), ('Snapshot fingerprint', fingerprint), ('Generated UTC', generated)]
    for row, (key, value) in enumerate(info, 4):
        _cell(guide, row, 1, key)
        guide.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        _cell(guide, row, 2, value)
        guide.row_dimensions[row].height = 54
    guide.column_dimensions['A'].width = 26
    guide.column_dimensions['B'].width = 70
    guide.column_dimensions['C'].width = 45
    guide.sheet_view.showGridLines = False
    if sampling:
        qa=book.create_sheet('Weekly checks')
        _banner(qa,'Weekly original and classification checks','Monitoring only. Pending findings and short samples do not establish accuracy. Decisions are saved in the workbench.',8)
        headers=['Week','Stage','Record','Source','Original position / judgment','Status','Verdict','Reviewer',
            'Requested','Sampled','Evidence / note','Source SHA256','Item ID','Revision','Retained snapshot / history','Text part']
        rows=[]
        for batch in sampling['batches']:
            rows.append([batch['week'],batch['stage'],'Batch','','','See item states','','',batch['requested'],batch['sampled'],
                _text({'shortfall':max(0,batch['requested']-batch['sampled']),'missing_strata':batch['missing_strata'],'catalog_errors':batch['source_catalog_errors']}),'','',1,_text(batch),1])
        for item in sampling['items']:
            detail=_text(item)
            for part,text in enumerate(_chunks(detail),1):
                rows.append([item['week'],item['stage'],'Item' if part==1 else 'Continuation',item['source']['source_id'],item['title'],
                    item['status'],item['verdict'],item['actor'],'','',item['history'][-1]['note'] if item['history'] else '',
                    item['source']['content_hash'],item['id'],item['revision'],text,part])
        safe=[]
        for row in rows:
            chunks=[_chunks(v) for v in row]
            for index in range(max(map(len,chunks))):
                safe.append([c[index] if index<len(c) else '' for c in chunks])
        _table(qa,4,headers,safe,'WeeklySampling')
        import math
        for number,row in enumerate(safe,5):
            lines=max(sum(max(1,math.ceil(len(line)/56)) for line in str(row[col] or '').split('\n')) for col in (4,10))
            qa.row_dimensions[number].height=min(260,max(36,lines*15+12))
        for col,width in enumerate([15,12,16,16,46,22,18,22,14,14,60,28,28,14,70,12],1):
            qa.column_dimensions[get_column_letter(col)].width=width
    book.active = 0
    return book


def save_workbook(book, target):
    # Emit SDK-compatible CT_Font child order. Native Excel accepts both orders,
    # but Open XML schema validation requires size/color before the font name.
    raw = BytesIO()
    book.save(raw)
    order = 'b i strike condense extend outline shadow u vertAlign sz color name family charset scheme'.split()
    with ZipFile(raw) as source, ZipFile(target, 'w', ZIP_DEFLATED) as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'xl/styles.xml':
                root = etree.fromstring(data)
                for font in root.find('{*}fonts'):
                    font[:] = sorted(font, key=lambda child: order.index(etree.QName(child).localname))
                data = etree.tostring(root)
            dest.writestr(item, data)


def sync_workbook(store, records, registry_hash, target: Path, assert_current=lambda: None, registry_kind='workbook'):
    """Serialize exports, capture a coherent DB snapshot, and atomically replace changed views."""
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    state_root=operational_root(store.root)
    state_root.mkdir(parents=True,exist_ok=True)
    with (state_root / 'workbook.lock').open('a+b') as lock:
        lock_file(lock, blocking=True)
        marker = state_root / 'workbook.json'
        try: prior=json.loads(marker.read_text())
        except (OSError,ValueError): prior={}
        with store.connect() as db:
            db.execute('BEGIN')
            policy = store.policy(db)
            cursor = db.execute('SELECT COALESCE(MAX(sequence),0) FROM events').fetchone()[0]
            metadata=[r[0] for r in db.execute('SELECT data FROM documents ORDER BY rowid')]
            state_fingerprint=sha256(json.dumps([RENDER_VERSION,records,registry_hash,registry_kind,metadata,policy,cursor],sort_keys=True,default=str).encode()).hexdigest()
            if prior.get('state_fingerprint')==state_fingerprint and target.exists() and sha256(target.read_bytes()).hexdigest()==prior.get('sha256'):
                assert_current()
                return {**prior,'path':str(target),'status':'current'}
            if (target.parent / ('~$' + target.name)).exists():
                raise ValueError('Close the System2 Excel register before refreshing it. Reviews remain saved.')
            documents = [store._load(db, r[0]) for r in db.execute('SELECT id FROM documents ORDER BY rowid')]
            from ..review.weekly_sampling import exists
            sampling=({'batches':[json.loads(r[0]) for r in db.execute('SELECT data FROM sampling_batches ORDER BY week,stage')],
                'items':[json.loads(r[0]) for r in db.execute('SELECT data FROM sampling_items ORDER BY week,stage,rowid')]} if exists(db) else None)
        fingerprint = sha256(json.dumps([SCHEMA, RENDER_VERSION, records, documents, policy, cursor, registry_hash, registry_kind,sampling], sort_keys=True,
                                       ensure_ascii=False, default=str).encode()).hexdigest()
        marker = state_root / 'workbook.json'
        try:
            prior = json.loads(marker.read_text())
        except (OSError, ValueError):
            prior = {}
        assert_current()
        if prior.get('fingerprint') == fingerprint and target.exists() and sha256(target.read_bytes()).hexdigest() == prior.get('sha256'):
            return {**prior, 'path': str(target), 'status': 'current'}
        if (target.parent / ('~$' + target.name)).exists():
            raise ValueError('Close the System2 Excel register before refreshing it. Reviews remain saved.')
        book = build_workbook(records, documents, policy, registry_hash, cursor, fingerprint, registry_kind,sampling)
        fd, tempname = tempfile.mkstemp(prefix='.requirement-register-', suffix='.xlsx', dir=target.parent)
        os.close(fd)
        temp = Path(tempname)
        try:
            save_workbook(book, temp)
            book.close()
            assert_current()
            with temp.open('r+b') as handle:
                os.fsync(handle.fileno())
            if (target.parent / ('~$' + target.name)).exists():
                raise ValueError('Excel opened during synchronization. The previous snapshot and saved reviews are preserved.')
            temp.replace(target)
            result = {'schema': SCHEMA, 'state_fingerprint':state_fingerprint, 'fingerprint': fingerprint, 'sha256': sha256(target.read_bytes()).hexdigest(),
                      'event_cursor': cursor, 'policy_revision': policy['revision'], 'registry_sha256':registry_hash,
                      'registry_kind':registry_kind,'path': str(target), 'status': 'current'}
            marker_tmp = marker.with_suffix('.tmp')
            marker_tmp.write_text(json.dumps(result))
            marker_tmp.replace(marker)
            return result
        finally:
            temp.unlink(missing_ok=True)
