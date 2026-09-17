"""Bounded local metadata suggestions; never creates a material or review decision."""
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()[:1000]


class PageMetadata(HTMLParser):
    def __init__(self):
        super().__init__(); self.meta={}; self.title=[]; self.heading=[]; self.lines=[]; self.tags=[]; self.skip=0
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs); self.tags.append(tag)
        if tag in ('script','style','noscript','nav','footer','aside'): self.skip+=1
        if tag=='html' and attrs.get('lang'):self.meta['language']=attrs['lang']
        if tag=='meta': self.meta[(attrs.get('name') or attrs.get('property') or '').lower()]=attrs.get('content','')
        if tag in ('p','div','tr','br','h1','h2'): self.lines.append('\n')
    def handle_endtag(self, tag):
        if tag in ('script','style','noscript','nav','footer','aside'): self.skip=max(0,self.skip-1)
        if tag in self.tags: self.tags=self.tags[:len(self.tags)-1-self.tags[::-1].index(tag)]
        if tag in ('p','div','tr','h1','h2'): self.lines.append('\n')
    def handle_data(self, data):
        if self.skip:return
        if 'title' in self.tags:self.title.append(data)
        if 'h1' in self.tags:self.heading.append(data)
        self.lines.append(data)


def inspect(path, filename='', official_url='', retrieved=False):
    path=Path(path); suffix=path.suffix.lower(); meta={}; title=''; text=''; warnings=[];coverage={'units':'document','total_units':1,'inspected_units':1,'unreadable_units':[],'truncated':False}
    if suffix in ('.html','.htm'):
        page=PageMetadata(); page.feed(path.read_bytes().decode('utf-8',errors='replace'))
        meta=page.meta; title=clean(meta.get('og:title') or ' '.join(page.heading) or ' '.join(page.title));text=''.join(page.lines)
    elif suffix=='.pdf':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(path) as doc:
            raw=doc.get_metadata_dict();meta={k.lower():v for k,v in raw.items()};title=clean(meta.get('title'))
            parts=[]
            coverage.update(units='pages',total_units=len(doc),inspected_units=min(len(doc),300),truncated=len(doc)>300)
            for i in range(min(len(doc),300)):
                page=doc[i]
                try:
                    tp=page.get_textpage()
                    try:
                        part=tp.get_text_range();parts.append(part)
                        if not part.strip():coverage['unreadable_units'].append(i+1)
                    finally:tp.close()
                finally:page.close()
            text='\n'.join(parts)
            warnings.append('Basic inspection uses native text from up to 300 pages. It does not perform OCR or verify legal applicability.')
            if not text.strip():warnings.append('No readable text found. This file may require OCR; fill in missing information manually.')
    elif suffix=='.xlsx':
        # This is the uploaded original, never the generated source-register preview.
        from openpyxl import load_workbook
        with ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>80_000_000:raise ValueError('Spreadsheet is too large for basic inspection.')
        workbook=load_workbook(path,read_only=True,data_only=False)
        try:
            meta={'title':workbook.properties.title or '', 'publisher':workbook.properties.creator or ''};title=clean(meta['title']);parts=[];cells=0;size=0
            coverage.update(units='sheets',total_units=len(workbook.worksheets),inspected_units=0)
            for sheet in workbook.worksheets:
                coverage['inspected_units']+=1;sheet_text=[]
                for row in sheet.iter_rows(values_only=True):
                    line=' '.join(str(v) for v in row if v is not None);cells+=len(row);size+=len(line);sheet_text.append(line)
                    if cells>=50000 or size>=250000:coverage['truncated']=True;break
                part='\n'.join(sheet_text);parts.append(part)
                if not part.strip():coverage['unreadable_units'].append(sheet.title)
                if coverage['truncated']:break
            text='\n'.join(parts)
        finally:workbook.close()
    else:raise ValueError('Choose a PDF, HTML or XLSX original.')
    if re.search(r'^(?:access denied|just a moment|404(?: not found)?|403(?: forbidden)?|checking your browser)\b',title,re.I):raise ValueError('The URL returned an access or error page, not the source document. Upload the original or use its direct URL.')
    lines=[clean(line) for line in text[:100000].splitlines() if clean(line)]
    document_name=re.search(r'Document Name\s*:?\s*(.+?)\s+Document ID\b',text[:30000],re.I)
    if document_name:title=clean(document_name.group(1))
    if not title:title=next((line for line in lines[:35] if 8<len(line)<240),'')
    if not title:
        title=clean(Path(filename).stem) if filename else ''
        warnings.append('No document title was found. Check the suggested filename-based title.')
    publisher=clean(meta.get('publisher') or meta.get('citation_publisher') or meta.get('dc.publisher'))
    version=clean(meta.get('citation_publication_date') or meta.get('article:modified_time') or meta.get('dc.date'))
    embedded_version=re.search(r'\bVersion\s*:?\s*(V?\d+(?:\.\d+){1,3})\b',text[:30000],re.I)
    if not version and embedded_version:version=embedded_version.group(1)
    owner=re.search(r'The (.+?) (?:\([A-Z]+\) )?is the owner of this document',text[:30000])
    if not publisher and owner:publisher=clean(owner.group(1))
    for i,line in enumerate(lines[:120]):
        if not publisher:
            m=re.match(r'(?:Publisher|Issued by|Utgiver|Departement)\s*:?\s*(.*)',line,re.I)
            if m:publisher=clean(m.group(1) or (lines[i+1] if i+1<len(lines) else ''))
        if not version:
            m=re.match(r'(?:Version|Revision|Last amended|Sist endret|Versjon)\s*:?\s*(.*)',line,re.I)
            if m:
                candidate=clean(m.group(1) or (lines[i+1] if i+1<len(lines) else ''))
                if re.search(r'\d',candidate):version=candidate
    haystack=(title+' '+publisher+' '+' '.join(lines[:20])).lower()
    code='Z_Pending_Classification'
    if re.search(r'\b(?:aquaculture stewardship council|globalg\.a\.p|asc farm standard)\b',haystack):code='C_Certification_Scheme'
    elif re.search(r'\b(?:iso\s*\d|iec\s*\d|standards norway|standard norge)\b',haystack):code='B_Standards_Body'
    elif re.search(r'\b(?:regulation|forskrift|lov om|act relating|departement|european parliament)\b',haystack):code='A_Public_Authority'
    coverage['truncated']=coverage['truncated'] or len(text)>250000
    if not text.strip() and not coverage['unreadable_units']:coverage['unreadable_units']=[1]
    fields={'source_title':title,'issuer':publisher,'version':version,'folder_code':code,'file_format':'html' if suffix=='.htm' else suffix.lstrip('.')}
    from .source_analysis import analyze
    fields,inspection=analyze(fields,text[:250000],meta,coverage,official_url=official_url,retrieved=retrieved)
    return {'fields':fields,'inspection':inspection,'warnings':warnings,'method':inspection['method'],'requires_review':True}


if __name__=='__main__':
    try:
        request=json.load(sys.stdin);result=inspect(request['path'],request.get('filename',''),request.get('official_url',''),request.get('retrieved',False))
        print(json.dumps({'ok':True,'data':result},ensure_ascii=False))
    except Exception as exc:print(json.dumps({'ok':False,'error':str(exc)}))
