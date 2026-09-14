"""Build a separate scale/soak fixture without touching prior acceptance data."""
import hashlib
import json
from pathlib import Path
import sqlite3
import time

import material_fixture

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / 'workbench/runtime/round2-acceptance'


def build():
    material_fixture.TARGET = TARGET
    material_fixture.build()
    from reportlab.pdfgen import canvas
    from openpyxl import Workbook
    data = TARGET / 'Data/A_Public_Authority'
    pdf = data / 'TS004-001_scale.pdf'
    doc = canvas.Canvas(str(pdf))
    for page in range(1, 301):
        doc.drawString(35, 780, f'ENGINEERING SCALE FIXTURE / Page {page}')
        for line in range(24):
            doc.drawString(35, 750-line*24, f'Section {page}.{line+1}: retain original evidence and human corrections.')
        doc.showPage()
    doc.save()
    excel = data / 'TS005-001_scale.xlsx'
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Measurements')
    for row in range(1, 20001):
        sheet.append([f'R{row}C{col}' for col in range(1, 31)])
    book.create_sheet('Empty sheet')
    hidden = book.create_sheet('Hidden notes'); hidden.sheet_state = 'hidden'; hidden.append(['Retain this original scope'])
    book.save(excel)
    html = data / 'TS006-001_long.html'
    html.write_text('<!doctype html><html><body><h1>ENGINEERING LONG HTML</h1>' + ''.join(
        f'<h2>Chapter {i}</h2><p id="p{i}">Original paragraph {i}: maintain the source record.</p>'
        for i in range(1, 1001)) + '</body></html>')
    manifest = json.loads((TARGET/'manifest.json').read_text())
    with sqlite3.connect(TARGET/'governance/governance.sqlite') as db:
        template = json.loads(db.execute("SELECT data FROM sources WHERE id='TS003'").fetchone()[0])
        for position, (sid, path) in enumerate([('TS004',pdf),('TS005',excel),('TS006',html)], 20):
            record = json.loads(json.dumps(template)); fingerprint = hashlib.sha256(path.read_bytes()).hexdigest()
            snapshot = sid+'-001'; rel = 'A_Public_Authority/'+path.name
            values = dict(source_id=sid, snapshot_id=snapshot, source_title='ENGINEERING SCALE '+sid,
                          stored_filename=path.name, file_format=path.suffix[1:], content_hash=fingerprint)
            for key, value in values.items(): record[key] = ['value', value]
            db.execute('INSERT INTO sources VALUES(?,?,?,?,?)',(sid,position,1,'INCLUDE',json.dumps(record)))
            db.execute('INSERT INTO artifacts VALUES(?,?,?)',(rel,fingerprint,path.stat().st_size))
            db.execute('INSERT INTO source_versions VALUES(?,?,?,?,?)',(sid,snapshot,rel,fingerprint,1))
            manifest['sources'].append(dict(source_id=sid,path=str(path),sha256=fingerprint,provenance='synthetic_scale'))
    (TARGET/'manifest.json').write_text(json.dumps(manifest,indent=2))
    print(json.dumps({'fixture':str(TARGET),'pdf_pages':300,'excel_rows':20000,'excel_columns':30,'at':time.time()}))


if __name__ == '__main__':
    build()
