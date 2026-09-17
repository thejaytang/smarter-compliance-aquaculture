"""Real local launch, source/material/Requirement exchange and Git update.

Only synthetic sources and temporary Git repositories are used. Component
adapters run in the declared project environments; no domain service is mocked.
Windows runs the shipped .cmd launcher. Browser pointer/layout QA is separate.
"""
from contextlib import closing
from hashlib import sha256
from http.client import HTTPConnection
import json
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys
from tempfile import TemporaryDirectory
import time
import uuid
import zipfile

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / 'workbench/deployment'))
from initialize_local import initialize
from restore_initial_data import restore

TEXT = 'Trained staff must check equipment after a storm.'
ACTORS = ('Weijie Tang', 'Ana Jokic')


def run(args, cwd=ROOT, **kwargs):
    result = subprocess.run(list(map(str, args)), cwd=cwd, capture_output=True,
                            text=True, encoding='utf-8', timeout=180, **kwargs)
    if result.returncode:
        raise RuntimeError(f'{args[0]} failed: {result.stdout}\n{result.stderr}')
    return result.stdout.strip()


def python(component):
    return ROOT / component / '.venv' / ('Scripts/python.exe' if os.name == 'nt' else 'bin/python')


def git(cwd, *args):
    return run(['git', *args], cwd=cwd)


def initial_seed(root):
    source = root / 'seed-source'
    config = source / 'system1/Code/config'
    config.mkdir(parents=True)
    for name in ('config', 'schedule'):
        shutil.copy2(ROOT / 'system1/Code/config' / (name + '.example.json'), config / (name + '.example.json'))
    initialize(source)
    book = source / 'system1/Requirement_Source_Registry.xlsx'
    shutil.copy2(ROOT / 'system1/Code/tests/fixtures/source_registry.xlsx', book)
    original = source / 'system1/Data/C_Certification_Scheme/CS901-001_Synthetic.html'
    original.parent.mkdir(parents=True)
    original.write_text('<!doctype html><html lang="en"><head><title>Synthetic source</title></head>'
                        '<body><article><h1>Inspection</h1><p>' + TEXT + '</p></article></body></html>', encoding='utf-8')
    script = '''from pathlib import Path
import sys
from openpyxl import load_workbook
import source_updater as u
from system1.governance_store import GovernanceStore
p=Path(sys.argv[1]);w=load_workbook(p/'system1/Requirement_Source_Registry.xlsx');s=w['Source Register'];h=u.workbook_headers(s,2)
for k,v in {'stored_filename':'CS901-001_Synthetic.html','file_format':'html','content_hash':sys.argv[2]}.items():s.cell(3,h[k],v)
book=p/'system1/Requirement_Source_Registry.xlsx';u.save_workbook_atomic(w,book,book.stat().st_mtime_ns);w.close()
c=u.read_config(p/'system1/Code/config/config.json');GovernanceStore.migrate(c,c['governance_db'])
'''
    run([python('system1/Code'), '-c', script, source, sha256(original.read_bytes()).hexdigest()],
        env={**os.environ, 'PYTHONPATH': str(ROOT / 'system1/Code/src'), 'PYTHONUTF8': '1'})
    files = [book, original, source / 'system1/Code/runtime/governance.sqlite',
             source / 'system1/Code/runtime/governance-migration-input.xlsx']
    manifest = dict(schema='source-initial-data/1', files={
        p.relative_to(source).as_posix(): dict(bytes=p.stat().st_size, sha256=sha256(p.read_bytes()).hexdigest()) for p in files})
    archive = root / 'synthetic-initial-data.zip'
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
        for p in files: z.write(p, p.relative_to(source).as_posix())
        z.writestr('manifest.json', json.dumps(manifest))
    return archive, sha256(archive.read_bytes()).hexdigest()


class Peer:
    def __init__(self, root, actor):
        self.root, self.actor = root, actor
        self.state_file = root / 'workbench/runtime/server.json'
        self.csrf = self.cookie = ''
        self.port = None

    def launch(self):
        args = ['--no-browser', '--root', str(self.root / 'workbench'),
                '--system-root', str(ROOT / 'system1'), '--config', str(self.root / 'system1/Code/config/config.json')]
        if os.name == 'nt':
            command = ['cmd.exe', '/d', '/c', 'call', str(ROOT / 'workbench/deployment/Open Workbench.cmd'), *args]
        else:
            command = [python('workbench'), '-m', 'local_workbench', *args]
        run(command, env={**os.environ, 'PYTHONPATH': str(ROOT / 'workbench/src'), 'PYTHONUTF8': '1'})
        self.port = json.loads(self.state_file.read_text())['port']
        state = self.call('/api/state')
        self.csrf = state['csrf']
        self.call('/api/actor', dict(name=self.actor))
        assert len(state['sources']) == 2
        assert self.call('/health')['root'] == str((self.root / 'workbench').resolve())

    def call(self, path, body=None, raw=False):
        headers = {'Origin': f'http://127.0.0.1:{self.port}', 'X-CSRF-Token': self.csrf,
                   'Content-Type': 'application/zip' if isinstance(body, bytes) else 'application/json',
                   'X-Material-API-Version': '3'}
        if self.cookie: headers['Cookie'] = self.cookie
        client = HTTPConnection('127.0.0.1', self.port, timeout=120)
        payload = body if isinstance(body, bytes) else json.dumps(body) if body is not None else None
        try:
            client.request('GET' if body is None else 'POST', path, payload, headers)
            response = client.getresponse(); data = response.read()
            if response.getheader('Set-Cookie'): self.cookie = response.getheader('Set-Cookie').split(';')[0]
            if response.status != 200: raise AssertionError(f'{path}: {response.status} {data.decode()}')
            return data if raw else json.loads(data)
        finally: client.close()

    def save(self, path, **body):
        return self.call(path, dict(request_id=str(uuid.uuid4()), **body))

    def stop(self):
        if not self.state_file.exists(): return
        self.call('/api/stop', {})
        deadline = time.monotonic() + 45
        while self.state_file.exists() and time.monotonic() < deadline: time.sleep(.2)
        assert not self.state_file.exists(), 'Service did not release its stores after shutdown'

    def export(self):
        result = self.save('/api/sync/export')
        return self.call('/api/sync/download?id=' + result['id'], raw=True)

    def receive(self, raw):
        plan = self.call('/api/sync/import', raw)
        assert not plan['conflicts'], plan['conflicts']
        self.call('/api/sync/apply', dict(id=plan['id'], explicit_confirmation=True))


def fingerprints(root):
    files = [p for p in root.rglob('*') if p.is_file() and '.git' not in p.parts
             and (p.suffix in ('.sqlite', '.html') or p.name in ('config.json', 'schedule.json'))]
    for path in files:
        if path.suffix == '.sqlite':
            with closing(sqlite3.connect(path)) as db:
                assert db.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
                assert not db.execute('PRAGMA foreign_key_check').fetchall()
    return {p.relative_to(root).as_posix(): sha256(p.read_bytes()).hexdigest() for p in files}


def exercise(root):
    # A real local Git remote models the same app-only update applied by colleagues.
    remote = root / 'application.git'; git(root, 'init', '--bare', str(remote))
    producer = root / 'developer'; git(root, 'clone', str(remote), str(producer))
    git(producer, 'config', 'user.email', 'test@example.invalid'); git(producer, 'config', 'user.name', 'Integration test')
    git(producer, 'checkout', '-b', 'main')
    shutil.copy2(ROOT / '.gitignore', producer / '.gitignore')
    for rel in ('system1/Data/STORAGE.md', 'system1/Code/config/config.example.json', 'system1/Code/config/schedule.example.json'):
        target = producer / rel; target.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(ROOT / rel, target)
    (producer / 'application-version.txt').write_text('1')
    git(producer, 'add', '.'); git(producer, 'commit', '-m', 'Initial application'); git(producer, 'push', 'origin', 'main')
    archive, checksum = initial_seed(root)
    peers = []
    try:
        for index, actor in enumerate(ACTORS):
            target = root / ('peer ' + str(index))
            git(root, 'clone', '--branch', 'main', str(remote), str(target))
            restore(target, archive, checksum)
            peer = Peer(target, actor); peers.append(peer); peer.launch()
        a, b = peers
        material = a.save('/api/material/open', source_id='CS901')
        ref = dict(material['scope'][0]['location'], scope_id=material['scope'][0]['id'])
        material = a.save('/api/material/save', material_id=material['id'], expected_revision=material['revision'],
                          blocks=[dict(id='synthetic-passage', type='text', text=TEXT, source_refs=[ref])])['material']
        doc = a.save('/api/requirements/step', action='start', material_id=material['id'],
                     material_revision=material['revision'], block_id='synthetic-passage')['document']
        uid = next(iter(doc['units']))
        b.receive(a.export())
        received = b.call('/api/requirements/session?id=' + doc['id'])
        saved = b.save('/api/requirements/step', action='assign', session_id=doc['id'],
                       expected_revision=received['revision'], unit_id=uid, field='Subject', start=0, end=13)['document']
        context = b.call('/api/interpretations/context', dict(unit_id=uid))
        fields = {k: dict(value='', basis='unresolved', references=[], gaps=[]) for k in
                  ('scope', 'scope_information', 'condition', 'condition_information', 'demand', 'verification')}
        fields['verification']['value'] = 'Synthetic colleague evidence proposal'
        b.save('/api/interpretations/save', unit_id=uid, expected_revision=0,
               context_fingerprint=context['fingerprint'], fields=fields)
        a.receive(b.export())
        returned = a.call('/api/requirements/session?id=' + doc['id'])
        assert returned['units'][uid]['Subject'] == 'Trained staff'
        assert returned['edited_by'] == ACTORS[1]
        interpretation = a.call('/api/interpretations?unit_id=' + uid)
        assert interpretation['history'][0]['edited_by'] == ACTORS[1]
        assert interpretation['fields']['verification']['value'] == 'Synthetic colleague evidence proposal'
        assert interpretation['context']['requirement']['text'] == TEXT
        print('Real source/material/Requirement/interpretation round trip passed.', flush=True)
        for peer in peers: peer.stop()
        before = [fingerprints(p.root) for p in peers]
        assert all(git(p.root, 'status', '--porcelain', '--untracked-files=all') == '' for p in peers)
        (producer / 'application-version.txt').write_text('2')
        git(producer, 'add', '.'); git(producer, 'commit', '-m', 'Application update'); git(producer, 'push', 'origin', 'main')
        for peer, prior in zip(peers, before):
            git(peer.root, 'pull', '--ff-only')
            assert (peer.root / 'application-version.txt').read_text() == '2'
            assert fingerprints(peer.root) == prior, 'Application update changed local work'
            peer.launch()
            assert peer.call('/api/requirements/session?id=' + doc['id'])['edited_by'] == ACTORS[1]
            assert peer.call('/api/interpretations?unit_id=' + uid)['history'][0]['edited_by'] == ACTORS[1]
        return dict(status='passed', platform=sys.platform, peers=2, source_count=2,
                    real_component_adapters=True, initial_restore=True, native_windows_cmd=os.name == 'nt',
                    collaboration_round_trip=True, git_pull_preserved_local_files=True, restart_readback=True)
    finally:
        for peer in reversed(peers): peer.stop()


if __name__ == '__main__':
    with TemporaryDirectory(prefix='workbench-handoff-') as name:
        result = exercise(Path(name))
    print(json.dumps(result, indent=2))
