from pathlib import Path
from copy import copy
from tempfile import TemporaryDirectory
from unittest import TestCase
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.views import Selection
from system1.excel_output import prepare_output, normalize_font_order


class ExcelOutputTests(TestCase):
    def test_presentation_repairs_preserve_cells_fonts_and_other_package_parts(self):
        wb=Workbook();ws=wb.active;ws.title='Source Register'
        ws.append(['Generated snapshot']);ws.append(['source_id','stored_filename','selection_status'])
        filename='PA005-001_Regulation_on_the_technical_requirements_for_aquaculture_facilities_and_transport.html'
        ws.append(['PA005',filename,'=IF(1=1,"INCLUDE","PENDING")'])
        ws['B3'].font=Font(name='Carlito',size=11,bold=True,color='FF112233')
        ws.column_dimensions['B'].width=35;ws.row_dimensions[3].height=15
        ws.page_setup.horizontalDpi=ws.page_setup.verticalDpi=0
        ws.freeze_panes='B3';ws.sheet_view.selection.extend([Selection(pane='bottomRight'),Selection(pane='bottomRight')])
        values=[[c.value for c in row] for row in ws]
        prepare_output(wb,{'sheet_name':ws.title});height=ws.row_dimensions[3].height
        self.assertGreater(height,30);self.assertLessEqual(len(ws.sheet_view.selection),3)
        prepare_output(wb,{'sheet_name':ws.title});self.assertEqual(height,ws.row_dimensions[3].height)
        with TemporaryDirectory() as tmp:
            target=Path(tmp)/'view.xlsx';wb.save(target)
            with ZipFile(target) as z:before={n:z.read(n) for n in z.namelist()}
            normalize_font_order(target)
            with ZipFile(target) as z:after={n:z.read(n) for n in z.namelist()}
            self.assertEqual({k:v for k,v in before.items() if k!='xl/styles.xml'},
                             {k:v for k,v in after.items() if k!='xl/styles.xml'})
            def fonts(raw):
                return [sorted(ET.tostring(c).decode() for c in f) for f in ET.fromstring(raw).find('{*}fonts')]
            self.assertEqual(fonts(before['xl/styles.xml']),fonts(after['xl/styles.xml']))
            loaded=load_workbook(target);self.assertEqual([[c.value for c in row] for row in loaded.active],values)
            self.assertEqual(copy(loaded.active['B3'].font),copy(ws['B3'].font))
            loaded.close()
        wb.close()
