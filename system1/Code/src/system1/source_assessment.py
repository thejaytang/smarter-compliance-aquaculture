"""Optional machine source assessments, separate from named human decisions.

Existing HIGH/MEDIUM/LOW values are never converted into machine confidence.
Evidence rules are explicit source-version-bound assertions with validation IDs.
Absent rules and uncalibrated API proposals remain pending.
"""
from __future__ import annotations
from hashlib import sha256
import json
import math
from pathlib import Path
import sqlite3
from system1.sqlite_support import connect as connect_sqlite
import time

FIELDS = ('authority_quality','scope_relevance','version_currency','traceability','access_permission')


class Assessments:
    def __init__(self, database):
        self.database = Path(database)
        self.database.parent.mkdir(parents=True, exist_ok=True)
        with connect_sqlite(self.database) as db:
            db.executescript('''CREATE TABLE IF NOT EXISTS assessments(source_id TEXT PRIMARY KEY,data TEXT);
                CREATE TABLE IF NOT EXISTS assessment_events(sequence INTEGER PRIMARY KEY,data TEXT);
                CREATE TABLE IF NOT EXISTS assessment_policy(id INTEGER PRIMARY KEY CHECK(id=1),revision INTEGER,threshold REAL);
                CREATE TABLE IF NOT EXISTS assessment_holds(source_id TEXT PRIMARY KEY,actor TEXT,created REAL);''')
            db.execute('INSERT OR IGNORE INTO assessment_policy VALUES(1,1,.95)')

    def policy(self, revision, threshold):
        if type(threshold) not in (int,float) or not math.isfinite(threshold) or not 0 < threshold <= 1:
            raise ValueError('invalid_assessment_threshold')
        with connect_sqlite(self.database) as db:
            db.execute('BEGIN IMMEDIATE')
            old = db.execute('SELECT revision,threshold FROM assessment_policy').fetchone()
            if revision < old[0] or (revision == old[0] and threshold != old[1]):
                raise ValueError('stale_assessment_policy')
            if (revision, threshold) != old:
                db.execute('UPDATE assessment_policy SET revision=?,threshold=?', (revision,threshold))
                db.execute('INSERT INTO assessment_events(data) VALUES(?)',(json.dumps({'event_type':'policy','revision':revision,'threshold':threshold,'created':time.time()}),))
        return {'status':'applied','revision':revision}

    def hold(self, source_id, actor):
        if not actor:raise ValueError('named_draft_owner_required')
        with connect_sqlite(self.database) as db:
            db.execute('INSERT OR IGNORE INTO assessment_holds VALUES(?,?,?)',(source_id,actor,time.time()))
        return {'status':'applied'}

    def evaluate(self, source, evidence_rules, provider=None):
        values = {}
        for field in FIELDS:
            matching = [r for r in evidence_rules if r.get('field') == field
                        and r.get('source_id') == source['source_id']
                        and r.get('source_sha256') == source.get('content_hash')
                        and r.get('source_revision') == source.get('source_revision')]
            rule = matching[0] if len(matching) == 1 else {}
            value = {'rating': rule.get('rating','UNKNOWN'), 'confidence':rule.get('confidence'),
                     'evidence':rule.get('evidence',[]), 'calibration_version':rule.get('calibration_version'),
                     'method':rule.get('method','no_validated_rule')}
            # Invalid fields never inherit a positive rating or synthetic percentage.
            confidence = value['confidence']
            if value['rating'] not in {'HIGH','MEDIUM','LOW','UNKNOWN'} or (confidence is not None and (
                type(confidence) not in (int,float) or not math.isfinite(confidence) or not 0 <= confidence <= 1)):
                raise ValueError('invalid_validated_assessment_rule')
            values[field] = value
        if source.get('_snapshot_verified') and source.get('official_url') and source.get('snapshot_id'):
            values['traceability'] = {'rating':'HIGH','confidence':1.0,
                'evidence':[source['snapshot_id'],source['content_hash'],source['official_url']],
                'calibration_version':'exact-snapshot-hash-proof/1','method':'verified_local_snapshot_identity/1'}
        assessment = {'source_id':source['source_id'],'source_revision':source['source_revision'],
                      'source_sha256':source.get('content_hash'),'dimensions':values,'created':time.time(),
                      'mode':'NO_API','proposals':[]}
        if provider:
            try:
                # Optional proposals have no authority to modify evidence-calibrated ratings.
                assessment['proposals'] = provider.assess(source)
                assessment['mode'] = 'API_CONNECTED'
            except Exception as exc:
                assessment['provider_error'] = type(exc).__name__
                assessment['mode'] = 'NO_API_FALLBACK'
        assessment['id'] = sha256(json.dumps(assessment,sort_keys=True).encode()).hexdigest()
        with connect_sqlite(self.database) as db:
            db.execute('INSERT OR REPLACE INTO assessments VALUES(?,?)', (source['source_id'],json.dumps(assessment)))
            db.execute('INSERT INTO assessment_events(data) VALUES(?)',(json.dumps(assessment),))
        return assessment

    def project(self, sources, tasks, history):
        with connect_sqlite(self.database) as db:
            stored = {r[0]:json.loads(r[1]) for r in db.execute('SELECT source_id,data FROM assessments')}
            revision, threshold = db.execute('SELECT revision,threshold FROM assessment_policy').fetchone()
            held_ids={r[0] for r in db.execute('SELECT source_id FROM assessment_holds')}
        # A record of a named source decision takes precedence over machine proposals.
        human_ids = {h['source_id'] for h in history if h.get('operator') and h.get('operation_type') != 'RANDOM_QA_CHECK'}
        open_ids = {t['source_id'] for t in tasks if t.get('human_issue')} | held_ids
        for source in sources:
            a = stored.get(source['source_id'])
            if not a:
                continue
            source['machine_assessment'] = dict(a,policy_revision=revision,threshold=threshold)
            if source['source_id'] in human_ids or source.get('operator_selection_decision') in {'INCLUDE','EXCLUDE'}:
                source['selection_origin'] = 'human'
                continue
            current = a['source_revision'] == source.get('source_revision') and a['source_sha256'] == source.get('content_hash')
            valid = current and source['source_id'] not in open_ids and all(
                v['rating']=='HIGH' and v['confidence'] is not None and v['confidence'] >= threshold
                and v['evidence'] and v['calibration_version'] for v in a['dimensions'].values())
            valid = valid and all(source.get(k)==v for k,v in {
                'snapshot_status':'STORED','source_status':'CURRENT','download_status':'SUCCESS'}.items())
            valid = valid and all(source.get(k) for k in ('issuer','acquisition_channel','provenance_status','content_hash'))
            valid = valid and source.get('provenance_status') != 'UNVERIFIED'
            if valid:
                import source_updater
                proposed=dict(source,operator_selection_decision='INCLUDE',
                              **{field:value['rating'] for field,value in a['dimensions'].items()})
                valid=source_updater.selection_from_scores(proposed)=='INCLUDE'
            if valid:
                source['effective_selection']='INCLUDE'; source['selection_origin']='machine'
                source['assessment_id']=a['id']
        automated={s['source_id'] for s in sources if s.get('selection_origin')=='machine' and s.get('effective_selection')=='INCLUDE'}
        # The owning workbook retains unresolved compatibility rows. This live view
        # reopens the same task ID on threshold increase without creating duplicates.
        tasks[:]=[t for t in tasks if not (t['source_id'] in automated and t.get('operation_type')=='SELECTION_REVIEW'
                   and t.get('trigger')=='SELECTION_PENDING' and not t.get('human_issue') and not t.get('operator'))]
        return sources


class ConfiguredProvider:
    """An optional JSON evidence-assessment gateway, disabled without explicit config."""
    def __init__(self, config):
        self.config=config;self.calls=0

    def assess(self, source):
        import os
        import urllib.request
        from urllib.parse import urlsplit
        cfg=self.config
        if not cfg.get('enabled'):
            return []
        url=urlsplit(cfg['endpoint'])
        if url.scheme!='https' or url.username or url.password or not cfg.get('model'):
            raise ValueError('explicit_https_endpoint_and_model_required')
        if source['source_id'] not in cfg.get('allowed_source_ids',[]) or self.calls>=cfg.get('max_calls',0):
            raise ValueError('source_or_budget_not_authorized')
        allowed={'source_id','source_title','issuer','official_url','version','document_type','requirement_role',
                 'inclusion_rationale','applicability_reference','source_family','provenance_status'}
        evidence={k:source.get(k) for k in cfg.get('allowed_fields',[]) if k in allowed}
        body=json.dumps({'schema_version':'source-assessment-suggestions/1','model':cfg['model'],'evidence':evidence}).encode()
        if not evidence or len(body)>cfg.get('max_bytes',0):
            raise ValueError('evidence_not_authorized_or_too_large')
        headers={'Content-Type':'application/json'}
        if cfg.get('api_key_env'):
            key=os.environ.get(cfg['api_key_env'])
            if not key:raise ValueError('credential_not_configured')
            headers['Authorization']='Bearer '+key
        class NoRedirect(urllib.request.HTTPRedirectHandler):
            def redirect_request(self,*args,**kwargs):return None
        self.calls+=1
        request=urllib.request.Request(cfg['endpoint'],data=body,headers=headers,method='POST')
        with urllib.request.build_opener(NoRedirect).open(request,timeout=min(cfg.get('timeout',20),60)) as response:
            raw=response.read(1_000_001)
        if len(raw)>1_000_000:raise ValueError('response_size_limit')
        suggestions=json.loads(raw)['suggestions']
        if not isinstance(suggestions,list) or any(not isinstance(s,dict) or s.get('field') not in FIELDS
            or s.get('rating') not in {'HIGH','MEDIUM','LOW'} or not s.get('reason') or not s.get('evidence_fields')
            or set(s['evidence_fields'])-set(evidence) for s in suggestions):
            raise ValueError('invalid_or_unsupported_assessment_suggestion')
        return suggestions


def sync_confidence_sheet(workbook, database):
    """Program-managed machine evidence, kept separate from human source decisions."""
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    database=Path(database)
    if not database.exists():return
    with connect_sqlite(database) as db:
        assessments={r[0]:json.loads(r[1]) for r in db.execute('SELECT source_id,data FROM assessments')}
        policy=db.execute('SELECT revision,threshold FROM assessment_policy').fetchone()
    name='Machine confidence';ws=workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
    ws.delete_rows(1,ws.max_row)
    headers=['Source ID']+[f.replace('_',' ').title()+' confidence' for f in FIELDS]+['Calibration status','Policy revision','Threshold','Source SHA256','Assessment ID','Assessment method / evidence']
    ws.append(headers)
    source=workbook['Source Register']; cols={c.value:c.column for c in source[2]}
    for rr in range(3,source.max_row+1):
        sid=source.cell(rr,cols['source_id']).value
        if not sid:continue
        a=assessments.get(sid,{})
        import human_operations as h
        import source_updater as u
        current=bool(a and a.get('source_sha256')==source.cell(rr,cols['content_hash']).value and a.get('source_revision')==h.source_record_fingerprint(u.record_from_row(source,rr,cols)))
        dims=a.get('dimensions',{})
        scores=[dims.get(f,{}).get('confidence') if current else None for f in FIELDS]
        state='; '.join(f.replace('_',' ')+': '+(dims.get(f,{}).get('calibration_version') or 'Not calibrated') for f in FIELDS) if current else 'No current version-bound assessment'
        detail='\n'.join(f+': '+json.dumps(dims.get(f,{}),ensure_ascii=False) for f in FIELDS) if current else ''
        ws.append([sid,*scores,state,policy[0],policy[1],a.get('source_sha256'),a.get('id'),detail])
    from openpyxl.worksheet.views import Selection
    if ws.freeze_panes!='B2':ws.freeze_panes='B2'
    # Rebuilding the sheet must not accumulate extra selections on every export.
    ws.sheet_view.selection=[Selection(pane='topRight',activeCell='B1',sqref='B1'),
                             Selection(pane='bottomLeft',activeCell='A2',sqref='A2'),
                             Selection(pane='bottomRight',activeCell='B2',sqref='B2')]
    ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
    for col in range(1,len(headers)+1):
        ws.column_dimensions[get_column_letter(col)].width=23 if col<7 else 38
        if col>=10:ws.column_dimensions[get_column_letter(col)].hidden=True
    for row in ws:
        for c in row:c.font=Font(name='Aptos',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
        if row[0].row>1:
            for c in list(row)[1:6]+[row[8]]:c.number_format='0.0%'
    for c in ws[1]:c.font=Font(name='Aptos',size=11,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='173B54')
    ws.row_dimensions[1].height=45
