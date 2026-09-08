"""Browser-first workbook presentation; never changes source or review records."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

BROWSER_MARKER = "system1-browser-workbench"
GUIDE_MARKER = "system1-browser-guide-v1"
SECTIONS = {4: "START IN THE WORKBENCH", 11: "WHAT THIS WORKBOOK CONTAINS",
            18: "REVIEW STATUS", 24: "COMMON REVIEW ACTIONS", 32: "SOURCE RULES",
            37: "AUTOMATIC PROCESSING", 42: "RECOVERY AND HELP"}
INSTRUCTION_ROWS = {
    5: ("1. Open", "Project-root launcher", "Save and close Excel, then double-click Open Workbench.command in the project root. It opens the browser workbench."),
    6: ("2. Identify yourself", "Current reviewer", "Choose Ana Jokic, Daniel Restad or Weijie Tang once. Switch reviewer when another person takes over."),
    7: ("3. Review", "System1 → Pending review", "Choose a source from the left queue. Open its official link or original, then complete the action requested on screen."),
    8: ("4. Record", "Decisions apply automatically", "The last required rating submits a source assessment. Other tasks have their own save button. Do not run a program after each review."),
    9: ("5. Follow up", "Review history / Overview", "Completed reviews move to Review history. Unresolved issues stay pending. Overview contains the live Dashboard and weekly accuracy chart."),
    12: ("Instructions", "Read-only guidance", "This guide points to the browser workflow. Full instructions: workbench/USER_GUIDE.md in the Requirement Workstream."),
    13: ("Categories", "Controlled vocabulary", "Program-maintained definitions and dropdown values. These are reference data, not a human review form."),
    14: ("Source Register", "Canonical source state", "System1 manages source identities, snapshots, scores and selection here. Submit human changes in the browser, not directly in these cells."),
    15: ("Dashboard", "Retired; hidden", "The browser Overview replaces the Excel Dashboard. Its old formulas and charts remain hidden for compatibility; do not unhide it for daily work."),
    16: ("Human Operation Desktop", "Program-owned; hidden", "The browser replaces this editing surface. The hidden sheet still stores adapter staging and audit history. Do not delete it or edit its rows."),
    19: ("Pending review", "Action remains unresolved", "Includes source assessment, missing originals, reported issues and weekly checks. A previous score or draft is not a completed decision."),
    20: ("Submitted / Waiting", "Not applied yet", "The workbench shows application progress. If Excel is open or another task holds the lock, processing waits and resumes automatically."),
    21: ("Applied / Review history", "Recorded decision", "The reviewer, original note and result are preserved. An applied partial correction can still leave a source pending for other issues."),
    22: ("INCLUDE / PENDING / EXCLUDE", "Effective source selection", "Selection is separate from file storage and download success. Inclusion requires a valid snapshot and completed governance gates."),
    25: ("Source assessment", "Rate five dimensions", "Click H / M / L for each requested dimension. Light fill is the previous score; dark fill is your current choice. The fifth rating submits automatically."),
    26: ("Correct a link", "Save corrected links", "Open the suggested URL, correct the appropriate field and save. Saving a URL does not download or verify a new original. Reported issues require explicit verification."),
    27: ("Missing original", "Choose file", "Upload an authorised PDF, HTML or XLSX in the workbench. After the format check, confirm identity, completeness and permission, then submit."),
    28: ("Notes / Keep pending", "Preserve unresolved questions", "Notes save as drafts without completing a task. Explain the reason when keeping a source pending, removing it from scope or reporting an error."),
    29: ("Weekly random check", "Correct / Report an error", "Compare the sampled source with its original. An error creates a follow-up task; it does not count as a correct result."),
    30: ("New candidate", "Maintainer-assisted intake", "A dedicated browser candidate form is not connected yet. Ask the maintainer to handle the named-human acceptance; never use ordinary ratings as candidate acceptance."),
    33: ("Identity and retrieval", "Two distinct URLs", "official_url identifies the publisher page; retrieval_url identifies the downloadable original. Preserve the original format and authoritative language."),
    34: ("Existing snapshots", "Preserve valid originals", "Failed retrieval or replacement must retain the current valid snapshot and failure evidence. An authorised upload is required for restricted material; never bypass a paywall."),
    35: ("Downstream use", "Source eligibility is governed here", "Downstream parsing consumes the controlled source state and original. Stored files and source approval do not prove extraction accuracy or legal compliance."),
    38: ("Weekly QA", "Service-owned schedule", "Every Monday, sample 5 eligible items per connected system in Europe/Oslo time. After downtime, create only the current week's missing batch. Keep incomplete weeks blank in accuracy charts."),
    39: ("Browser decisions", "No manual rerun", "The running workbench applies submitted reviews automatically. Weekly QA is part of the service, not a Codex task. Closing the page does not stop the service."),
    40: ("Maintenance schedules", "Separate, optional", "Source retrieval and maintenance schedules require explicit setup by a maintainer. Opening the workbench does not start Full Source Check, downloads or parsing."),
    43: ("Workbook repair prompt", "Stop and preserve evidence", "Do not accept automatic repair. Keep the file and contact the maintainer to restore a verified backup under the normal locks."),
    44: ("Excel is open", "Save and close Excel", "Submitted decisions remain waiting. The service retries automatically when the workbook is available. Do not force a second writer."),
    45: ("Unconfirmed submission", "Retry submission", "Use the existing request in the workbench so its request ID is preserved. For an outdated task, reload and review the current source version."),
    46: ("Project documentation", "One entry point", "Start with the project-root README.md. Agent rules are in the root AGENTS.md; current System1 facts are in system1/PROJECT_STATE.md."),
}


def browser_mode(workbook):
    return BROWSER_MARKER in (workbook.properties.keywords or "")


def prepare_browser_workbook(workbook):
    if not browser_mode(workbook):
        return
    if GUIDE_MARKER not in (workbook.properties.keywords or "") or "Instructions" not in workbook:
        if "Instructions" in workbook:
            workbook.remove(workbook["Instructions"])
        sheet = workbook.create_sheet("Instructions", 0)
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        for column, width in (("A", 29), ("B", 33), ("C", 96)):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "C5"
        sheet.merge_cells("A1:C1")
        sheet["A1"] = "SYSTEM1 · SOURCE REGISTRY"
        sheet["A1"].font = Font(name="Calibri", size=20, bold=True, color="17374D")
        sheet.row_dimensions[1].height = 35
        sheet.merge_cells("A2:C2")
        sheet["A2"] = "Reference workbook · Review sources and use the Dashboard in the browser workbench"
        sheet["A2"].font = Font(name="Calibri", size=11, color="546779")
        sheet.row_dimensions[2].height = 25
        for row, title in SECTIONS.items():
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            cell = sheet.cell(row, 1, title)
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17374D")
            cell.alignment = Alignment(vertical="center", indent=1)
            sheet.row_dimensions[row].height = 25
        for row, values in INSTRUCTION_ROWS.items():
            sheet.row_dimensions[row].height = 47
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column, value)
                cell.font = Font(name="Calibri", size=11, bold=column == 1, color="253746")
                cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
                cell.fill = PatternFill("solid", fgColor="EEF3F7" if row % 2 else "FFFFFF")
        sheet.print_options.horizontalCentered = True
        sheet.print_area = "A1:C46"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=.25, right=.25, top=.3, bottom=.3)
        sheet.oddFooter.center.text = "System1 reference guide | &P / &N"
        workbook.properties.keywords = (workbook.properties.keywords or "") + ";" + GUIDE_MARKER
    for name in ("Dashboard", "Human Operation Desktop"):
        if name in workbook:
            workbook[name].sheet_state = "veryHidden"
            workbook[name].sheet_view.tabSelected = False
    # Refresh the launcher reference during the existing controlled save.
    workbook["Instructions"]["C5"] = INSTRUCTION_ROWS[5][2]
    workbook["Instructions"].sheet_state = "visible"
    workbook.active = workbook.sheetnames.index("Instructions")
    for sheet in workbook:
        sheet.sheet_view.tabSelected = sheet.title == "Instructions"
    for view in workbook.views:
        if view.firstSheet >= len(workbook.worksheets) or workbook.worksheets[view.firstSheet].sheet_state != "visible":
            view.firstSheet = 0
