import contextlib
import io
import uuid
from unittest.mock import patch
from openpyxl import load_workbook
from test_operator_journeys import OperatorJourneyTests
import human_operations as h
import source_updater as u
from system1 import workbench_bridge as bridge
from system1.cli import run_cycle


class WorkbenchBridgeTests(OperatorJourneyTests):
    def task(self, issue=False):
        wb=load_workbook(self.workbook);src=wb["Source Register"];sh=u.workbook_headers(src,2)
        src.cell(3,sh["scope_relevance"],"MEDIUM");src.cell(3,sh["operator_selection_decision"],"PENDING")
        if issue:
            ops=wb[h.HUMAN_SHEET];oh=h.operation_headers(ops)
            values=h.review_operation_values({"review_id":"bridge-issue","source_id":"PA001","source_title":"Existing",
              "trigger":"HUMAN_REPORTED_ISSUE","issue_codes":["HUMAN_REPORTED_ISSUE"],
              "human_reported_issue":{"reason":"Wrong link","action":"Verify","evidence_id":"fixture-1"}},
              u.record_from_row(src,3,sh),u.local_now({"timezone":"Europe/Oslo"}))
            h.append_operation(ops,oh,values)
        wb.save(self.workbook);wb.close()
        with contextlib.redirect_stdout(io.StringIO()):
            run_cycle(self.config_path,False)
        return bridge.read(self.config_path)["tasks"][0]

    def request(self, task, **extra):
        return {"request_id":str(uuid.uuid4()),"actor":"Ana","task_id":task["operation_id"],
                "revision":task["revision"],"source_revision":task["source_revision"],"action":"assess",
                "scores":{k:"HIGH" for k in u.SCORE_FIELDS},"note":"Human assessment",**extra}

    def test_bridge_roundtrip_and_repeated_request(self):
        req=self.request(self.task())
        with patch("source_updater.run_updates",side_effect=AssertionError("No downloads")):
            self.assertEqual(bridge.apply(self.config_path,req)["status"],"applied")
            before=self._source_record()
            self.assertEqual(before["manual_updated_by"],"Ana")
            self.assertEqual(before["operator_selection_decision"],"INCLUDE")
            self.assertEqual(bridge.apply(self.config_path,req)["status"],"applied")
            self.assertEqual(before,self._source_record())

    def test_review_apply_updates_formula_caches_without_opening_excel(self):
        req=self.request(self.task())
        self.assertEqual(bridge.apply(self.config_path,req)['status'],'applied')
        values=load_workbook(self.workbook,data_only=True)
        formulas=load_workbook(self.workbook,data_only=False)
        self.assertEqual(values['Source Register']['E3'].value,'INCLUDE')
        self.assertEqual(formulas['Source Register']['E3'].data_type,'f')
        values.close();formulas.close()
        snapshot=bridge.read(self.config_path)
        self.assertEqual(next(s for s in snapshot['sources'] if s['source_id']=='PA001')['effective_selection'],'INCLUDE')
        task=self.task(issue=True)
        exclude=self.request(task,action='exclude',note='Explicit exclusion in isolated test')
        self.assertEqual(bridge.apply(self.config_path,exclude)['status'],'applied')
        values=load_workbook(self.workbook,data_only=True)
        self.assertEqual(values['Source Register']['E3'].value,'EXCLUDE')
        values.close()

    def test_bridge_partial_score_preserves_human_issue(self):
        req=self.request(self.task(issue=True))
        self.assertEqual(bridge.apply(self.config_path,req)["status"],"applied")
        self.assertEqual(len(h.pending_human_reported_issues(self.config_path)),1)
        pending=bridge.read(self.config_path)["tasks"][0]
        self.assertTrue(pending["human_issue"])
        self.assertEqual(bridge.apply(self.config_path,self.request(pending,action="verify",issue_verified=True,note="Verified original"))["status"],"applied")
        self.assertEqual(h.pending_human_reported_issues(self.config_path),[])

    def test_bridge_rejects_old_revision_without_write(self):
        request=self.request(self.task())
        wb=load_workbook(self.workbook);ws=wb["Source Register"];hs=u.workbook_headers(ws,2)
        ws.cell(3,hs["version"],"new");wb.save(self.workbook);wb.close()
        before=self.workbook.read_bytes()
        with self.assertRaisesRegex(ValueError,"STALE"):bridge.apply(self.config_path,request)
        self.assertEqual(before,self.workbook.read_bytes())

    def test_bridge_waits_on_excel_and_requires_actor(self):
        request=self.request(self.task());request["actor"]=""
        with self.assertRaisesRegex(ValueError,"操作人"):bridge.apply(self.config_path,request)
        request["actor"]="Ana"
        lock=self.workbook.with_name("~$"+self.workbook.name);lock.write_text("open")
        before=self.workbook.read_bytes()
        with self.assertRaises(Exception):bridge.apply(self.config_path,request)
        self.assertEqual(before,self.workbook.read_bytes())
        lock.unlink()
        self.assertEqual(bridge.apply(self.config_path,request)["status"],"applied")

    def test_bridge_recovers_staged_request(self):
        request=self.request(self.task())
        with patch("system1.workbench_bridge.run_cycle",side_effect=RuntimeError("simulated stop")):
            with self.assertRaises(RuntimeError):bridge.apply(self.config_path,request)
        self.assertEqual(bridge.apply(self.config_path,request)["status"],"applied")
        self.assertEqual(self._source_record()["manual_updated_by"],"Ana")

    def weekly_qa_task(self):
        import json
        cfg=json.loads(self.config_path.read_text())
        cfg['random_qa']={'enabled':True,'schedule':'MONDAY_WITH_CATCH_UP','weekly_sample_size':5}
        self.config_path.write_text(json.dumps(cfg))
        h.generate_random_qa(self.config_path)
        return next(t for t in bridge.read(self.config_path)['tasks'] if t['operation_type']=='RANDOM_QA_CHECK')

    def test_weekly_qa_correct_applies_once_without_downloads(self):
        task=self.weekly_qa_task()
        req=self.request(task,action='qa',verdict='CORRECT',actor='Ana Jokic')
        with patch('source_updater.run_updates',side_effect=AssertionError('No downloads')):
            self.assertEqual(bridge.apply(self.config_path,req)['status'],'applied')
            self.assertEqual(bridge.apply(self.config_path,req)['status'],'applied')
        records=[t for t in bridge.read(self.config_path)['history'] if t['operation_id']==task['operation_id']]
        self.assertEqual(len(records),1)
        self.assertEqual(records[0]['operator'],'Ana Jokic')
        self.assertEqual(records[0]['decision'],'CORRECT')
        self.assertIn('batch_week',h.parse_payload(records[0]['payload_json']))

    def test_weekly_qa_error_survives_routine_as_correction_task(self):
        task=self.weekly_qa_task()
        req=self.request(task,action='qa',verdict='INCORRECT',actor='Daniel Restad',note='The issuer is wrong; verify the original.')
        self.assertEqual(bridge.apply(self.config_path,req)['status'],'applied')
        snapshot=bridge.read(self.config_path)
        self.assertTrue(any(t['source_id']==task['source_id'] for t in snapshot['tasks']))
        self.assertEqual(next(t for t in snapshot['history'] if t['operation_id']==task['operation_id'])['decision'],'INCORRECT')
        with contextlib.redirect_stdout(io.StringIO()):run_cycle(self.config_path,False)
        remaining=next(t for t in bridge.read(self.config_path)['tasks'] if t['source_id']==task['source_id'])
        self.assertEqual(remaining['human_issue']['reason'],req['note'])
        verified=self.request(remaining,action='verify',issue_verified=True,note='Checked and corrected the issuer against the original.')
        self.assertEqual(bridge.apply(self.config_path,verified)['status'],'applied')
        self.assertFalse(any(t['source_id']==task['source_id'] for t in bridge.read(self.config_path)['tasks']))

    def test_bridge_manual_file_validation_and_replay(self):
        task=self.task()
        upload_root=self.root/"uploads";upload_root.mkdir()
        supplied=upload_root/"replacement.html"
        supplied.write_bytes(self.current_bytes.replace(b"current aquaculture",b"reviewed aquaculture"))
        request=self.request(task,action="manual",identity_verified=True,permission_verified=True,
            upload_path=str(supplied),upload_root=str(upload_root),upload_hash=u.sha256_file(supplied))
        self.assertEqual(bridge.apply(self.config_path,request)["status"],"applied")
        source=self._source_record()
        self.assertEqual(source["content_hash"],u.sha256_file(supplied))
        self.assertEqual(source["manual_updated_by"],"Ana")
        snapshot=source["snapshot_id"]
        self.assertEqual(bridge.apply(self.config_path,request)["status"],"applied")
        self.assertEqual(self._source_record()["snapshot_id"],snapshot)



for _name in list(OperatorJourneyTests.__dict__):
    if _name.startswith("test_"):setattr(WorkbenchBridgeTests,_name,None)
del OperatorJourneyTests
