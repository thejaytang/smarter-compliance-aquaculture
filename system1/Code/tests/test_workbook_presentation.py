"""Presentation migration must not rewrite governed data or lose cached results."""
import shutil
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

import source_updater as updater
from system1.workbook_presentation import prepare_browser_workbook


def business_content(workbook):
    return {sheet.title: [[(cell.value, cell.data_type) for cell in row]
                           for row in sheet.iter_rows()]
            for sheet in workbook if sheet.title != "Instructions"}


class BrowserWorkbookPresentationTest(unittest.TestCase):
    def test_controlled_save_preserves_records_and_retires_legacy_surfaces(self):
        source = Path(__file__).resolve().parents[2] / "Requirement_Source_Registry.xlsx"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / source.name
            shutil.copy2(source, path)
            book = load_workbook(path)
            before = business_content(book)
            tables = {sheet.title: list(sheet.tables) for sheet in book}
            chart_count = len(book["Dashboard"]._charts)
            # Exercise migration even after the production template has migrated.
            book.properties.keywords = "system1-browser-workbench"
            book["Instructions"]["C6"] = "Edit yellow cells in Human Operation Desktop"
            book["Dashboard"].sheet_state = "visible"
            book["Human Operation Desktop"].sheet_state = "visible"
            book.active = book.sheetnames.index("Dashboard")
            updater.save_workbook_atomic(book, path, path.stat().st_mtime_ns)
            book.close()
            for attempt in range(2):
                book = load_workbook(path)
                self.assertEqual(business_content(book), before)
                self.assertEqual({s.title: list(s.tables) for s in book}, tables)
                self.assertEqual(len(book["Dashboard"]._charts), chart_count)
                self.assertEqual(book.active.title, "Instructions")
                self.assertEqual([s.title for s in book if s.sheet_state == "visible"],
                                 ["Instructions", "Categories", "Source Register"])
                self.assertEqual(book["Dashboard"].sheet_state, "veryHidden")
                self.assertEqual(book["Human Operation Desktop"].sheet_state, "veryHidden")
                self.assertIn("Current reviewer", book["Instructions"]["B6"].value)
                self.assertIn("Do not run a program", book["Instructions"]["C8"].value)
                cached = load_workbook(path, data_only=True)
                formulas = [(s.title, c.coordinate) for s in book for row in s for c in row
                            if c.data_type == "f"]
                self.assertTrue(formulas)
                for sheet, cell in formulas:
                    self.assertNotEqual(cached[sheet][cell].data_type, "e", (sheet, cell))
                # Empty-string formulas legitimately read as None in openpyxl.
                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                with ZipFile(path) as archive:
                    nodes = [c for name in archive.namelist()
                             if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                             for c in ET.fromstring(archive.read(name)).iter(ns + "c")
                             if c.find(ns + "f") is not None]
                    self.assertEqual(len(nodes), len(formulas))
                    self.assertTrue(all(c.find(ns + "v") is not None for c in nodes))
                cached.close()
                if attempt == 0:
                    # A later compatibility writer cannot expose the old UI again.
                    book["Dashboard"].sheet_state = "visible"
                    book["Human Operation Desktop"].sheet_state = "visible"
                    book.active = book.sheetnames.index("Human Operation Desktop")
                    updater.save_workbook_atomic(book, path, path.stat().st_mtime_ns)
                book.close()

    def test_unmarked_legacy_fixture_is_not_migrated(self):
        book = Workbook()
        book.active.title = "Instructions"
        book.active["A1"] = "Legacy fixture"
        book.create_sheet("Dashboard")
        book.create_sheet("Human Operation Desktop")
        book.active = 1
        prepare_browser_workbook(book)
        self.assertEqual(book["Instructions"]["A1"].value, "Legacy fixture")
        self.assertEqual(book.active.title, "Dashboard")
        self.assertTrue(all(s.sheet_state == "visible" for s in book))
        book.close()
