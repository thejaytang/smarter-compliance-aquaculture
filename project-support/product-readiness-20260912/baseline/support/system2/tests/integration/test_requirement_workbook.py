from copy import deepcopy
from hashlib import sha256
from openpyxl import load_workbook
import pytest
from pdf_extraction.delivery.requirement_workbook import NAME, build_workbook, sync_workbook
from .test_two_stage_workflow import system, current, decide, threshold


def export(system, tmp_path, records=None):
    store = system[0]
    records = records if records is not None else [dict(current(system)['source'], source_title='Source for testing')]
    path = tmp_path / NAME
    result = sync_workbook(store, records, 'registry', path)
    return result, load_workbook(path), path


def test_source_sheets_include_unstarted_and_preserve_no_jobs(system, tmp_path):
    records = [dict(current(system)['source']), {'source_id':'NEXT', 'selection_status':'INCLUDE', 'source_title':'Not started'},
               {'source_id':'NO', 'selection_status':'EXCLUDE'}]
    _, wb, _ = export(system, tmp_path, records)
    assert 'NEXT' in wb.sheetnames and 'NO' not in wb.sheetnames
    assert wb['NEXT']['A11'].value.startswith('No extracted items')
    assert 'Not started' in wb['NEXT']['A2'].value
    assert len(system[0].snapshot()['documents']) == 1
    wb.close()


def test_source_version_kind_and_hash_are_exported_without_review_events(system,tmp_path):
    before=current(system)
    records=[dict(before['source'],source_title='First checked title')]
    first=sync_workbook(system[0],records,'a'*64,tmp_path/NAME,registry_kind='sqlite_snapshot')
    records[0]['source_title']='Corrected source title'
    second=sync_workbook(system[0],records,'b'*64,tmp_path/NAME,registry_kind='sqlite_snapshot')
    assert second['event_cursor']==first['event_cursor']
    assert second['sha256']!=first['sha256'] and second['registry_sha256']=='b'*64
    assert second['registry_kind']=='sqlite_snapshot' and current(system)==before
    book=load_workbook(tmp_path/NAME)
    values=[(row[0],row[1]) for row in book['Read Me'].iter_rows(values_only=True) if len(row)>1]
    assert ('Registry identity kind','sqlite_snapshot') in values
    assert ('Registry SHA256','b'*64) in values
    assert 'Corrected source title' in book.worksheets[1]['A1'].value
    book.close()


def test_gate_delivery_then_correction_is_reflected(system, tmp_path):
    threshold(system, .98, .98)
    _, wb, path = export(system, tmp_path)
    ws = wb.worksheets[1]
    assert not any(r[5].value == 'Available' for r in list(ws.rows)[10:])
    wb.close()
    decide(system, 'accept_content', 'coverage')
    decide(system, 'accept_content')
    decide(system, 'classify', classification='requirement')
    _, wb, path = export(system, tmp_path)
    assert any(r[5].value == 'Available' for r in list(wb.worksheets[1].rows)[10:])
    wb.close()
    decide(system, 'correct', fields={'body':'Fish farms shall retain records.'}, note='Corrected source transcription')
    _, wb, _ = export(system, tmp_path)
    rows = list(wb.worksheets[1].rows)[10:]
    assert not any(r[5].value == 'Available' for r in rows)
    changed = next(r for r in rows if r[1].value == 'Fish farms shall retain records.')
    assert changed[12].value == 'Fish farms shall keep records.'
    wb.close()


def test_unchanged_refresh_and_excel_lock_preserve_file(system, tmp_path):
    first, wb, path = export(system, tmp_path)
    wb.close()
    stat = path.stat().st_mtime_ns
    second, wb, _ = export(system, tmp_path)
    wb.close()
    assert first['sha256'] == second['sha256'] and path.stat().st_mtime_ns == stat
    decide(system, 'draft', draft={'note':'unfinished'})
    lock = path.parent / ('~$' + path.name)
    lock.touch()
    with pytest.raises(ValueError, match='Close the System2'):
        export(system, tmp_path)
    assert sha256(path.read_bytes()).hexdigest() == first['sha256']
    lock.unlink()
    changed, wb, _ = export(system, tmp_path)
    wb.close()
    assert changed['fingerprint'] != first['fingerprint']


def test_literal_long_text_and_invalid_sheet_names(system, tmp_path):
    doc = deepcopy(current(system))
    text = '=not_a_formula()' + '😀' * 35000
    doc['source']['source_id'] = 'Invalid / Source: with [characters] and long name'
    doc['units'][1]['original']['fields']['body'] = text
    wb = build_workbook([dict(doc['source'])], [doc], {'revision':1}, 'hash', 0, 'f')
    path = tmp_path / 'long.xlsx'
    wb.save(path)
    wb.close()
    wb = load_workbook(path)
    ws = wb.worksheets[1]
    assert len(ws.title) <= 31 and '/' not in ws.title
    rows = [r for r in list(ws.rows)[10:] if r[17].value == doc['units'][1]['id']]
    assert ''.join(r[1].value or '' for r in rows) == text
    assert all(r[1].data_type == 's' for r in rows)
    wb.close()


def test_withdrawal_keeps_source_sheet_but_no_available_items(system, tmp_path):
    decide(system, 'accept_content', 'coverage')
    decide(system, 'accept_content')
    decide(system, 'classify', classification='requirement')
    system[0].reconcile_sources({})
    source = dict(current(system)['source'], selection_status='EXCLUDE')
    _, wb, _ = export(system, tmp_path, [source])
    ws = wb.worksheets[1]
    assert 'Withdrawn' in ws['A2'].value
    assert all(r[5].value == 'Inactive version / source' for r in list(ws.rows)[10:])
    wb.close()


def test_source_date_values_and_failed_atomic_refresh(system, tmp_path, monkeypatch):
    from datetime import datetime
    import pdf_extraction.delivery.requirement_workbook as exporter
    records = [dict(current(system)['source'], retrieved_at=datetime(2026, 9, 9))]
    first, wb, path = export(system, tmp_path, records)
    wb.close()
    decide(system, 'draft', draft={'note':'new draft'})
    def broken(*args):
        raise OSError('simulated disk failure')
    monkeypatch.setattr(exporter, 'save_workbook', broken)
    with pytest.raises(OSError, match='disk failure'):
        export(system, tmp_path, records)
    assert sha256(path.read_bytes()).hexdigest() == first['sha256']
    assert not list(tmp_path.glob('.requirement-register-*.xlsx'))


def test_excel_opened_during_generation_preserves_previous_snapshot(system, tmp_path, monkeypatch):
    import pdf_extraction.delivery.requirement_workbook as exporter
    first, wb, path = export(system, tmp_path)
    wb.close()
    marker = system[0].root / 'workbook.json'
    old_marker = marker.read_bytes()
    decide(system, 'draft', draft={'note':'saved while output is unavailable'})
    save = exporter.save_workbook
    lock = path.parent / ('~$' + path.name)
    def opened_during_save(book, target):
        save(book, target)
        lock.touch()
    monkeypatch.setattr(exporter, 'save_workbook', opened_during_save)
    with pytest.raises(ValueError, match='Excel opened during'):
        export(system, tmp_path)
    assert sha256(path.read_bytes()).hexdigest() == first['sha256']
    assert marker.read_bytes() == old_marker
    assert not list(tmp_path.glob('.requirement-register-*.xlsx'))
    lock.unlink()
    monkeypatch.setattr(exporter, 'save_workbook', save)
    recovered, wb, _ = export(system, tmp_path)
    wb.close()
    assert recovered['event_cursor'] > first['event_cursor']


def test_policy_reassessment_and_human_confirmation_export(system, tmp_path):
    first, wb, _ = export(system, tmp_path)
    assert wb.worksheets[1]['F12'].value == 'Available'
    wb.close()
    threshold(system, .98, .98)
    changed, wb, _ = export(system, tmp_path)
    assert wb.worksheets[1]['F12'].value == 'Not delivered'
    assert changed['policy_revision'] > first['policy_revision']
    wb.close()
    decide(system, 'accept_content', 'coverage')
    decide(system, 'accept_content')
    decide(system, 'classify', classification='requirement')
    threshold(system, 1, 1)
    _, wb, _ = export(system, tmp_path)
    assert wb.worksheets[1]['F12'].value == 'Available'
    assert wb.worksheets[1]['N12'].value == .96
    wb.close()


def test_overlap_evidence_is_visible_but_never_available(system):
    doc=deepcopy(current(system));u=next(u for u in doc['units'] if u['id']=='clause');u['evidence_only']=True
    # Even an inconsistent old publication must not promote a retained evidence copy.
    doc['published'][u['id']]={'fields':u['original']['fields']}
    b=build_workbook([doc['source']],[doc],{'revision':1},'registry',0,'evidence-copy')
    rows=[r for r in b.worksheets[1].iter_rows(min_row=11,values_only=True) if r[17]=='clause']
    assert rows and all(r[5]=='Evidence only' for r in rows)
    assert any(r[1]=='Fish farms shall keep records.' for r in rows)
    b.close()
