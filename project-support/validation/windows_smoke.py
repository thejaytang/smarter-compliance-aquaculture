"""Local Windows checkout diagnostics; no business writes or remote calls."""
import ast
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import winreg

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT / 'workbench'), str(ROOT / 'workbench/backend/application')]
os.environ['PYTHONPYCACHEPREFIX'] = str(ROOT / 'workbench/runtime/cache/python')
report = {'platform': sys.platform, 'checks': {}}

def check(name, fn):
    try:
        report['checks'][name] = {'status': 'passed', 'details': fn()}
    except Exception as exc:
        report['checks'][name] = {'status': 'failed', 'error': str(exc)}

def syntax():
    files = subprocess.check_output(['git', 'ls-files', '-z'], cwd=ROOT).decode().split('\0')
    counts = {'python': 0, 'json': 0}
    for name in files:
        path = ROOT / name
        if name.endswith('.py'):
            ast.parse(path.read_bytes(), filename=name)
            counts['python'] += 1
        elif name.endswith('.json'):
            json.loads(path.read_text(encoding='utf-8'))
            counts['json'] += 1
    return counts

def seed():
    from local_workbench.workspace_package import validate
    raw = (ROOT / 'workbench/initial-data/workspace-20260917.zip').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == 'a6978fb67ff00bdf55d9a99e1a1488f86ebdca1e3e55c4d0a618e1ec91774a75'
    _, manifest = validate(raw, allow_delivery=True)
    return {'files': len(manifest['files']), 'databases': len(manifest['databases']),
            'originals': len(manifest['originals']), 'bindings': len(manifest['bindings'])}

def native_pdf():
    import numpy as np
    import cv2
    import pypdfium2 as pdfium
    from pypdf import PdfWriter
    from lxml import etree
    import openpyxl
    import jsonschema
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pdfium.PdfDocument(buffer.getvalue()) as doc:
        page = doc[0]
        bitmap = page.render(scale=1)
        try:
            assert bitmap.to_pil().size == (72, 72)
        finally:
            bitmap.close()
            page.close()
    assert cv2.cvtColor(np.zeros((2, 2, 3), dtype=np.uint8), cv2.COLOR_BGR2GRAY).shape == (2, 2)
    assert etree.fromstring(b'<root/>').tag == 'root'
    book = openpyxl.Workbook()
    book.active['A1'] = 'Windows smoke'
    excel = io.BytesIO()
    book.save(excel)
    book.close()
    excel.seek(0)
    loaded = openpyxl.load_workbook(excel)
    assert loaded.active['A1'].value == 'Windows smoke'
    loaded.close()
    return 'PDF render, OpenCV, NumPy, lxml, Excel roundtrip passed'

check('tracked_source_syntax', syntax)
check('initial_package_validation', seed)
check('native_dependencies', native_pdf)
with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r'SYSTEM\CurrentControlSet\Control\FileSystem') as key:
    report['long_paths_enabled'] = winreg.QueryValueEx(key, 'LongPathsEnabled')[0]
report['optional_tools'] = {name: shutil.which(name) for name in ('tesseract', 'pdftoppm', 'node', 'npm')}
report['layout_ready'] = (ROOT / 'workbench/runtime/state/layout.json').exists()
report['automated_suite_included'] = (ROOT / 'workbench/tests/run_checks.py').exists()
target = Path(__file__).with_name('windows-smoke-results.json')
target.write_text(json.dumps(report, indent=2), encoding='utf-8')
print(json.dumps(report, indent=2))
