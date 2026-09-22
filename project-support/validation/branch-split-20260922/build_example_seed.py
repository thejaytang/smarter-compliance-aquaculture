"""Build and verify a fresh PE002 seed using owning migration/import services."""
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
import sys
import uuid
from local_workbench.server import Application
from local_workbench.full_snapshot import FullSnapshot
from local_workbench.workspace_migration import migrate
from local_workbench.workspace_package import capture, validate
from backend.system3.requirements import Requirements
from backend.system3.interpretations import Interpretations

ROOT = Path(__file__).resolve().parents[3]
MID = 'c3fd36db7f5b246d56f3c1bbd3f21ca5'


def build(destination, logical):
    destination.mkdir()  # A fresh isolated destination only; never reuse live data.
    config = destination / 'system1/Code/config'
    config.mkdir(parents=True)
    for name in ('config', 'schedule'):
        value = json.loads((ROOT / 'workbench/config/system1' / (name + '.example.json')).read_text())
        if name == 'config':
            value['governance_db'] = '../runtime/governance.sqlite'
            value['random_qa']['enabled'] = False
        else:
            value['enabled'] = False
        (config / (name + '.json')).write_text(json.dumps(value))
    # Retain the tested workbook structure, clear its two fictional source rows
    # in a new copy, then let System1 create the empty owning database.
    script = '''from pathlib import Path
import sys
from openpyxl import load_workbook
from system1.governance_store import GovernanceStore
import source_updater as u
root=Path(sys.argv[1]);out=Path(sys.argv[2])
w=load_workbook(root/'workbench/tests/system1/fixtures/source_registry.xlsx')
for sheet,start in [('Source Register',3),('Human Operation Desktop',3),('Machine confidence',2)]:
    for row in w[sheet].iter_rows(min_row=start):
        for cell in row:cell.value=None
book=out/'system1/Requirement_Source_Registry.xlsx';w.save(book);w.close()
cfg=u.read_config(out/'system1/Code/config/config.json')
GovernanceStore.migrate(cfg,cfg['governance_db'])
'''
    python = ROOT / 'workbench/backend/system1/.venv' / ('Scripts/python.exe' if os.name == 'nt' else 'bin/python')
    env = dict(os.environ, PYTHONUTF8='1', PYTHONPATH=os.pathsep.join(map(str, [ROOT/'workbench', ROOT/'workbench/backend/system1/src'])))
    subprocess.run([str(python), '-c', script, str(ROOT), str(destination)], env=env, check=True)
    migrate(destination)
    app = Application(destination/'workbench', code_root=ROOT, start_workers=False)
    actor = 'Weijie Tang'
    try:
        assert not app.collaboration.source_records()
        sync = FullSnapshot(app.collaboration)
        preview = sync.receive(actor, logical.read_bytes())
        assert not preview['conflicts'] and not preview['validation_errors'], preview
        result = sync.apply(actor, {'id':preview['id'], 'explicit_confirmation':True})
        assert result['status'] == 'applied'
        assert [s['source_id'] for s in app.collaboration.source_records()] == ['PE002']
        material = app.collaboration.read_material(actor, MID)
        assert len(material['blocks']) == 85
        sessions = Requirements(app.collaboration).listing(actor, MID)['sessions']
        assert len(sessions) == 21
        units = [u for s in sessions for u in Requirements(app.collaboration).read(actor, s['id'])['units']]
        interpretations = Interpretations(app.collaboration)
        assert all(interpretations.read(actor, u)['revision'] > 0 for u in units)
        raw, manifest = capture(app.collaboration, actor, str(uuid.uuid4()))
        validate(raw)
    finally:
        app.close()
    seed = logical.with_name('example-seed-20260922.zip')
    seed.write_bytes(raw)
    report = {'status':'prepared', 'seed':seed.name, 'sha256':hashlib.sha256(raw).hexdigest(), 'bytes':len(raw),
              'source_ids':['PE002'], 'materials':[MID], 'blocks':85, 'active_requirements':len(sessions),
              'active_interpretations':len(units), 'databases':list(manifest['databases']),
              'fresh_restore_and_restart':'pending'}
    seed.with_suffix('.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    build(Path(sys.argv[1]), Path(sys.argv[2]))
