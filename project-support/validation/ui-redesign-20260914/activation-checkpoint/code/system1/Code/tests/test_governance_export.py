from unittest.mock import patch
from openpyxl import load_workbook
from test_governance_store import GovernanceStoreTests
import source_updater as u
from system1 import governance_export as exporter


class GovernanceExportTests(GovernanceStoreTests):
    def setup_export(self):
        cfg,store=self.migrate();cfg['governance_db']=store.path
        return cfg,store

    def edit(self,cfg,store,value):
        wb=store.workbook(cfg);hs=u.workbook_headers(wb[cfg['sheet_name']],2)
        wb[cfg['sheet_name']].cell(3,hs['issuer']).value=value
        store.save(wb,cfg,wb._governance_revision);wb.close()

    def test_saved_state_failed_excel_retry_and_no_readback(self):
        cfg,store=self.setup_export()
        first=exporter.sync(cfg);prior=cfg['workbook'].read_bytes()
        self.edit(cfg,store,'Verified issuer')
        self.assertEqual(exporter.status(cfg)['status'],'pending')
        lock=self.workbook.with_name('~$'+self.workbook.name);lock.touch()
        with self.assertRaisesRegex(ValueError,'Close the source Excel'):exporter.sync(cfg)
        status=exporter.status(cfg)
        self.assertEqual((status['status'],status['saved_revision'],status['exported_revision']),('failed',2,1))
        self.assertEqual(cfg['workbook'].read_bytes(),prior)
        lock.unlink();exporter.sync(cfg)
        self.assertEqual(exporter.status(cfg)['status'],'current')
        wb=load_workbook(self.workbook,data_only=True);hs=u.workbook_headers(wb[cfg['sheet_name']],2)
        self.assertEqual(wb[cfg['sheet_name']].cell(3,hs['issuer']).value,'Verified issuer')
        self.assertTrue(wb[cfg['sheet_name']].protection.sheet)
        wb[cfg['sheet_name']].cell(3,hs['issuer']).value='Unsubmitted Excel value';wb.save(self.workbook);wb.close()
        with self.assertRaisesRegex(ValueError,'changed or is not ready'):exporter.cached(cfg)
        exporter.sync(cfg)
        self.assertEqual(store.snapshot()['sources'][0]['record']['issuer'],'Verified issuer')
        self.assertEqual(exporter.cached(cfg)['revision'],2)

    def test_open_during_generation_preserves_last_valid_snapshot(self):
        cfg,store=self.setup_export();exporter.sync(cfg)
        prior=self.workbook.read_bytes();marker=exporter.paths(cfg)[0].read_bytes()
        self.edit(cfg,store,'Pending export')
        save=u.save_workbook_atomic;lock=self.workbook.with_name('~$'+self.workbook.name)
        def opening(*args):
            result=save(*args);lock.touch();return result
        with patch('source_updater.save_workbook_atomic',side_effect=opening):
            with self.assertRaisesRegex(ValueError,'Excel opened during'):exporter.sync(cfg)
        self.assertEqual(self.workbook.read_bytes(),prior)
        self.assertEqual(exporter.paths(cfg)[0].read_bytes(),marker)
        self.assertFalse(list(self.workbook.parent.glob('.source-registry-*.xlsx')))

    def test_new_decision_during_export_remains_pending_for_next_snapshot(self):
        cfg,store=self.setup_export();save=u.save_workbook_atomic
        def concurrent(*args):
            result=save(*args);self.edit(cfg,store,'Newer than current output');return result
        with patch('source_updater.save_workbook_atomic',side_effect=concurrent):exporter.sync(cfg)
        status=exporter.status(cfg)
        self.assertEqual((status['status'],status['saved_revision'],status['exported_revision']),('pending',2,1))
        exporter.sync(cfg)
        self.assertEqual(exporter.status(cfg)['status'],'current')


for name in list(GovernanceStoreTests.__dict__):
    if name.startswith('test_'):setattr(GovernanceExportTests,name,None)
del GovernanceStoreTests
