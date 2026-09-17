from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from pdf_extraction.contracts.source import Snapshot
from pdf_extraction.evidence.material_reader import inspect_original, read_material
from pdf_extraction.orchestration.material_parser import parse_material


def snapshot(tmp_path, filename, raw):
    root = tmp_path / 'originals'
    root.mkdir(exist_ok=True)
    path = root / filename
    path.write_bytes(raw)
    source = Snapshot(source_id='TS001',snapshot_id='TS001-001',relative_path=filename,
        content_hash=sha256(raw).hexdigest(),file_format=path.suffix.lstrip('.'),
        operator_selection_decision='INCLUDE',selection_status='INCLUDE',snapshot_status='STORED',
        source_status='CURRENT',download_status='SUCCESS',registry_sha256='0'*64)
    return root,path,source


def excel_bytes():
    book = Workbook()
    ws = book.active
    ws.title = 'Values'
    ws['A1'] = 'Merged source title'
    ws.merge_cells('A1:C2')
    ws['A3'] = 'stored'
    ws['B3'] = '=1+2'
    ws['Z105'] = 'Far right and below'
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions['B'].hidden = True
    book.create_sheet('Empty')
    book.create_sheet('Hidden').sheet_state='hidden'
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def pdf_bytes():
    buffer = BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(30,790,'First page text')
    c.showPage()
    c.showPage()
    c.drawString(30,790,'Third page text')
    c.showPage()
    c.drawString(30,790,'Fourth page after chunk boundary')
    c.save()
    return buffer.getvalue()


def test_reader_rejects_changed_bytes_and_bounds(tmp_path):
    _,path,source = snapshot(tmp_path,'sample.xlsx',excel_bytes())
    with pytest.raises(ValueError,match='original_version_changed'):
        read_material(path,'f'*64)
    with pytest.raises(ValueError,match='original_cell_window_out_of_range'):
        read_material(path,source.content_hash,row=1000000)
    with pytest.raises(ValueError,match='original_sheet_not_found'):
        read_material(path,source.content_hash,sheet='Unknown')


def test_excel_full_navigation_merge_formula_and_empty_hidden_scope(tmp_path):
    _,path,source = snapshot(tmp_path,'sample.xlsx',excel_bytes())
    manifest = inspect_original(path,source.content_hash)
    assert [s['id'] for s in manifest['scope']] == ['sheet:Values','sheet:Empty','sheet:Hidden']
    first = read_material(path,source.content_hash,row=2,column=2,row_count=2,column_count=2)
    assert first['cells'][0][0]['merge_anchor']=='A1'
    assert first['cells'][0][0]['merge_value']=='Merged source title'
    formula = first['cells'][1][0]
    assert formula['formula']=='=1+2' and formula['cache_status']=='missing'
    assert formula['hidden_row'] and formula['hidden_column']
    last = read_material(path,source.content_hash,row=105,column=26)
    assert last['cells'][0][0]['value']=='Far right and below'
    empty = read_material(path,source.content_hash,sheet='Empty')
    assert empty['cells'][0][0]['value'] is None
    hidden = read_material(path,source.content_hash,sheet='Hidden')
    assert hidden['state']=='hidden'


def test_html_complete_sanitized_selectable_and_stable_anchors(tmp_path):
    raw = b'''<!doctype html><html><head><style>@import 'https://evil.invalid/style'; p{display:none}</style><script>alert(1)</script></head><body onload="evil()"><h1>Title</h1><p>Paragraph <b>bold</b></p><table><tr><td colspan="2">Merged cell</td></tr></table><img src="https://evil.invalid/pixel" alt="Figure"><a href="javascript:evil()">link text</a><iframe src="https://evil.invalid"></iframe><details><summary>Open me</summary><p>Hidden static text</p></details></body></html>'''
    root,path,source=snapshot(tmp_path,'sample.html',raw)
    read=read_material(path,source.content_hash)
    assert 'Hidden static text' in read['html'] and 'Merged cell' in read['html']
    assert '<script' not in read['html'] and '<iframe' not in read['html']
    assert 'evil.invalid' not in read['html'] and 'onload' not in read['html']
    assert 'display:none' not in read['html'] and 'href=' not in read['html']
    assert '<details' in read['html'] and 'open="open"' in read['html']
    assert read['anchors']==read_material(path,source.content_hash)['anchors']
    parsed=parse_material(source,root,tmp_path/'candidate')
    anchors={a['id'] for a in read['anchors']}
    assert all(r['anchor'] in anchors for b in parsed['blocks'] for r in b['source_refs'])
    assert any(b['type']=='heading' and b['text']=='Title' for b in parsed['blocks'])
    table=next(b['table'] for b in parsed['blocks'] if b['type']=='table')
    assert table['rows']==[['Merged cell','']] and table['merges']==[{'row':0,'col':0,'rowspan':1,'colspan':2}]
    assert any(b['type']=='image' for b in parsed['blocks'])
    assert path.read_bytes()==raw


def test_parser_excel_keeps_far_cells_formulas_merges_and_original(tmp_path):
    raw=excel_bytes()
    root,path,source=snapshot(tmp_path,'sample.xlsx',raw)
    result=parse_material(source,root,tmp_path/'candidate')
    assert len(result['scope'])==3
    tables=[b['table'] for b in result['blocks'] if b['type']=='table']
    assert any('=1+2' in row for t in tables for row in t['rows'])
    assert any('Far right and below' in row for t in tables for row in t['rows'])
    assert any(t['merges'] for t in tables)
    assert result['requirement_status']=='not_connected'
    assert (tmp_path/'candidate'/'excel-canonical.json').is_file()
    assert path.read_bytes()==raw
    with pytest.raises(FileExistsError): parse_material(source,root,tmp_path/'candidate')


def test_pdf_reader_and_parser_include_blank_scope_and_chunk_mapping(tmp_path):
    raw=pdf_bytes()
    root,path,source=snapshot(tmp_path,'sample.pdf',raw)
    manifest=inspect_original(path,source.content_hash)
    assert manifest['pages']==4 and len(manifest['scope'])==4
    reader=read_material(path,source.content_hash,page=2)
    assert reader['page']==2 and not reader['selectable_native']
    assert reader['image'].startswith('data:image/png;base64,')
    reader=read_material(path,source.content_hash,page=4)
    assert 'Fourth page' in reader['native_text']
    result=parse_material(source,root,tmp_path/'candidate')
    assert result['covered_scope']==['page:1','page:2','page:3','page:4']
    fourth=[b for b in result['blocks'] if 'Fourth page' in b['text']]
    assert fourth and fourth[0]['source_refs'][0]['page_index']==3
    assert any('Page 2' in w for w in result['warnings'])
    assert path.read_bytes()==raw
    assert not any('requirement-assembly' in p.name for p in (tmp_path/'candidate').iterdir())
    assert json.loads((tmp_path/'candidate'/'pdf-native-0002.json').read_text())['original_page_indices']==[3]


def test_pdf_partial_window_is_honest(tmp_path):
    root,_,source=snapshot(tmp_path,'sample.pdf',pdf_bytes())
    result=parse_material(source,root,tmp_path/'candidate',page_indices={2})
    assert result['status']=='partial' and result['covered_scope']==['page:3']
    assert len(result['scope'])==4
    assert not any('First page' in b['text'] for b in result['blocks'])


def test_pdf_ruled_table_and_image_remain_structural_candidates(tmp_path):
    from PIL import Image
    from reportlab.lib.utils import ImageReader
    buffer=BytesIO()
    c=canvas.Canvas(buffer)
    c.drawString(30,790,'Text before the table')
    for x in (40,180,360): c.line(x,500,x,650)
    for y in (500,550,600,650): c.line(40,y,360,y)
    for x,y,text in [(50,620,'Left heading'),(190,620,'Right heading'),(50,570,'A1'),(190,570,'B1'),(50,520,'A2'),(190,520,'B2')]:
        c.drawString(x,y,text)
    c.drawImage(ImageReader(Image.new('RGB',(40,40),'red')),40,400,width=40,height=40)
    c.save()
    root,path,source=snapshot(tmp_path,'table.pdf',buffer.getvalue())
    parsed=parse_material(source,root,tmp_path/'candidate')
    tables=[b for b in parsed['blocks'] if b['type']=='table']
    assert tables
    assert any('A1' in value for b in tables for row in b['table']['rows'] for value in row)
    assert any(b['type']=='image' and b['image']['source_ref']['page_index']==0 for b in parsed['blocks'])
    assert any('Text before' in b['text'] for b in parsed['blocks'])
    assert all(b['source_refs'][0]['page_index']==0 for b in parsed['blocks'])
