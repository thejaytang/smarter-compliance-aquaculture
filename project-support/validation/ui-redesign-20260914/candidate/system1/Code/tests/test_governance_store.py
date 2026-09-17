import json
import shutil
import sqlite3
from unittest.mock import patch
from openpyxl import load_workbook
from test_operator_journeys import OperatorJourneyTests
import source_updater as u
from system1.governance_store import GovernanceStore, business_rows


class GovernanceStoreTests(OperatorJourneyTests):
    def migrate(self):
        config = u.read_config(self.config_path)
        wb = load_workbook(self.workbook)
        wb.security = None  # The real imported registry has no workbookProtection element.
        wb.properties.keywords = 'system1-browser-workbench'
        u.save_workbook_atomic(wb, self.workbook, u.workbook_mtime(self.workbook)); wb.close()
        store = GovernanceStore.migrate(config, self.runtime/'governance/state.sqlite')
        return config, store

    def test_exact_business_round_trip_and_cached_date_types(self):
        config, store = self.migrate()
        before = store.snapshot()
        original = load_workbook(self.workbook)
        rebuilt = store.workbook(config)
        self.assertEqual(business_rows(original, config), business_rows(rebuilt, config))
        self.assertEqual(store.save(rebuilt, config, 1), 1)
        self.assertEqual(store.snapshot(), before)
        original.close(); rebuilt.close()
        values = store.workbook(config, data_only=True)
        source = before['sources'][0]['record']
        headers = u.workbook_headers(values[config['sheet_name']], 2)
        for key in ('selection_status', 'current_snapshot_date', 'needs_human_action'):
            self.assertEqual(values[config['sheet_name']].cell(3, headers[key]).value, source[key])
        values.close()
        with self.assertRaisesRegex(ValueError, 'already exists'): GovernanceStore.migrate(config, store.path)

    def test_doctor_distinguishes_database_recovery_from_derived_excel(self):
        from system1.doctor import run_doctor
        _, store = self.migrate()
        raw=json.loads(self.config_path.read_text()); raw['governance_db']=str(store.path)
        self.config_path.write_text(json.dumps(raw))
        def checks():return {c['name']:c for c in run_doctor(self.config_path)['checks']}
        self.assertEqual(checks()['governance_database']['status'],'PASS')
        with patch('system1.doctor.excel_appears_open',return_value=True):
            self.assertEqual(checks()['excel_lock']['status'],'WARN')
        self.workbook.unlink()
        self.assertEqual(checks()['workbook']['status'],'WARN')
        self.assertEqual(checks()['governance_database']['status'],'PASS')
        companion=store.path.parent/store.snapshot()['state']['archive']
        original=companion.read_bytes(); companion.write_bytes(b'changed')
        self.assertEqual(checks()['governance_database']['status'],'FAIL')
        companion.write_bytes(original); companion.unlink()
        self.assertEqual(checks()['governance_database']['status'],'FAIL')
        store.path.unlink()
        self.assertEqual(checks()['governance_database']['status'],'FAIL')
        self.assertFalse(store.path.exists())

    def test_business_adapter_matches_full_view_without_sharing_mutable_state(self):
        config,store=self.migrate()
        full=store.workbook(config);first=store.business_workbook(config);other=store.business_workbook(config)
        self.assertEqual(business_rows(full,config),business_rows(first,config))
        headers=u.workbook_headers(first[config['sheet_name']],config['header_row'])
        first[config['sheet_name']].cell(3,headers['issuer']).value='Only this transaction'
        self.assertNotEqual(business_rows(first,config),business_rows(other,config))
        self.assertEqual(business_rows(full,config),business_rows(other,config))
        store.save(first,config,1)
        fresh=store.business_workbook(config)
        self.assertEqual(fresh[config['sheet_name']].cell(3,headers['issuer']).value,'Only this transaction')
        for wb in (full,first,other,fresh):wb.close()

    def test_stale_views_append_only_history_and_atomic_failure(self):
        config, store = self.migrate()
        first = store.workbook(config); stale = store.workbook(config)
        hs = u.workbook_headers(first[config['sheet_name']], 2)
        first[config['sheet_name']].cell(3, hs['issuer']).value = 'Checked issuer'
        self.assertEqual(store.save(first, config, 1, actor='Ana'), 2)
        with self.assertRaisesRegex(ValueError, 'STALE'): store.save(stale, config, 1)
        with self.assertRaisesRegex(ValueError, 'STALE'): store.save(stale, config, 2)
        # Failing after a record update must roll back that update and its history.
        first[config['sheet_name']].cell(3, hs['content_hash']).value = 'changed identity'
        before = store.snapshot()
        with self.assertRaisesRegex(ValueError, 'version hash'): store.save(first, config, 2)
        self.assertEqual(store.snapshot(), before)
        with store.connect() as db:
            with self.assertRaisesRegex(sqlite3.IntegrityError, 'append-only'): db.execute('DELETE FROM history')
            self.assertEqual(db.execute('SELECT actor FROM history ORDER BY sequence DESC LIMIT 1').fetchone()[0], 'Ana')
        first.close(); stale.close()

    def test_derived_excel_edits_and_locks_cannot_change_authority(self):
        config, store = self.migrate()
        baseline = store.snapshot()
        external = load_workbook(self.workbook)
        hs = u.workbook_headers(external[config['sheet_name']], 2)
        external[config['sheet_name']].cell(3, hs['issuer']).value = 'Unsubmitted Excel edit'
        external.save(self.workbook)
        with self.assertRaisesRegex(ValueError, 'readback is forbidden'): store.save(external, config, 1)
        external.close()
        view = store.workbook(config)
        self.assertEqual(store.snapshot(), baseline)
        lock = self.workbook.with_name('~$'+self.workbook.name); lock.touch()
        view[config['sheet_name']].cell(3, hs['issuer']).value = 'Database decision while Excel locked'
        self.assertEqual(store.save(view, config, 1), 2)
        self.assertEqual(store.snapshot()['sources'][0]['record']['issuer'], 'Database decision while Excel locked')
        view.close()

    def test_interrupted_migration_is_unpublished_and_retryable(self):
        config = u.read_config(self.config_path)
        wb = load_workbook(self.workbook)
        u.save_workbook_atomic(wb, self.workbook, u.workbook_mtime(self.workbook)); wb.close()
        path = self.runtime/'atomic/state.sqlite'
        with patch.object(GovernanceStore, '_versions', side_effect=RuntimeError('interrupted import')):
            with self.assertRaisesRegex(RuntimeError, 'interrupted'): GovernanceStore.migrate(config, path)
        self.assertFalse(path.exists())
        self.assertEqual((path.parent/'state-migration-input.xlsx').read_bytes(), self.workbook.read_bytes())
        self.assertFalse(list(path.parent.glob('.state.sqlite*')))
        store = GovernanceStore.migrate(config, path)
        self.assertEqual(store.revision(), 1)

    def test_original_change_during_import_prevents_publication(self):
        config = u.read_config(self.config_path)
        wb = load_workbook(self.workbook)
        u.save_workbook_atomic(wb, self.workbook, u.workbook_mtime(self.workbook)); wb.close()
        path = self.runtime/'atomic/state.sqlite'
        versions = GovernanceStore._versions
        def changing(*args, **kwargs):
            versions(*args, **kwargs)
            (config['source_root']/'unexpected-file.txt').write_text('Concurrent acquisition')
        with patch.object(GovernanceStore, '_versions', side_effect=changing):
            with self.assertRaisesRegex(ValueError, 'changed during migration'): GovernanceStore.migrate(config, path)
        self.assertFalse(path.exists())

    def test_backup_template_failure_does_not_publish_a_broken_backup(self):
        config, store = self.migrate()
        archive = store.path.parent/store.snapshot()['state']['archive']
        archive.write_bytes(b'corrupted template')
        destination = self.runtime/'broken/state.sqlite'
        with self.assertRaisesRegex(ValueError, 'template hash'): store.backup(destination)
        self.assertFalse(destination.exists())

    def test_recovery_retains_history_and_does_not_overwrite_live_state(self):
        config, store = self.migrate()
        destination = self.runtime/'recovered/state.sqlite'
        store.backup(destination)
        archive = store.snapshot()['state']['archive']
        shutil.copy2(store.path.parent/archive, destination.parent/archive)
        recovered = GovernanceStore(destination)
        self.assertEqual(recovered.snapshot(), store.snapshot())
        view = recovered.workbook(config)
        hs = u.workbook_headers(view[config['sheet_name']], 2)
        view[config['sheet_name']].cell(3, hs['issuer']).value = 'Recovery rehearsal'
        recovered.save(view, config, 1);view.close()
        self.assertEqual(store.revision(), 1)
        self.assertEqual(recovered.revision(), 2)
        with self.assertRaises(FileExistsError): store.backup(destination)


for _name in list(OperatorJourneyTests.__dict__):
    if _name.startswith('test_'): setattr(GovernanceStoreTests, _name, None)
del OperatorJourneyTests
