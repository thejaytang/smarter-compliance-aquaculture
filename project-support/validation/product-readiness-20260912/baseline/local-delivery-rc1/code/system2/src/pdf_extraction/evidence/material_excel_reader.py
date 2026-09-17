"""Bounded-memory OOXML reader. Never evaluates formulas or external links."""
from io import BytesIO
import posixpath
import re
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile
from lxml import etree
from ..formats.excel import NS, REL, POLICY


def tag(name):
    return '{' + NS + '}' + name


def coordinate(value):
    match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', value or '')
    if not match: raise ValueError('xlsx_cell_coordinate_invalid')
    letters, number = match.groups()
    column = 0
    for letter in letters: column = column * 26 + ord(letter) - 64
    row = int(number)
    if column > 16384 or row > 1048576: raise ValueError('xlsx_cell_coordinate_out_of_bounds')
    return row, column


def archive_reader(raw, sheet=None, window=None, anchor_requests=frozenset(), inventory_only=False):
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist(); names = {entry.filename for entry in entries}
        if len(entries) > POLICY['max_parts'] or len(names) != len(entries) or sum(e.file_size for e in entries) > POLICY['max_expanded_bytes']:
            raise ValueError('xlsx_container_limits_or_duplicate_parts')
        if any(n.startswith('/') or '..' in n.split('/') or '\\' in n for n in names): raise ValueError('xlsx_unsafe_part_path')
        def tree(name):
            node = etree.fromstring(archive.read(name), parser)
            if node.getroottree().docinfo.doctype: raise ValueError('xlsx_doctype_not_supported')
            return node
        book = tree('xl/workbook.xml')
        if book.tag != tag('workbook'): raise ValueError('xlsx_namespace_not_supported')
        relationships = {}
        for rel in tree('xl/_rels/workbook.xml.rels'):
            rid = rel.get('Id')
            if not rid or rid in relationships: raise ValueError('xlsx_duplicate_or_missing_relationship')
            relationships[rid] = rel
        def resolve(rel):
            target = rel.get('Target', ''); uri = urlsplit(target)
            if rel.get('TargetMode') == 'External' or uri.scheme or uri.netloc or uri.query or uri.fragment or not target:
                raise ValueError('xlsx_relationship_not_local')
            target = unquote(uri.path)
            if '\\' in target or '\x00' in target: raise ValueError('xlsx_relationship_invalid_path')
            part = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('xl', target))
            if part.startswith('../') or part not in names: raise ValueError('xlsx_relationship_part_missing_or_outside')
            return part
        strings = []
        string_rels = [r for r in relationships.values() if r.get('Type') == REL + '/sharedStrings']
        if len(string_rels)>1: raise ValueError('xlsx_shared_string_relationship_ambiguous')
        if string_rels and not inventory_only:
            shared = tree(resolve(string_rels[0]))
            if shared.tag != tag('sst'): raise ValueError('xlsx_shared_string_part_type_invalid')
            strings = [''.join(s.xpath('./s:t/text() | ./s:r/s:t/text()',namespaces={'s':NS})) for s in shared]
        entries = book.find(tag('sheets'))
        if entries is None or not len(entries): raise ValueError('xlsx_no_sheets')
        infos = []; selected = None; used = set()
        for entry in entries:
            rel = relationships[entry.get('{' + REL + '}id')]
            if rel.get('Type') != REL + '/worksheet': raise ValueError('xlsx_nonworksheet_sheet_not_supported')
            part = resolve(rel)
            if part in used: raise ValueError('xlsx_sheet_part_invalid')
            used.add(part)
            info = dict(name=entry.get('name'), state=entry.get('state','visible'), part=part, rows=1, columns=1)
            if inventory_only:
                # Opening needs whole worksheet identities, including blank/hidden
                # sheets. Dimensions are measured when that reader is requested.
                info.update(rows=None,columns=None,dimensions_pending=True)
                infos.append(info)
                continue
            wanted = window is not None and info['name'] == (sheet or entries[0].get('name'))
            cells = {}; merges = []; hidden_rows = set(); hidden_columns = set(); sheet_data = False; shared_formulas = {}
            # Clear complete rows, keeping memory independent of worksheet length.
            with archive.open(part) as stream:
                events = etree.iterparse(stream, events=('start','end'), resolve_entities=False, no_network=True, load_dtd=False)
                for event, element in events:
                    if event == 'start':
                        if element.getparent() is None and element.tag != tag('worksheet'): raise ValueError('xlsx_sheet_type_invalid')
                        if element.tag == tag('sheetData'): sheet_data = True
                        continue
                    if element.tag == tag('dimension'):
                        try:
                            _,_,x2,y2 = range_boundaries(element.get('ref'))
                            info['rows'] = max(info['rows'], y2); info['columns'] = max(info['columns'], x2)
                        except (ValueError,TypeError): pass
                    elif element.tag == tag('col') and element.get('hidden') in {'1','true'}:
                        first_column,last_column = int(element.get('min')),int(element.get('max'))
                        if not 1 <= first_column <= last_column <= 16384: raise ValueError('xlsx_column_range_out_of_bounds')
                        hidden_columns.update(range(first_column,last_column+1))
                    elif element.tag == tag('mergeCell'):
                        region = element.get('ref'); x1,y1,x2,y2 = range_boundaries(region)
                        if not (1 <= x1 <= x2 <= 16384 and 1 <= y1 <= y2 <= 1048576): raise ValueError('xlsx_merge_out_of_bounds')
                        info['rows'] = max(info['rows'],y2); info['columns'] = max(info['columns'],x2)
                        merges.append(dict(range=region,row=y1,column=x1,rowspan=y2-y1+1,colspan=x2-x1+1,anchor=f'{get_column_letter(x1)}{y1}'))
                    elif element.tag == tag('row'):
                        r = int(element.get('r','0'))
                        if element.get('hidden') in {'1','true'}: hidden_rows.add(r)
                        for cell in element:
                            if cell.tag != tag('c'): raise ValueError('xlsx_unknown_cell_structure')
                            rr, cc = coordinate(cell.get('r'))
                            if rr != r: raise ValueError('xlsx_cell_coordinate_out_of_bounds')
                            info['rows'] = max(info['rows'],rr); info['columns'] = max(info['columns'],cc)
                            formula = cell.find(tag('f')) if wanted else None
                            if formula is not None and formula.get('t') == 'shared' and formula.text:
                                shared_formulas[formula.get('si')] = dict(anchor=cell.get('r'), reference=formula.get('ref'))
                            if not wanted: continue
                            row,col,nrows,ncols = window
                            # Retain merge anchors outside visible window as necessary.
                            if not (row <= rr < row+nrows and col <= cc < col+ncols) and cell.get('r') not in anchor_requests: continue
                            value_node = cell.find(tag('v')); value = value_node.text if value_node is not None else None; formula = cell.find(tag('f')); text = None
                            if cell.get('t') == 's':
                                try:
                                    index = int(value)
                                    if index < 0: raise ValueError()
                                    text = strings[index]
                                except (ValueError,TypeError,IndexError): raise ValueError('xlsx_shared_string_invalid')
                            elif cell.get('t') == 'inlineStr': text = ''.join(cell.xpath('./s:is/s:t/text() | ./s:is/s:r/s:t/text()',namespaces={'s':NS}))
                            elif cell.get('t') == 'str': text = value
                            address = cell.get('r')
                            if address in cells: raise ValueError('xlsx_cell_coordinate_invalid_or_duplicate')
                            cells[address] = dict(address=address,value=text if text is not None else value,formula='='+formula.text if formula is not None and formula.text else None,formula_present=formula is not None,formula_attributes=dict(formula.attrib) if formula is not None else None,formula_kind=formula.get('t','normal') if formula is not None else None,cached=value if formula is not None else None,cache_status='not_formula' if formula is None else 'present' if value is not None or (value_node is not None and cell.get('t') == 'str') else 'missing')
                        element.clear()
                        while element.getprevious() is not None: del element.getparent()[0]
                if events.root.getroottree().docinfo.doctype: raise ValueError('xlsx_doctype_not_supported')
            if not sheet_data: raise ValueError('xlsx_missing_sheet_data')
            infos.append(info)
            if wanted:
                for address,cell in cells.items():
                    if not cell['formula_present']: continue
                    attributes = cell['formula_attributes']
                    cell['formula_reference'] = attributes.get('ref')
                    cell['formula_anchor'] = address if cell['formula'] else None
                    if cell['formula_kind'] == 'shared':
                        anchor = shared_formulas.get(attributes.get('si'), {})
                        cell['formula_anchor'] = anchor.get('anchor')
                        cell['formula_reference'] = anchor.get('reference') or attributes.get('ref')
                        if not cell['formula']:
                            cell['formula_note'] = ('Shared formula follower; expression is stored at '+anchor['anchor']+'. No expression is expanded or calculated.') if anchor else 'Shared formula follower; the anchor was not found in this saved worksheet. No expression is invented or calculated.'
                    elif cell['formula_kind'] == 'array':
                        cell['formula_note'] = 'Stored array formula and source range. Results are not calculated.'
                    elif not cell['formula']:
                        cell['formula_note'] = 'Formula record has no expression text in this cell. Original formula attributes are retained; no expression is invented.'
                selected = (info,cells,merges,hidden_rows,hidden_columns)
        if window is None: return {'sheets':infos}
        if selected is None: raise ValueError('original_sheet_not_found')
        info,cells,merges,hidden_rows,hidden_columns = selected
        row,col,nrows,ncols = window
        for value, maximum in ((row,info['rows']),(col,info['columns']),(nrows,200),(ncols,100)):
            if type(value) is not int or not 1<=value<=maximum: raise ValueError('original_cell_window_out_of_range')
        needed_anchors = {m['anchor'] for m in merges if m['row'] < row+nrows and m['row']+m['rowspan'] > row and m['column'] < col+ncols and m['column']+m['colspan'] > col and m['anchor'] not in cells}
        if needed_anchors and not anchor_requests:
            return archive_reader(raw, sheet=sheet, window=window, anchor_requests=needed_anchors)
        grid=[]
        for rr in range(row,min(info['rows']+1,row+nrows)):
            line=[]
            for cc in range(col,min(info['columns']+1,col+ncols)):
                address=f'{get_column_letter(cc)}{rr}'
                cell=dict(cells.get(address,dict(address=address,value=None,formula=None,cached=None,cache_status='not_formula')))
                merge=next((m for m in merges if m['row']<=rr<m['row']+m['rowspan'] and m['column']<=cc<m['column']+m['colspan']),None)
                cell.update(hidden_row=rr in hidden_rows,hidden_column=cc in hidden_columns,merge_anchor=merge['anchor'] if merge else None,merge_range=merge['range'] if merge else None)
                if merge:
                    cell['merge_value']=cells.get(merge['anchor'],{}).get('value')
                    if merge['anchor'] not in cells: cell['merge_value_notice']='Navigate to the merge anchor to read its value.'
                line.append(cell)
            grid.append(line)
        return dict(sheets=infos,sheet=info['name'],state=info['state'],part=info['part'],row=row,column=col,row_count=len(grid),column_count=len(grid[0]) if grid else 0,rows=info['rows'],columns=info['columns'],cells=grid,merged=merges,label='Bound original workbook · navigate every worksheet and cell range')
