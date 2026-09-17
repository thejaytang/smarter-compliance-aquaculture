"""Original-side inspection scopes, built without reading extraction units.

These are sampling positions, not independent accuracy labels or a verifier.
PDF pages include pages with no output. HTML text nodes and spreadsheet cells
retain their positions even when the extractor has no corresponding record.
"""
from hashlib import sha256
from pathlib import Path
from ..contracts.hashing import digest

SCHEMA = 'original-sampling/1'


def catalog(path, expected_hash):
    path = Path(path)
    raw = path.read_bytes()
    if sha256(raw).hexdigest() != expected_hash:
        raise ValueError('original_sampling_source_hash_changed')
    suffix = path.suffix.lower()
    scopes, unverified = [], []
    if suffix == '.pdf':
        from pypdf import PdfReader
        from io import BytesIO
        reader = PdfReader(BytesIO(raw))
        for page in range(len(reader.pages)):
            scopes.append({'title': f'Original PDF page {page+1}', 'page_index': page,
                           'references': [{'page_index': page}], 'original_text': None})
    elif suffix in {'.html', '.htm'}:
        from lxml import html
        tree = html.document_fromstring(raw)
        def locator(node):
            parts = []
            while node is not None and isinstance(node.tag, str):
                previous = list(node.itersiblings(preceding=True))
                parts.append(f'{node.tag}:nth-of-type({1+sum(n.tag==node.tag for n in previous)})')
                node = node.getparent()
            return ' > '.join(reversed(parts))
        nodes = []
        for text in tree.xpath('//body//text()'):
            parent = text.getparent()
            ancestors = [parent, *parent.iterancestors()]
            if not str(text).strip() or any(n.tag in {'script','style'} for n in ancestors):
                continue
            # A tail belongs after its element; retain that distinction explicitly.
            nodes.append({'text':str(text), 'locator':locator(parent),
                          'text_position':'tail' if text.is_tail else 'text'})
        for start in range(0, len(nodes), 20):
            group = nodes[start:start+20]
            scopes.append({'title':f'Original HTML text positions {start+1}–{start+len(group)}',
                'references':[{'locator':n['locator'],'text_position':n['text_position']} for n in group],
                'original_text':'\n'.join(n['text'] for n in group), 'original_nodes':group})
        for i, node in enumerate(tree.xpath('//body//img | //body//svg'), 1):
            scopes.append({'title':f'Original HTML image {i}', 'references':[{'locator':locator(node)}],
                           'original_text':None, 'visual_only':True})
        if tree.xpath('//script | //iframe | //object | //embed | //link[@href]'):
            unverified.append('Script-generated, embedded or externally loaded content is outside the retained local markup catalog.')
    elif suffix == '.xlsx':
        from io import BytesIO
        from zipfile import ZipFile
        from lxml import etree
        import posixpath
        from openpyxl import load_workbook
        with ZipFile(BytesIO(raw)) as archive:
            rels=etree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
            targets={r.get('Id'):posixpath.normpath('xl/'+r.get('Target')) if not r.get('Target').startswith('/') else r.get('Target').lstrip('/') for r in rels}
            book=etree.fromstring(archive.read('xl/workbook.xml'))
            parts={s.get('name'):targets[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')] for s in book.findall('.//{*}sheet')}
            if any(n.startswith('xl/drawings/') for n in archive.namelist()):
                unverified.append('Spreadsheet drawing/image layout is not represented by the cell catalog; inspect it in the original view.')
        workbook=load_workbook(BytesIO(raw),read_only=True,data_only=False)
        try:
            for sheet in workbook:
                cells=[{'locator':parts[sheet.title]+'#'+c.coordinate,'text':str(c.value)}
                       for row in sheet.iter_rows() for c in row if c.value is not None]
                for start in range(0,len(cells),40):
                    group=cells[start:start+40]
                    scopes.append({'title':f'{sheet.title}: original cells {start+1}–{start+len(group)}',
                        'references':[{'locator':c['locator']} for c in group],
                        'original_text':'\n'.join(c['locator'].split('#')[-1]+': '+c['text'] for c in group),
                        'original_cells':group,'worksheet':sheet.title})
                if not cells:
                    scopes.append({'title':f'{sheet.title}: empty/visual worksheet',
                        'references':[{'locator':parts[sheet.title]+'#A1'}], 'original_text':'', 'visual_only':True})
        finally:
            workbook.close()
    else:
        raise ValueError('original_sampling_format_not_supported')
    for scope in scopes:
        scope.update(id=digest([SCHEMA,expected_hash,scope])[:32],schema=SCHEMA,
                     source_sha256=expected_hash,unverified=list(unverified))
    return {'schema':SCHEMA,'source_sha256':expected_hash,'scopes':scopes,'unverified':unverified}


def matches(scope, references):
    """Locate overlap, including large existing units covering a sampled position."""
    if scope.get('page_index') is not None:
        return any(r.get('page_index') == scope['page_index'] for r in references)
    for original in scope['references']:
        a=original.get('locator','').split('::')[0]
        for ref in references:
            b=ref.get('locator','').split('::')[0]
            if a and b and (a==b or a.startswith(b+' > ') or b.startswith(a+' > ')):
                return True
    return False
