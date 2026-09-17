"""Explicit, local structural extraction into a separate human-review candidate.

This adapter deliberately does not call the legacy domain-producing pipeline.
It reuses pure HTML/OOXML parsers, positioned NativeExtractor and ruled-table
geometry; candidate artifacts are immutable and never human confirmations.
"""
from __future__ import annotations

from contextlib import closing
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re

from ..evidence.material_reader import _manifest, html_anchor
from ..intake.registry import read_snapshot

VERSION = 'material-structural-parser/1'
POLICY = {'adapter': VERSION, 'semantic_processing': False, 'external_models': False, 'ocr': False, 'pdf_window_pages': 3, 'human_adoption_required': True}
CONFIG_HASH = sha256(json.dumps(POLICY, sort_keys=True).encode()).hexdigest()


def _write(path, payload):
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)


def _block(source, locator, kind, text='', **extra):
    return {'id': 'block-' + sha256((source.snapshot_id + ':' + locator).encode()).hexdigest()[:24],
            'type': kind, 'text': text, 'level': 1, 'parent_id': None,
            'numbering': '', 'dependencies': [], 'source_refs': [], **extra}


def _html_blocks(raw, source, output):
    from bs4 import BeautifulSoup, Tag, NavigableString, Comment
    from ..formats.html import parse_html, dom_path
    from ..formats.html_tables import build_table
    from ..formats.html_profiles import HtmlProfileError
    artifacts, warnings = [], []
    try:
        doc = parse_html(raw, source, {'template': 'auto', 'encoding': 'utf-8-sig'})
        _write(output / 'html-canonical.json', doc.model_dump(mode='json'))
        artifacts.append('html-canonical.json')
        warnings.extend(doc.issues)
    except HtmlProfileError:
        warnings.append('The saved HTML has no supported template profile; the complete static DOM is retained as a generic structural candidate for manual review.')
    soup = BeautifulSoup(raw, 'lxml')
    paths = {id(el): dom_path(el) for el in soup.find_all(True)}
    ids = {id(el): 'node-' + sha256(paths[id(el)].encode()).hexdigest()[:20] for el in soup.find_all(True)}
    node_by_id = {ids[id(el)]: el for el in soup.find_all(True)}
    blocks, stack, counts = [], [], {}
    def ref(el):
        return {'scope_id': 'html:document', 'locator': paths[id(el)], 'anchor': html_anchor(paths[id(el)])}
    def append(el, kind, text='', **extra):
        locator = paths[id(el)]
        counts[locator] = counts.get(locator, 0) + 1
        b = _block(source, locator + ':' + str(counts[locator]), kind, text, source_refs=[ref(el)], **extra)
        if kind == 'heading':
            level = int(el.name[1])
            while stack and stack[-1][0] >= level: stack.pop()
            b['level'] = level
        if stack:
            b['parent_id'] = stack[-1][1]
            b['dependencies'] = [stack[-1][1]]
        if kind == 'heading': stack.append((b['level'], b['id']))
        match = re.match(r'^(\d+(?:\.\d+)*[.)]?|§\s*\d+[a-z]?)\s', text)
        if match: b['numbering'] = match[1]
        blocks.append(b)
        return b
    def walk(el):
        if el.name in {'script', 'style', 'head', 'noscript'}: return
        if el.name == 'img':
            append(el, 'image', el.get('alt', ''), image={'source_ref': ref(el), 'attachment': None, 'attribution': 'Image element in the saved original; external image assets are not fetched.'})
            return
        if el.name == 'table':
            table = build_table(el, ids)
            if table.column_count is None:
                warnings.append('Ambiguous HTML table spans retained as row text; repair table geometry manually at ' + paths[id(el)])
                rows = [[c.get_text('', strip=False).strip() for c in r.find_all(['td', 'th'], recursive=False)] for r in el.find_all('tr') if r.find_parent('table') is el]
                width = max([1] + [len(r) for r in rows])
                rows = [r + [''] * (width - len(r)) for r in rows]
                merges = []
            else:
                rows = [[''] * max(1, table.column_count) for _ in range(max(1, table.row_count))]
                merges = []
                for cell in table.cells:
                    rows[cell.row][cell.column] = node_by_id[cell.node_id].get_text('', strip=False).strip()
                    if cell.rowspan > 1 or cell.colspan > 1:
                        merges.append({'row': cell.row, 'col': cell.column, 'rowspan': cell.rowspan, 'colspan': cell.colspan})
            caption = el.find('caption', recursive=False)
            append(el, 'table', '', table={'rows': rows or [['']], 'merges': merges, 'notes': [caption.get_text(' ', strip=True)] if caption else []})
            for img in el.find_all('img'): walk(img)
            return
        # Preserve all static body text, including outside known profile roots.
        semantic = {'p', 'li', 'dt', 'dd', 'pre', 'blockquote', 'figcaption', 'caption', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'}
        nested = el.find(['table', 'img'] + list(semantic))
        if el.name in semantic and nested is None:
            text = el.get_text('', strip=False).strip()
            if text: append(el, 'heading' if re.fullmatch('h[1-6]', el.name) else 'text', text)
            return
        buffer = []
        def flush():
            text = ''.join(buffer).strip()
            if text: append(el, 'heading' if re.fullmatch('h[1-6]', el.name) else 'text', text)
            buffer.clear()
        for child in el.children:
            if isinstance(child, Comment): continue
            if isinstance(child, NavigableString): buffer.append(str(child))
            elif isinstance(child, Tag):
                if child.name == 'br': buffer.append('\n')
                elif child.name in {'span', 'a', 'b', 'strong', 'i', 'em', 'u', 's', 'small', 'sup', 'sub', 'code'} and not child.find(['table', 'img'] + list(semantic)):
                    buffer.append(child.get_text('', strip=False))
                else:
                    flush(); walk(child)
        flush()
    walk(soup.body or soup)
    _write(output / 'html-static-candidate.json', {'source_sha256': source.content_hash, 'blocks': blocks, 'warnings': warnings})
    artifacts.append('html-static-candidate.json')
    return blocks, warnings, artifacts


def _excel_blocks(raw, source, output):
    from ..formats.excel import parse_excel
    from ..evidence.material_reader import _sheet_info
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
    doc = parse_excel(raw, source)
    _write(output / 'excel-canonical.json', doc.model_dump(mode='json'))
    blocks, warnings = [], list(doc.issues)
    for sheet in doc.sheets:
        scope_id = 'sheet:' + sheet.name
        heading = _block(source, sheet.part, 'heading', sheet.name, source_refs=[{'scope_id': scope_id, 'sheet': sheet.name, 'locator': sheet.part + '#A1', 'cell_range': 'A1'}])
        blocks.append(heading)
        # Sparse sheets are tiled; avoid allocating their million-row blank rectangle.
        tiles = {}
        for cell in sheet.cells:
            r, c = coordinate_to_tuple(cell.coordinate)
            tiles.setdefault(((r-1)//100, (c-1)//30), []).append(cell)
        for tile, cells in sorted(tiles.items()):
            points = [coordinate_to_tuple(c.coordinate) for c in cells]
            r1, c1 = min(p[0] for p in points), min(p[1] for p in points)
            r2, c2 = max(p[0] for p in points), max(p[1] for p in points)
            selected_merges = []
            for region in sheet.merged_ranges:
                x1, y1, x2, y2 = range_boundaries(region)
                if y1 <= r2 and y2 >= r1 and x1 <= c2 and x2 >= c1:
                    if x2-x1 < 300 and y2-y1 < 1000:
                        r1, c1, r2, c2 = min(r1,y1),min(c1,x1),max(r2,y2),max(c2,x2)
                        selected_merges.append((x1,y1,x2,y2))
                    else: warnings.append('Large merged range requires original review: ' + sheet.name + '!' + region)
            if (r2-r1+1)*(c2-c1+1) > 50000:
                raise ValueError('excel_candidate_table_limit')
            rows = [[''] * (c2-c1+1) for _ in range(r2-r1+1)]
            notes = []
            for cell in cells:
                r,c = coordinate_to_tuple(cell.coordinate)
                rows[r-r1][c-c1] = '=' + cell.formula if cell.formula is not None else (cell.text if cell.text is not None else cell.value or '')
                if cell.formula is not None: notes.append(cell.coordinate + ' saved value: ' + (cell.value if cell.value is not None else '[cache unavailable]') + '. Formula has not been calculated.')
            merges = [{'row': y1-r1, 'col': x1-c1, 'rowspan': y2-y1+1, 'colspan': x2-x1+1} for x1,y1,x2,y2 in selected_merges]
            cell_range = f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}'
            block = _block(source, sheet.part+'#'+cell_range, 'table', source_refs=[{'scope_id': scope_id, 'sheet': sheet.name, 'locator': sheet.part+'#'+f'{get_column_letter(c1)}{r1}', 'cell_range': cell_range}], parent_id=heading['id'], dependencies=[heading['id']], table={'rows':rows,'merges':merges,'notes':notes})
            blocks.append(block)
    warnings.append('All worksheets remain in required scope, including blank and hidden sheets. Charts, drawings, comments and native formatting require original inspection or manual image/notes supplementation; formulas are source text, not calculated results.')
    return blocks, warnings, ['excel-canonical.json']



def _pdf_tables(page, original_page, source, reference, output):
    """Reuse the existing local ruled-grid detector and merged-cell assembler.

    Words are assigned only by their retained native geometry. Unknown/unruled
    structures stay as text, ready for manual table insertion and correction.
    """
    import numpy as np
    from ..config import TableSettings
    from ..layout.detector import detect_table_boxes
    from ..parsers.table import _grid_positions, _merged_cell_specs, _compact_grid_specs
    bitmap = original_page.render(scale=1.5)
    image_path = output / f"pdf-page-{reference['page']:04d}.png"
    try:
        image = bitmap.to_pil().convert('RGB')
        image.save(image_path)
        array = np.asarray(image)
    finally:
        bitmap.close()
    candidates = []
    for number, box in enumerate(detect_table_boxes(image_path, TableSettings(backend='native', recognition_pipeline_enabled=False, img2table_enabled=False, gmft_enabled=False))):
        crop = array[box.y0:box.y1, box.x0:box.x1]
        xs, ys = _grid_positions(crop)
        if len(xs) < 2 or len(ys) < 2 or (len(xs)-1)*(len(ys)-1)>10000:
            continue
        specs, xs, ys = _compact_grid_specs(_merged_cell_specs(crop, xs, ys), xs, ys)
        rows = [['']*(len(xs)-1) for _ in range(len(ys)-1)]
        merges = []
        claimed = set()
        for row, col, rowspan, colspan in specs:
            x0,y0,x1,y1 = (box.x0+xs[col])/1.5,(box.y0+ys[row])/1.5,(box.x0+xs[col+colspan])/1.5,(box.y0+ys[row+rowspan])/1.5
            words = [w for w in page.words if x0 <= (w.bbox_points[0]+w.bbox_points[2])/2 < x1 and y0 <= (w.bbox_points[1]+w.bbox_points[3])/2 < y1]
            # Native parser order is retained within each geometric cell.
            rows[row][col] = ' '.join(w.text for w in words).strip()
            claimed.update(w.id for w in words)
            if rowspan>1 or colspan>1:
                merges.append({'row':row,'col':col,'rowspan':rowspan,'colspan':colspan})
        if not claimed:
            continue
        bbox=[(box.x0+xs[0])/1.5,(box.y0+ys[0])/1.5,(box.x0+xs[-1])/1.5,(box.y0+ys[-1])/1.5]
        block=_block(source, f"page:{reference['page_index']}:table:{number}", 'table', source_refs=[{**reference,'bbox':bbox}], table={'rows':rows,'merges':merges,'notes':[]})
        candidates.append((bbox,block))
    return candidates, image_path.name


def _pdf_blocks(raw, source, output, page_indices):
    from ..ingest.native_extractor import NativeExtractor
    from ..config import AppConfig
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(raw) as inventory:
        total = len(inventory)
    selected = sorted(set(range(total)) if page_indices is None else set(page_indices))
    if not selected or any(type(p) is not int or not 0 <= p < total for p in selected): raise ValueError('pdf_page_selection_invalid')
    config = AppConfig.from_yaml(Path(__file__).resolve().parents[3] / 'config' / 'pdf-intake-positioned.yaml')
    blocks, warnings, artifacts = [], [], []
    with pdfium.PdfDocument(raw) as original:
        for offset in range(0, len(selected), 3):
            mapping = selected[offset:offset+3]
            chunk = output / f'pdf-window-{offset//3+1:04d}.pdf'
            # PDFium already supports the readable source encryption modes.
            # Avoid an optional pypdf AES dependency for immutable window copies.
            with pdfium.PdfDocument.new() as window:
                window.import_pages(original, pages=mapping)
                with chunk.open('xb') as stream: window.save(stream)
            try:
                native = NativeExtractor(config.native_extraction.backend).extract(chunk)
            except Exception as exc:
                native = NativeExtractor('pdfium').extract(chunk)
                warnings.append('Positioned PDF parser fallback to PDFium for pages ' + ', '.join(str(p+1) for p in mapping) + ': ' + type(exc).__name__)
            evidence = {'source_sha256': source.content_hash, 'config_hash': sha256(config.model_dump_json().encode()).hexdigest(), 'adapter_config_hash': CONFIG_HASH, 'original_page_indices': mapping, 'backend': native.backend, 'version': native.version, 'pages': [asdict(p) for p in native.pages]}
            evidence_name = f'pdf-native-{offset//3+1:04d}.json'
            _write(output / evidence_name, evidence)
            artifacts.extend([chunk.name, evidence_name])
            for local, page in enumerate(native.pages):
                index = mapping[local]
                scope_id = f'page:{index+1}'
                reference = {'scope_id':scope_id, 'page_index':index, 'page':index+1}
                lines = page.text_lines or page.words
                with closing(original[index]) as original_page:
                    tables, page_image = _pdf_tables(page, original_page, source, reference, output)
                    artifacts.append(page_image)
                    page_blocks = [b for _, b in tables]
                    for n,line in enumerate(lines):
                        text = line.text.strip()
                        if not text: continue
                        x0,y0,x1,y1 = line.bbox_points
                        # Suppress only lines fully contained by a table; partial
                        # overlaps remain visible for human association review.
                        if any(b[0]<=x0 and b[1]<=y0 and b[2]>=x1 and b[3]>=y1 for b,_ in tables):
                            continue
                        block = _block(source, f'page:{index}:line:{n}', 'text', text, source_refs=[{**reference, 'bbox':list(line.bbox_points), 'native_id':line.id}])
                        page_blocks.append(block)
                    page_blocks.sort(key=lambda b:(b['source_refs'][0]['bbox'][1],b['source_refs'][0]['bbox'][0]))
                    blocks.extend(page_blocks)
                    # Keep embedded source images beside native structural text.
                    image_count = 0
                    for obj in original_page.get_objects():
                        if obj.type != pdfium.raw.FPDF_PAGEOBJ_IMAGE: continue
                        left,bottom,right,top = obj.get_bounds()
                        bbox = [left,page.height_points-top,right,page.height_points-bottom]
                        ref = {**reference,'bbox':bbox}
                        blocks.append(_block(source, f'page:{index}:image:{image_count}', 'image', image={'source_ref':ref,'attachment':None,'attribution':'Embedded image in the bound original PDF.'}, source_refs=[ref]))
                        image_count += 1
                    if not lines:
                        warnings.append(f'Page {index+1} has no positioned native text. Inspect the complete original page and manually transcribe or supplement omitted text. No OCR/model was run.')
                        if not image_count:
                            blocks.append(_block(source,f'page:{index}:untranscribed','image', image={'source_ref':reference,'attachment':None,'attribution':'Complete original page; inspect for scanned or blank content.'},source_refs=[reference]))
    warnings.append('PDF candidate uses the existing positioned native parser in mapped three-page windows. Ruled-table grids use existing local geometry assembly. Reading order, headings and unruled tables may need manual correction; source page images and complete native evidence are retained. No Requirement classifier, semantic processor, OCR model or external service runs.')
    return blocks,warnings,artifacts,[f'page:{p+1}' for p in selected]


def parse_material(source, source_root, output, *, page_indices=None):
    """Parse one captured snapshot, returning a non-adopted review candidate.

    ``output`` must be a fresh job-owned directory outside the source root. The
    captured bytes are pinned before any parser works; subsequent source changes
    are the owning store's staleness concern, never relabelled by this adapter.
    """
    from ..contracts.source import Snapshot
    from ..formats.router import route_format
    source = source if isinstance(source, Snapshot) else Snapshot.model_validate(source)
    source_root, output = Path(source_root), Path(output)
    if output.resolve().is_relative_to(source_root.resolve()): raise ValueError('output_must_be_outside_source_root')
    raw = read_snapshot(source, source_root)
    kind = route_format(source.relative_path, source.file_format, raw)
    output.mkdir(parents=True, exist_ok=False)
    suffix = Path(source.relative_path).suffix.lower()
    with (output / ('source' + suffix)).open('xb') as stream: stream.write(raw)
    manifest = _manifest(raw, suffix)
    if kind == 'html': blocks,warnings,artifacts = _html_blocks(raw,source,output)
    elif kind == 'excel': blocks,warnings,artifacts = _excel_blocks(raw,source,output)
    else: blocks,warnings,artifacts,covered = _pdf_blocks(raw,source,output,page_indices)
    if kind != 'pdf': covered = [s['id'] for s in manifest['scope']]
    candidate = {'schema_version':VERSION,'parser_version':VERSION,'source_sha256':source.content_hash,
                 'snapshot_id':source.snapshot_id,'config_hash':CONFIG_HASH,'blocks':blocks,'scope':manifest['scope'], 'covered_scope':covered,
                 'warnings':list(dict.fromkeys(manifest['warnings']+warnings)), 'unresolved':[],
                 'canonical_artifacts':artifacts,'status':'candidate_available' if len(covered)==len(manifest['scope']) else 'partial',
                 'review_policy':'human_material_confirmation_required', 'requirement_status':'not_connected'}
    _write(output / 'candidate.json',candidate)
    return candidate
