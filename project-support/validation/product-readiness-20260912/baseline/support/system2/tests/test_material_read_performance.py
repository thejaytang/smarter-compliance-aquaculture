"""Isolated regression checks for projections and source-bound reader caching."""
from hashlib import sha256
import uuid
import json
import pytest
from pdf_extraction.review.materials import MaterialStore
from pdf_extraction.evidence.material_reader import read_material, inspect_original


def request(material, **kwargs):
    return dict(material_id=material['id'],expected_revision=material['revision'],request_id=str(uuid.uuid4()),actor='Isolated Engineer',**kwargs)


def opened(store, number=1):
    source=dict(source_id=f'S{number}',snapshot_id=f'S{number}-001',content_hash=str(number).zfill(64),file_format='pdf',relative_path='sample.pdf')
    return store.open(source,f'Material {number}',[dict(id='page:1',kind='page',label='Page 1',location={'page':1})])


def test_paged_projection_preserves_full_historical_and_candidate_details(tmp_path):
    store=MaterialStore(tmp_path); material=opened(store)
    body=[dict(id=f'b{i}',type='text',text='Human text '+str(i),source_refs=[{'scope_id':'page:1'}]) for i in range(100)]
    saved=store.save(request(material,blocks=body))['material']
    material=store.start_candidate(request(saved))['material']; candidate_id=material['candidates'][0]['id']
    store.finish_candidate(material['id'],candidate_id,body,complete=True)
    original_history=store.history(material['id']); compact=store.read_compact(material['id'])
    assert compact['blocks']==body
    assert 'blocks' not in compact['candidates'][0] and 'base_blocks' not in compact['candidates'][0]
    assert compact['candidates'][0]['block_count']==100
    assert store.candidate_detail(material['id'],candidate_id)['blocks']==body
    assert store.read(material['id'])['candidates'][0]['blocks']==body
    listing=store.list_page()
    assert listing['total']==1 and listing['materials'][0]['block_count']==100
    assert 'blocks' not in listing['materials'][0] and 'scope' not in listing['materials'][0]
    page=store.history_page(material['id'],limit=2)
    assert len(page['revisions'])==2 and page['has_more']
    assert all('blocks' not in revision for revision in page['revisions'])
    assert store.history(material['id'])==original_history
    other=opened(store,2)
    with pytest.raises(ValueError,match='candidate_not_found'): store.candidate_detail(other['id'],candidate_id)


def test_material_search_and_pagination_use_all_materials(tmp_path):
    store=MaterialStore(tmp_path)
    for n in range(1,56): opened(store,n)
    first=store.list_page(); second=store.list_page(offset=50)
    assert first['total']==55 and len(first['materials'])==50 and first['has_more']
    assert len(second['materials'])==5 and not second['has_more']
    assert store.list_page(query='Material 55')['materials'][0]['source']['source_id']=='S55'
    assert store.counts()==dict(total=55,current_reviewed=0,in_progress=55,historical=0,needs_followup=0)
    assert store.counts(['S55'])['needs_followup']==1
    for options in [dict(offset=-1),dict(limit=51),dict(query=['bad'])]:
        with pytest.raises(ValueError): store.list_page(**options)


def test_existing_history_gets_additive_projection_without_rewrite(tmp_path):
    store=MaterialStore(tmp_path); material=opened(store)
    with store.connect() as db:
        before=db.execute('SELECT data FROM material_revisions').fetchone()[0]
        for trigger in ['material_read_insert','material_read_update','material_revision_read_insert']:
            db.execute('DROP TRIGGER '+trigger)
        db.execute('DROP TABLE material_read_index');db.execute('DROP TABLE material_revision_index');db.execute('DELETE FROM material_read_schema')
    reopened=MaterialStore(tmp_path)
    assert reopened.list_page()['total']==1 and reopened.history_page(material['id'])['total']==1
    with reopened.connect() as db: assert db.execute('SELECT data FROM material_revisions').fetchone()[0]==before


def test_reader_cache_rechecks_source_and_recovers_corrupt_cache(tmp_path):
    path=tmp_path/'source.html';path.write_bytes(b'<html><body><p>Bound original</p></body></html>')
    fingerprint=sha256(path.read_bytes()).hexdigest(); original=read_material(path,fingerprint)
    assert read_material(path,fingerprint)==original
    cache=next((tmp_path/'.reader-cache').glob('*.json'))
    envelope=json.loads(cache.read_text());envelope['result']['html']='Silent damaged cached text';cache.write_text(json.dumps(envelope))
    assert read_material(path,fingerprint)==original
    cache.write_text('[]')
    assert read_material(path,fingerprint)==original
    cache.write_text('{broken')
    assert read_material(path,fingerprint)==original
    path.write_bytes(b'<html>Changed source</html>')
    with pytest.raises(ValueError,match='original_version_changed'): read_material(path,fingerprint)
    assert 'Changed source' in read_material(path,sha256(path.read_bytes()).hexdigest())['html']


def test_empty_excel_scope_remains_available(tmp_path):
    from openpyxl import Workbook
    book=Workbook(); book.active.title='Blank';book.create_sheet('Hidden').sheet_state='hidden'
    path=tmp_path/'blank.xlsx';book.save(path); fingerprint=sha256(path.read_bytes()).hexdigest()
    assert [s['id'] for s in inspect_original(path,fingerprint)['scope']]==['sheet:Blank','sheet:Hidden']
    result=read_material(path,fingerprint,sheet='Hidden')
    assert result['state']=='hidden' and result['cells'][0][0]['value'] is None


def test_large_review_provenance_is_excluded_from_summary_transport(tmp_path):
    store=MaterialStore(tmp_path); material=opened(store)
    # Synthetic saved metadata, isolated from business data, exercises projection size.
    material['review_impact']=dict(affected_blocks=['b'+str(i) for i in range(10000)],affected_scope=['page:1'],retained_scope=[],reason='linked_content_changed')
    material['review_checks']={'page:1':{'actor':'Isolated Engineer','at':'retained','content_revision':0}}
    with store.transaction() as db:
        material['revision']+=1;store._write(db,material,'saved','Isolated Engineer')
    summary=store.list_page()['materials'][0]
    assert 'review_checks' not in summary and 'affected_blocks' not in summary['review_impact']
    assert len(json.dumps(summary))<3000
    assert len(store.read(material['id'])['review_impact']['affected_blocks'])==10000


def test_conflict_detail_is_bound_to_material_and_summary_has_no_rejected_body(tmp_path):
    store=MaterialStore(tmp_path); first=opened(store); second=opened(store,2)
    store.save(request(first,blocks=[]))
    conflict=store.save(request(first,blocks=[dict(id='rejected',type='text',text='Unsaved private draft',source_refs=[])]))
    summary=store.read_compact(first['id'])['conflicts'][0]
    assert 'request' not in summary
    assert store.conflict_detail(first['id'],conflict['conflict_id'])['request']['blocks'][0]['text']=='Unsaved private draft'
    with pytest.raises(ValueError,match='conflict_not_found'): store.conflict_detail(second['id'],conflict['conflict_id'])


def test_missing_projection_is_rebuilt_even_with_current_marker(tmp_path):
    store=MaterialStore(tmp_path);material=opened(store)
    with store.connect() as db:
        original=db.execute('SELECT data FROM material_revisions').fetchone()[0]
        db.execute('DELETE FROM material_read_index')
    reopened=MaterialStore(tmp_path)
    assert reopened.list_page()['materials'][0]['id']==material['id']
    with reopened.connect() as db: assert db.execute('SELECT data FROM material_revisions').fetchone()[0]==original


def test_large_candidate_metadata_stays_in_detail_not_summary(tmp_path):
    store=MaterialStore(tmp_path);material=opened(store)
    material=store.start_candidate(request(material))['material'];cid=material['candidates'][0]['id']
    blocks=[dict(id='b'+str(i),type='text',text='candidate',source_refs=[{'scope_id':'page:1'}]) for i in range(10000)]
    warnings=['Warning '+str(i) for i in range(10)]
    store.finish_candidate(material['id'],cid,blocks,complete=True,warnings=warnings,metadata={'unresolved':['b'+str(i) for i in range(10000)]})
    summary=store.list_page()['materials'][0]['candidates'][0]
    assert summary['block_count']==10000 and summary['warning_count']==10 and len(summary['warnings'])==3
    assert len(json.dumps(summary))<3000 and 'differences' not in summary and 'metadata' not in summary
    assert store.candidate_detail(material['id'],cid)['warnings']==warnings


def test_shared_formula_followers_and_array_metadata_across_windows(tmp_path):
    from io import BytesIO
    from zipfile import ZipFile, ZIP_DEFLATED
    from lxml import etree
    from openpyxl import Workbook
    from pdf_extraction.formats.excel import NS
    book=Workbook();book.active.title='Values'
    for address,value in {'B2':20,'B3':30,'B150':1500,'D2':42,'D3':44}.items(): book.active[address]=value
    book.create_sheet('Empty').sheet_state='hidden'
    raw=BytesIO();book.save(raw);changed=BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as source,ZipFile(changed,'w',ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data=source.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                tree=etree.fromstring(data)
                for address,attrs,expression in [('B2',{'t':'shared','si':'7','ref':'B2:B150'},'A2*10'),('B3',{'t':'shared','si':'7'},None),('B150',{'t':'shared','si':'7'},None),('D2',{'t':'array','ref':'D2:D3'},'ROW(D2:D3)')]:
                    cell=tree.xpath(f'//s:c[@r="{address}"]',namespaces={'s':NS})[0]
                    formula=etree.Element('{'+NS+'}f',**attrs);formula.text=expression;cell.insert(0,formula)
                    if address=='B150': cell.remove(cell.find('{'+NS+'}v'))
                data=etree.tostring(tree)
            target.writestr(entry,data)
    path=tmp_path/'shared.xlsx';path.write_bytes(changed.getvalue());fingerprint=sha256(path.read_bytes()).hexdigest()
    follower=read_material(path,fingerprint,row=3,column=2,row_count=1,column_count=1)['cells'][0][0]
    assert follower['formula'] is None and follower['formula_present']
    assert follower['formula_attributes']=={'t':'shared','si':'7'}
    assert follower['formula_anchor']=='B2' and follower['formula_reference']=='B2:B150'
    assert follower['cached']=='30' and follower['cache_status']=='present'
    assert 'No expression is expanded or calculated' in follower['formula_note']
    anchor=read_material(path,fingerprint,row=2,column=2,row_count=1,column_count=1)['cells'][0][0]
    assert anchor['formula']=='=A2*10' and anchor['formula_anchor']=='B2'
    far=read_material(path,fingerprint,row=150,column=2,row_count=1,column_count=1)['cells'][0][0]
    assert far['formula'] is None and far['cached'] is None and far['cache_status']=='missing'
    assert far['formula_anchor']=='B2'
    array=read_material(path,fingerprint,row=2,column=4,row_count=1,column_count=1)['cells'][0][0]
    assert array['formula']=='=ROW(D2:D3)' and array['formula_kind']=='array' and array['formula_reference']=='D2:D3'
    assert array['formula_attributes']=={'t':'array','ref':'D2:D3'} and array['cached']=='42'
    empty=read_material(path,fingerprint,sheet='Empty')
    assert empty['state']=='hidden' and empty['cells'][0][0]['value'] is None
    assert path.read_bytes()==changed.getvalue()


def test_reader_version_change_never_reuses_previous_semantics(tmp_path,monkeypatch):
    from pdf_extraction.evidence import material_reader
    path=tmp_path/'original.html';path.write_bytes(b'<html><body>Current source text</body></html>')
    fingerprint=sha256(path.read_bytes()).hexdigest();current_version=material_reader.VERSION
    monkeypatch.setattr(material_reader,'VERSION','material-reader/prior')
    material_reader.read_material(path,fingerprint)
    old_cache=next((tmp_path/'.reader-cache').glob('*.json'))
    envelope=json.loads(old_cache.read_text());envelope['result']['html']='Previous interpretation of source'
    envelope['checksum']=sha256(json.dumps(envelope['result'],sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    old_cache.write_text(json.dumps(envelope))
    monkeypatch.setattr(material_reader,'VERSION',current_version)
    latest=material_reader.read_material(path,fingerprint)
    assert latest['schema_version']==current_version and 'Current source text' in latest['html']
    assert len(list((tmp_path/'.reader-cache').glob('*.json')))==2
