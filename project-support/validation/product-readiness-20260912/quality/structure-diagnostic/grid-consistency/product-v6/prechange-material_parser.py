"""Explicit, local structural extraction into a separate human-review candidate.

This adapter deliberately does not call the legacy domain-producing pipeline.
It reuses pure HTML/OOXML parsers, positioned NativeExtractor and ruled-table
geometry; candidate artifacts are immutable and never human confirmations.
"""
from __future__ import annotations

from contextlib import closing
from collections import defaultdict
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re

from ..evidence.material_reader import _manifest, html_anchor
from ..intake.registry import read_snapshot

VERSION = 'material-structural-parser/5'
POLICY = {'adapter': VERSION, 'semantic_processing': False, 'external_models': False, 'ocr': False, 'pdf_window_pages': 3, 'pdf_backend': 'pdfium', 'pdf_marker_order': 'same-row-before-body/1', 'pdf_paragraph_grouping': 'same-page-source-preserving/1', 'human_adoption_required': True}
CONFIG_HASH = sha256(json.dumps(POLICY, sort_keys=True).encode()).hexdigest()
NUMBERING_MARKER = re.compile(r'(?:\d+(?:\.\d+)+[.)]?|\d+[.)]|[a-z][.)]|[•●▪◦])')


def _write(path, payload):
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)


def _block(source, locator, kind, text='', **extra):
    return {'id': 'block-' + sha256((source.snapshot_id + ':' + locator).encode()).hexdigest()[:24],
            'type': kind, 'text': text, 'level': 1, 'parent_id': None,
            'numbering': '', 'dependencies': [], 'source_refs': [], **extra}


def _static_text(element):
    """Keep explicit HTML breaks without inventing spaces within inline words."""
    from bs4 import Tag, NavigableString, Comment
    parts = []
    boundaries = {'p', 'div', 'li', 'dt', 'dd', 'pre', 'blockquote', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'}
    for child in element.children:
        if isinstance(child, Comment):
            continue
        if isinstance(child, NavigableString):
            parts.append(str(child))
        elif isinstance(child, Tag):
            if child.name in {'script', 'style', 'head', 'noscript'}:
                continue
            if child.name == 'br':
                parts.append('\n')
            else:
                value = _static_text(child)
                if child.name in boundaries:
                    if parts and not parts[-1].endswith('\n'):
                        parts.append('\n')
                    parts.extend([value, '\n'])
                else:
                    parts.append(value)
    return ''.join(parts)


def _has_content(blocks):
    return any((b['type'] in {'text', 'heading'} and b.get('text', '').strip()) or
               any(str(value).strip() for row in b.get('table', {}).get('rows', []) for value in row)
               for b in blocks)


def _stabilize_marker_order(blocks, page_width):
    """Move a separate marker before its unambiguous adjacent same-row body.

    Native glyph heights can put a marker just below its body in a y sort.
    Preserve every block and source reference and every nonmoved block's order;
    do not merge text, infer headings or reorder columns. Ambiguous pairs stay.
    """
    positions = {b['id']: i for i, b in enumerate(blocks)}
    proposals = defaultdict(list)
    tables = [b['source_refs'][0]['bbox'] for b in blocks if b['type'] == 'table']
    for marker in blocks:
        if marker['type'] != 'text' or not NUMBERING_MARKER.fullmatch(marker['text'].strip()):
            continue
        mx0, my0, mx1, my1 = marker['source_refs'][0]['bbox']
        matches = []
        for body in blocks:
            if (body['type'] != 'text' or body['id'] == marker['id'] or not body['text'].strip()
                    or NUMBERING_MARKER.fullmatch(body['text'].strip())
                    or positions[body['id']] >= positions[marker['id']]):
                continue
            x0, y0, x1, y1 = body['source_refs'][0]['bbox']
            gap = x0 - mx1
            if not 0 <= gap <= min(page_width * .12, max(18., 6 * (y1 - y0))):
                continue
            # Same-row overlap reuses the existing native fragment criterion;
            # the local gap bound prevents pairing distant columns or footers.
            overlap = max(0., min(my1, y1) - max(my0, y0)) / max(.1, min(my1 - my0, y1 - y0))
            if overlap < .7:
                continue
            if any((tx0 <= (mx0 + mx1) / 2 <= tx1 and ty0 <= (my0 + my1) / 2 <= ty1)
                   or (tx0 <= (x0 + x1) / 2 <= tx1 and ty0 <= (y0 + y1) / 2 <= ty1)
                   for tx0, ty0, tx1, ty1 in tables):
                continue
            intervening = blocks[positions[body['id']] + 1:positions[marker['id']]]
            if any(b['type'] in {'table', 'image', 'heading'} for b in intervening):
                continue
            matches.append((gap, body, overlap))
        matches.sort(key=lambda item: item[0])
        if not matches or (len(matches) > 1 and matches[1][0] - matches[0][0] < 1):
            continue
        gap, body, overlap = matches[0]
        proposals[body['id']].append((marker, gap, overlap))
    accepted = {body: items[0] for body, items in proposals.items() if len(items) == 1}
    moved = {marker['id'] for marker, _, _ in accepted.values()}
    output, operations = [], []
    for block in blocks:
        if block['id'] in moved:
            continue
        if block['id'] in accepted:
            marker, gap, overlap = accepted[block['id']]
            output.append(marker)
            operations.append({'operation': 'marker_before_adjacent_same_row_text',
                               'marker_block_id': marker['id'], 'body_block_id': block['id'],
                               'horizontal_gap_points': gap, 'vertical_overlap_ratio': overlap,
                               'marker_source_refs': marker['source_refs'], 'body_source_refs': block['source_refs']})
        output.append(block)
    return output, operations


def _group_pdf_paragraphs(blocks, page_width, page_height, image_boxes=()):
    """Group adjacent native fragments; retain complete members as evidence.

    Reuse the existing same-page geometry/punctuation score, without its text
    rewriting or global sort. Explicit markers and structural/margin boundaries
    remain separate; page, table and image intersections are never traversed.
    """
    from ..assemble.paragraph_assembler import _score
    from ..models import Block, BlockType, Segment, BoundingBox, TextContent, Resolution, Page

    def ref(b):
        return b['source_refs'][0]

    def eligible(b):
        if b['type'] != 'text' or b.get('role') in {'header', 'footer', 'margin'}:
            return False
        return page_height*.10 < ref(b)['bbox'][1] and ref(b)['bbox'][3] < page_height*.90

    def clear(left, right):
        a, b = ref(left), ref(right)
        if a['page_index'] != b['page_index']:
            return False
        x0, y0 = min(a['bbox'][0], b['bbox'][0]), min(a['bbox'][1], b['bbox'][1])
        x1, y1 = max(a['bbox'][2], b['bbox'][2]), max(a['bbox'][3], b['bbox'][3])
        return not any(x0 < xx1 and x1 > xx0 and y0 < yy1 and y1 > yy0
                       for xx0, yy0, xx1, yy1 in obstacles)

    def model(b):
        return Block(id=b['id'], type=BlockType.PARAGRAPH,
            segments=[Segment(id=b['id'], page_index=ref(b)['page_index'],
                bbox=BoundingBox(**dict(zip(('x0', 'y0', 'x1', 'y1'), ref(b)['bbox']))))],
            content=TextContent(native_text=b['text'], resolved_text=b['text'],
                resolution=Resolution(selected_source='native', reason='native_fragment_geometry', confidence=0)))

    obstacles = [*image_boxes, *[ref(b)['bbox'] for b in blocks if b['type'] in {'table', 'image', 'heading'}]]
    pages = {ref(b)['page_index']: Page(page_index=ref(b)['page_index'], width=page_width,
        height=page_height, rotation=0, image_ref='', native_text_coverage=0, image_coverage=0,
        page_kind='born_digital') for b in blocks}
    output, operations, i = [], [], 0
    while i < len(blocks):
        first = blocks[i]; members = [first]; numbering = None; i += 1
        if eligible(first) and NUMBERING_MARKER.fullmatch(first['text'].strip()) and i < len(blocks):
            body = blocks[i]
            if eligible(body) and not NUMBERING_MARKER.fullmatch(body['text'].strip()) and clear(first, body):
                x0, y0, x1, y1 = ref(first)['bbox']; bx0, by0, bx1, by1 = ref(body)['bbox']
                overlap = max(0., min(y1, by1)-max(y0, by0)) / max(.1, min(y1-y0, by1-by0))
                if overlap >= .7 and 0 <= bx0-x1 <= min(page_width*.12, max(18., 6*(by1-by0))):
                    members.append(body); numbering = first['text'].strip(); i += 1
        if eligible(first) and (len(members) > 1 or not NUMBERING_MARKER.fullmatch(first['text'].strip())):
            while i < len(blocks):
                last, right = members[-1], blocks[i]
                # Short aligned labels can be neighboring table/header cells.
                # Leave them separate even when the generic geometry score agrees.
                if (len(last['text'].split()) < 5
                    or not eligible(right) or NUMBERING_MARKER.fullmatch(right['text'].strip())
                    or re.match(r'^(?:\d+(?:\.\d+)+[.)]?|\d+[.)]|[a-z][.)]|[•●▪◦])\s+', right['text'])
                    or not clear(last, right) or _score(model(last), model(right), pages) < .9):
                    break
                members.append(right); i += 1
        if len(members) == 1:
            output.append(first)
            continue
        identity = 'block-' + sha256(('paragraph:' + ':'.join(b['id'] for b in members)).encode()).hexdigest()[:24]
        group = {**first, 'id': identity, 'text': '\n'.join(b['text'] for b in members),
            'numbering': numbering or first.get('numbering', ''),
            'source_refs': [reference for b in members for reference in b['source_refs']]}
        output.append(group)
        operations.append({'operation': 'same_page_paragraph_group', 'group_block_id': identity,
            'numbering': group['numbering'], 'fragments': members})
    return output, operations


def _html_blocks(raw, source, output):
    from bs4 import BeautifulSoup, Tag, NavigableString, Comment
    from ..formats.html import parse_html, dom_path
    from ..formats.html_tables import build_table
    from ..formats.html_profiles import HtmlProfileError
    artifacts, warnings = [], []
    try:
        doc = parse_html(raw, source, {'template': 'auto', 'encoding': 'utf-8-sig'})
        _write(output / 'html-canonical.json', doc.model_dump(mode='json'))
        artifacts.append('html-canonical.json')
        warnings.extend(doc.issues)
    except HtmlProfileError:
        warnings.append('The saved HTML has no supported template profile; the complete static DOM is retained as a generic structural candidate for manual review.')
    soup = BeautifulSoup(raw, 'lxml')
    paths = {id(el): dom_path(el) for el in soup.find_all(True)}
    ids = {id(el): 'node-' + sha256(paths[id(el)].encode()).hexdigest()[:20] for el in soup.find_all(True)}
    node_by_id = {ids[id(el)]: el for el in soup.find_all(True)}
    blocks, stack, counts = [], [], {}
    def ref(el):
        return {'scope_id': 'html:document', 'locator': paths[id(el)], 'anchor': html_anchor(paths[id(el)])}
    def append(el, kind, text='', **extra):
        locator = paths[id(el)]
        counts[locator] = counts.get(locator, 0) + 1
        b = _block(source, locator + ':' + str(counts[locator]), kind, text, source_refs=[ref(el)], **extra)
        if kind == 'heading':
            level = int(el.name[1])
            while stack and stack[-1][0] >= level: stack.pop()
            b['level'] = level
        if stack:
            b['parent_id'] = stack[-1][1]
            b['dependencies'] = [stack[-1][1]]
        if kind == 'heading': stack.append((b['level'], b['id']))
        # Preserve a section's complete suffix even when its title follows a
        # punctuation mark without whitespace. Do not consume the first letter
        # of an ordinary title or invent support for hyphenated section IDs.
        match = re.match(r'^(§\s*\d+(?:\s*[a-z](?=[.)\s]|$))?)(?=[.)\s]|$)', text)
        if not match:
            match = re.match(r'^(\d+(?:\.\d+)*[.)]?)\s', text)
        if match: b['numbering'] = match[1]
        blocks.append(b)
        return b
    def walk(el):
        if el.name in {'script', 'style', 'head', 'noscript'}: return
        if el.name == 'img':
            append(el, 'image', el.get('alt', ''), image={'source_ref': ref(el), 'attachment': None, 'attribution': 'Image element in the saved original; external image assets are not fetched.'})
            return
        if el.name == 'table':
            table = build_table(el, ids)
            if table.column_count is None:
                warnings.append('Ambiguous HTML table spans retained as row text; repair table geometry manually at ' + paths[id(el)])
                rows = [[_static_text(c).strip() for c in r.find_all(['td', 'th'], recursive=False)] for r in el.find_all('tr') if r.find_parent('table') is el]
                width = max([1] + [len(r) for r in rows])
                rows = [r + [''] * (width - len(r)) for r in rows]
                merges = []
            else:
                rows = [[''] * max(1, table.column_count) for _ in range(max(1, table.row_count))]
                merges = []
                for cell in table.cells:
                    rows[cell.row][cell.column] = _static_text(node_by_id[cell.node_id]).strip()
                    if cell.rowspan > 1 or cell.colspan > 1:
                        merges.append({'row': cell.row, 'col': cell.column, 'rowspan': cell.rowspan, 'colspan': cell.colspan})
            caption = el.find('caption', recursive=False)
            append(el, 'table', '', table={'rows': rows or [['']], 'merges': merges, 'notes': [_static_text(caption).strip()] if caption else []})
            for img in el.find_all('img'): walk(img)
            return
        # Preserve all static body text, including outside known profile roots.
        semantic = {'p', 'li', 'dt', 'dd', 'pre', 'blockquote', 'figcaption', 'caption', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'}
        nested = el.find(['table', 'img'] + list(semantic))
        if el.name in semantic and nested is None:
            text = _static_text(el).strip()
            if text: append(el, 'heading' if re.fullmatch('h[1-6]', el.name) else 'text', text)
            return
        buffer = []
        def flush():
            text = ''.join(buffer).strip()
            if text: append(el, 'heading' if re.fullmatch('h[1-6]', el.name) else 'text', text)
            buffer.clear()
        for child in el.children:
            if isinstance(child, Comment): continue
            if isinstance(child, NavigableString): buffer.append(str(child))
            elif isinstance(child, Tag):
                if child.name == 'br': buffer.append('\n')
                elif child.name in {'span', 'a', 'b', 'strong', 'i', 'em', 'u', 's', 'small', 'sup', 'sub', 'code'} and not child.find(['table', 'img'] + list(semantic)):
                    buffer.append(_static_text(child))
                else:
                    flush(); walk(child)
        flush()
    walk(soup.body or soup)
    _write(output / 'html-static-candidate.json', {'source_sha256': source.content_hash, 'blocks': blocks, 'warnings': warnings})
    artifacts.append('html-static-candidate.json')
    return blocks, warnings, artifacts


def _excel_blocks(raw, source, output):
    from ..formats.excel import parse_excel
    from ..evidence.material_reader import _sheet_info
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
    doc = parse_excel(raw, source)
    _write(output / 'excel-canonical.json', doc.model_dump(mode='json'))
    blocks, warnings = [], list(doc.issues)
    for sheet in doc.sheets:
        scope_id = 'sheet:' + sheet.name
        heading = _block(source, sheet.part, 'heading', sheet.name, source_refs=[{'scope_id': scope_id, 'sheet': sheet.name, 'locator': sheet.part + '#A1', 'cell_range': 'A1'}])
        blocks.append(heading)
        # Sparse sheets are tiled; avoid allocating their million-row blank rectangle.
        tiles = {}
        for cell in sheet.cells:
            r, c = coordinate_to_tuple(cell.coordinate)
            tiles.setdefault(((r-1)//100, (c-1)//30), []).append(cell)
        for tile, cells in sorted(tiles.items()):
            points = [coordinate_to_tuple(c.coordinate) for c in cells]
            r1, c1 = min(p[0] for p in points), min(p[1] for p in points)
            r2, c2 = max(p[0] for p in points), max(p[1] for p in points)
            selected_merges = []
            for region in sheet.merged_ranges:
                x1, y1, x2, y2 = range_boundaries(region)
                if y1 <= r2 and y2 >= r1 and x1 <= c2 and x2 >= c1:
                    if x2-x1 < 300 and y2-y1 < 1000:
                        r1, c1, r2, c2 = min(r1,y1),min(c1,x1),max(r2,y2),max(c2,x2)
                        selected_merges.append((x1,y1,x2,y2))
                    else: warnings.append('Large merged range requires original review: ' + sheet.name + '!' + region)
            if (r2-r1+1)*(c2-c1+1) > 50000:
                raise ValueError('excel_candidate_table_limit')
            rows = [[''] * (c2-c1+1) for _ in range(r2-r1+1)]
            notes = []
            for cell in cells:
                r,c = coordinate_to_tuple(cell.coordinate)
                rows[r-r1][c-c1] = '=' + cell.formula if cell.formula is not None else (cell.text if cell.text is not None else cell.value or '')
                if cell.formula is not None: notes.append(cell.coordinate + ' saved value: ' + (cell.value if cell.value is not None else '[cache unavailable]') + '. Formula has not been calculated.')
            merges = [{'row': y1-r1, 'col': x1-c1, 'rowspan': y2-y1+1, 'colspan': x2-x1+1} for x1,y1,x2,y2 in selected_merges]
            cell_range = f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}'
            block = _block(source, sheet.part+'#'+cell_range, 'table', source_refs=[{'scope_id': scope_id, 'sheet': sheet.name, 'locator': sheet.part+'#'+f'{get_column_letter(c1)}{r1}', 'cell_range': cell_range}], parent_id=heading['id'], dependencies=[heading['id']], table={'rows':rows,'merges':merges,'notes':notes})
            blocks.append(block)
    warnings.append('All worksheets remain in required scope, including blank and hidden sheets. Charts, drawings, comments and native formatting require original inspection or manual image/notes supplementation; formulas are source text, not calculated results.')
    return blocks, warnings, ['excel-canonical.json']



def _pdf_tables(page, original_page, source, reference, output):
    """Reuse the existing local ruled-grid detector and merged-cell assembler.

    Words are assigned only by their retained native geometry. Unknown/unruled
    structures stay as text, ready for manual table insertion and correction.
    """
    import numpy as np
    from ..config import TableSettings
    from ..layout.detector import detect_table_boxes
    from ..parsers.table import _grid_positions, _merged_cell_specs, _compact_grid_specs
    bitmap = original_page.render(scale=1.5)
    image_path = output / f"pdf-page-{reference['page']:04d}.png"
    try:
        image = bitmap.to_pil().convert('RGB')
        image.save(image_path)
        array = np.asarray(image)
    finally:
        bitmap.close()
    candidates = []
    for number, box in enumerate(detect_table_boxes(image_path, TableSettings(backend='native', recognition_pipeline_enabled=False, img2table_enabled=False, gmft_enabled=False))):
        crop = array[box.y0:box.y1, box.x0:box.x1]
        xs, ys = _grid_positions(crop)
        if len(xs) < 2 or len(ys) < 2 or (len(xs)-1)*(len(ys)-1)>10000:
            continue
        specs, xs, ys = _compact_grid_specs(_merged_cell_specs(crop, xs, ys), xs, ys)
        rows = [['']*(len(xs)-1) for _ in range(len(ys)-1)]
        merges = []
        claimed = set()
        for row, col, rowspan, colspan in specs:
            x0,y0,x1,y1 = (box.x0+xs[col])/1.5,(box.y0+ys[row])/1.5,(box.x0+xs[col+colspan])/1.5,(box.y0+ys[row+rowspan])/1.5
            words = [w for w in page.words if x0 <= (w.bbox_points[0]+w.bbox_points[2])/2 < x1 and y0 <= (w.bbox_points[1]+w.bbox_points[3])/2 < y1]
            # Native parser order is retained within each geometric cell.
            rows[row][col] = ' '.join(w.text for w in words).strip()
            claimed.update(w.id for w in words)
            if rowspan>1 or colspan>1:
                merges.append({'row':row,'col':col,'rowspan':rowspan,'colspan':colspan})
        if not claimed:
            continue
        bbox=[(box.x0+xs[0])/1.5,(box.y0+ys[0])/1.5,(box.x0+xs[-1])/1.5,(box.y0+ys[-1])/1.5]
        block=_block(source, f"page:{reference['page_index']}:table:{number}", 'table', source_refs=[{**reference,'bbox':bbox}], table={'rows':rows,'merges':merges,'notes':[]})
        candidates.append((bbox,block))
    return candidates, image_path.name


def _pdf_blocks(raw, source, output, page_indices):
    from ..ingest.native_extractor import NativeExtractor
    from ..config import AppConfig
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(raw) as inventory:
        total = len(inventory)
    selected = sorted(set(range(total)) if page_indices is None else set(page_indices))
    if not selected or any(type(p) is not int or not 0 <= p < total for p in selected): raise ValueError('pdf_page_selection_invalid')
    config = AppConfig.from_yaml(Path(__file__).resolve().parents[3] / 'config' / 'pdf-intake-positioned.yaml')
    blocks, warnings, artifacts, usable, unresolved, ordering_repairs, paragraph_groups = [], [], [], [], [], [], []
    with pdfium.PdfDocument(raw) as original:
        for offset in range(0, len(selected), 3):
            mapping = selected[offset:offset+3]
            chunk = output / f'pdf-window-{offset//3+1:04d}.pdf'
            # PDFium already supports the readable source encryption modes.
            # Avoid an optional pypdf AES dependency for immutable window copies.
            with pdfium.PdfDocument.new() as window:
                window.import_pages(original, pages=mapping)
                with chunk.open('xb') as stream: window.save(stream)
            # The material route uses the already validated local positioned
            # backend directly. Importing a heavier backend cannot be bounded by
            # catching exceptions after it hangs; legacy intake keeps its config.
            native = NativeExtractor(POLICY['pdf_backend']).extract(chunk)
            if len(native.pages) != len(mapping):
                raise ValueError('native_pdf_page_count_mismatch')
            evidence = {'source_sha256': source.content_hash, 'config_hash': sha256(config.model_dump_json().encode()).hexdigest(), 'adapter_config_hash': CONFIG_HASH, 'material_backend_policy': POLICY['pdf_backend'], 'legacy_config_backend': config.native_extraction.backend, 'original_page_indices': mapping, 'backend': native.backend, 'version': native.version, 'pages': [asdict(p) for p in native.pages]}
            evidence_name = f'pdf-native-{offset//3+1:04d}.json'
            _write(output / evidence_name, evidence)
            artifacts.extend([chunk.name, evidence_name])
            for local, page in enumerate(native.pages):
                index = mapping[local]
                scope_id = f'page:{index+1}'
                reference = {'scope_id':scope_id, 'page_index':index, 'page':index+1}
                lines = page.text_lines or page.words
                with closing(original[index]) as original_page:
                    tables, page_image = _pdf_tables(page, original_page, source, reference, output)
                    artifacts.append(page_image)
                    page_blocks = [b for _, b in tables]
                    for n,line in enumerate(lines):
                        text = line.text.strip()
                        if not text: continue
                        x0,y0,x1,y1 = line.bbox_points
                        # Suppress only lines fully contained by a table; partial
                        # overlaps remain visible for human association review.
                        if any(b[0]<=x0 and b[1]<=y0 and b[2]>=x1 and b[3]>=y1 for b,_ in tables):
                            continue
                        block = _block(source, f'page:{index}:line:{n}', 'text', text, source_refs=[{**reference, 'bbox':list(line.bbox_points), 'native_id':line.id}])
                        page_blocks.append(block)
                    page_blocks.sort(key=lambda b:(b['source_refs'][0]['bbox'][1],b['source_refs'][0]['bbox'][0]))
                    page_blocks, repairs = _stabilize_marker_order(page_blocks, page.width_points)
                    ordering_repairs.extend(repairs)
                    image_boxes = []
                    for obj in original_page.get_objects():
                        if obj.type == pdfium.raw.FPDF_PAGEOBJ_IMAGE:
                            left, bottom, right, top = obj.get_bounds()
                            image_boxes.append([left, page.height_points-top, right, page.height_points-bottom])
                    page_blocks, groups = _group_pdf_paragraphs(page_blocks, page.width_points, page.height_points, image_boxes)
                    paragraph_groups.extend(groups)
                    blocks.extend(page_blocks)
                    has_content = _has_content(page_blocks)
                    if has_content:
                        usable.append(scope_id)
                    # Keep embedded source images beside native structural text.
                    image_count = 0
                    for obj in original_page.get_objects():
                        if obj.type != pdfium.raw.FPDF_PAGEOBJ_IMAGE: continue
                        left,bottom,right,top = obj.get_bounds()
                        bbox = [left,page.height_points-top,right,page.height_points-bottom]
                        ref = {**reference,'bbox':bbox}
                        blocks.append(_block(source, f'page:{index}:image:{image_count}', 'image', image={'source_ref':ref,'attachment':None,'attribution':'Embedded image in the bound original PDF.'}, source_refs=[ref]))
                        image_count += 1
                    if not has_content:
                        unresolved.append({'scope_id':scope_id, 'code':'native_text_unavailable',
                                           'message':f'Page {index+1} has no usable native text or table content. Inspect the original; a scan or a blank page still requires an explicit human check.'})
                        warnings.append(f'Page {index+1} has no positioned native text. Inspect the complete original page and manually transcribe or supplement omitted text. No OCR/model was run.')
                        if not image_count:
                            blocks.append(_block(source,f'page:{index}:untranscribed','image', image={'source_ref':reference,'attachment':None,'attribution':'Complete original page; inspect for scanned or blank content.'},source_refs=[reference]))
    _write(output / 'pdf-order-repairs.json', {'method': POLICY['pdf_marker_order'], 'source_sha256': source.content_hash, 'operations': ordering_repairs})
    artifacts.append('pdf-order-repairs.json')
    if ordering_repairs:
        warnings.append('Separate numbering or list markers were placed before adjacent same-row text using native geometry. Text and source references are unchanged; verify the associations against the original.')
    _write(output / 'pdf-paragraph-groups.json', {'method': POLICY['pdf_paragraph_grouping'], 'source_sha256': source.content_hash, 'groups': paragraph_groups})
    artifacts.append('pdf-paragraph-groups.json')
    if paragraph_groups:
        warnings.append('Adjacent same-page text fragments were grouped for editing. Original text, numbering and every source fragment remain in extraction evidence; check paragraph boundaries against the original.')
    warnings.append('PDF candidate uses the existing positioned PDFium parser in mapped three-page windows. Ruled-table grids use existing local geometry assembly. Reading order, headings and unruled tables may need manual correction; source page images and complete native evidence are retained. No Requirement classifier, semantic processor, OCR model or external service runs.')
    return blocks,warnings,artifacts,[f'page:{p+1}' for p in selected],usable,unresolved


def parse_material(source, source_root, output, *, page_indices=None):
    """Parse one captured snapshot, returning a non-adopted review candidate.

    ``output`` must be a fresh job-owned directory outside the source root. The
    captured bytes are pinned before any parser works; subsequent source changes
    are the owning store's staleness concern, never relabelled by this adapter.
    """
    from ..contracts.source import Snapshot
    from ..formats.router import route_format
    source = source if isinstance(source, Snapshot) else Snapshot.model_validate(source)
    source_root, output = Path(source_root), Path(output)
    if output.resolve().is_relative_to(source_root.resolve()): raise ValueError('output_must_be_outside_source_root')
    raw = read_snapshot(source, source_root)
    kind = route_format(source.relative_path, source.file_format, raw)
    output.mkdir(parents=True, exist_ok=False)
    suffix = Path(source.relative_path).suffix.lower()
    with (output / ('source' + suffix)).open('xb') as stream: stream.write(raw)
    manifest = _manifest(raw, suffix)
    if kind == 'html': blocks,warnings,artifacts = _html_blocks(raw,source,output)
    elif kind == 'excel': blocks,warnings,artifacts = _excel_blocks(raw,source,output)
    else: blocks,warnings,artifacts,covered,usable,unresolved = _pdf_blocks(raw,source,output,page_indices)
    if kind != 'pdf':
        covered = [s['id'] for s in manifest['scope']]
        usable, unresolved = (covered, []) if _has_content(blocks) else ([], [{'scope_id':s, 'code':'extracted_content_empty', 'message':'No usable text or table content was extracted. Inspect the original.'} for s in covered])
    unprocessed = [s['id'] for s in manifest['scope'] if s['id'] not in covered]
    status = 'empty' if not usable else 'partial' if unresolved or unprocessed else 'candidate_available'
    candidate = {'schema_version':VERSION,'parser_version':VERSION,'source_sha256':source.content_hash,
                 'snapshot_id':source.snapshot_id,'config_hash':CONFIG_HASH,'blocks':blocks,'scope':manifest['scope'], 'covered_scope':covered,
                 'processed_scope':covered, 'usable_scope':usable, 'unprocessed_scope':unprocessed,
                 'warnings':list(dict.fromkeys(manifest['warnings']+warnings)), 'unresolved':unresolved,
                 'canonical_artifacts':artifacts,'status':status,
                 'review_policy':'human_material_confirmation_required', 'requirement_status':'not_connected'}
    _write(output / 'candidate.json',candidate)
    return candidate
