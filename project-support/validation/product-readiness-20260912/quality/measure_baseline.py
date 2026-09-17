"""Immutable, local pre-change material-parser evidence. No business-store writes.

Run with System2's interpreter and PYTHONPATH=system2/src from the workstream root.
Frozen test Snapshot identities represent engineering fixtures, never decisions.
"""
from pathlib import Path
from hashlib import sha256
from collections import Counter
import importlib.metadata
import json
import platform
import sys
import time

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent


def digest(path):
    return sha256(path.read_bytes()).hexdigest()


def write(path, data):
    with path.open('x', encoding='utf-8') as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2)


def freeze():
    from openpyxl import load_workbook
    import pypdfium2 as pdfium
    specs = [
        ('asc-farm', 'system2/data/inputs/ASC-STD-001-ASC-Farm-Standard-V1.0.1-Aug-2025.pdf', [27,28], 'farm-standard-p028-p029.json'),
        ('asc-interpretation', 'system2/data/inputs/ASC-INT-001-ASC-Farm-Standard-Interpretation-Manual-V1.0-May-2025.pdf', [18,19,20], 'interpretation-manual-p019-p021.json'),
        ('asc-audit', 'system2/data/inputs/download-module-691-ASC-Salmon-Audit-Manual_v1.4.pdf', [0,1,2], 'audit-manual-p001-p003.json'),
        ('asc-salmon-cod', 'system2/data/inputs/ASC-STD-010-ASC-Salmon-and-Cod-Standard-V1.5-Oct-2025-1.pdf', [17], 'salmon-cod-standard-p018-v2.json'),
    ]
    samples=[]
    for ident,rel,pages,ann in specs:
        ap=ROOT/'system2/gold/requirements/annotations'/ann
        gold=json.loads(ap.read_text())
        samples.append(dict(id=ident,path=rel,page_indices=pages,split='development',family='ASC_PDF',annotation={'path':str(ap.relative_to(ROOT)), 'sha256':digest(ap), 'trust':'existing source-backed engineering annotation; not business expert Gold', 'counts':{k:len(gold[k]) for k in ('segments','requirements','excluded_regions')},'denominator_status':'UNMEASURED: complete content-unit and binary positive/negative denominator not yet validated'}))
    for p in sorted((ROOT/'system2/gold/pdfs').glob('*.pdf')):
        ap=ROOT/'system2/gold/annotations'/p.with_suffix('.json').name
        with pdfium.PdfDocument(p.read_bytes()) as pdf:
            pages=list(range(len(pdf)))
        samples.append(dict(id=p.stem,path=str(p.relative_to(ROOT)),page_indices=pages,split='mechanism_regression',family='synthetic_PDF',annotation={'path':str(ap.relative_to(ROOT)), 'sha256':digest(ap),'trust':'synthetic authored visible-content expectations; not real-format acceptance','denominator_status':'UNMEASURED for real-material quality; expected fields can diagnose mechanisms'}))
    for ident,pattern in [('pa002-html','PA002*.html'),('cs001-html','CS001*.html')]:
        p=next((ROOT/'system1/Data').rglob(pattern))
        samples.append(dict(id=ident,path=str(p.relative_to(ROOT)),split='development',family='Lovdata_HTML' if ident.startswith('pa') else 'ASC_HTML',annotation={'path':'system2/tests/fixtures/html/source-windows-20260907.json','sha256':digest(ROOT/'system2/tests/fixtures/html/source-windows-20260907.json'),'trust':'previous manually selected source windows only; no complete unit denominator','denominator_status':'UNMEASURED'}))
    xp=ROOT.parent/'03_GLOBALGAP_Standards/IFA v6 Smart _ AQ_en_20240827 - Prefilled_00.xlsx'
    workbook=load_workbook(xp,read_only=False,data_only=False)
    sheet=workbook['CL - IFA v6 Smart - AQ']
    cells=[{'coordinate':c.coordinate,'value':c.value,'data_type':c.data_type} for row in sheet.iter_rows(min_row=9,max_row=24,min_col=1,max_col=8) for c in row]
    write(HERE/'xlsx-source-window.json',{'source_sha256':digest(xp),'sheet':sheet.title,'cell_range':'A9:H24','cells':cells,'merged_ranges':[str(m) for m in sheet.merged_cells.ranges], 'trust':'independent openpyxl source inventory, frozen before prediction; visual/business review pending'})
    workbook.close()
    samples.append(dict(id='globalgap-xlsx',path=str(xp.relative_to(ROOT.parent)).join(['../','']),split='development',family='GLOBALGAP_XLSX',annotation={'path':str((HERE/'xlsx-source-window.json').relative_to(ROOT)),'sha256':digest(HERE/'xlsx-source-window.json'),'trust':'independent openpyxl source inventory; engineering cell correspondence only','frozen_scope':"'CL - IFA v6 Smart - AQ'!A9:H24",'nonempty_cell_denominator':sum(c['value'] is not None for c in cells),'all_cell_denominator':len(cells),'denominator_status':'cell equality measurable; semantic and global structure quality UNMEASURED'},parse_scope='entire bounded 174604-byte workbook; score only frozen cell window'))
    for ident,pattern,pages,exposure in [('pa057-reserved','PA057*.pdf',[4,5],'Prior diagnosis pp8-11 in convergence ledger; same document exposed.'),('cs010-reserved','CS010*.pdf',[2,3],'Entire 4-page document parsed in prior material acceptance; not independent.')]:
        p=next((ROOT/'system1/Data').rglob(pattern))
        samples.append(dict(id=ident,path=str(p.relative_to(ROOT)),page_indices=pages,split='reserved_do_not_parse',family='Norwegian_guidance_PDF' if ident.startswith('pa') else 'GLOBALGAP_PDF',annotation={'trust':'source-only selection; not annotated','denominator_status':'UNMEASURED'},holdout_status='HISTORICALLY_EXPOSED_NOT_INDEPENDENT',exposure=exposure,eligibility='Engineering reference only; live System1 eligibility not established'))
    for n,sample in enumerate(samples,1):
        path=ROOT/sample['path']
        sample.update(sha256=digest(path),bytes=path.stat().st_size,engineering_source_id=f'QS{n:03d}')
        if 'page_indices' in sample: sample['pdf_page_numbers']=[n+1 for n in sample['page_indices']]
    files=list((ROOT/'system2/src/pdf_extraction').rglob('*.py'))+[ROOT/'system2/config/pdf-intake-positioned.yaml',ROOT/'system2/pyproject.toml',ROOT/'system2/uv.lock']
    fp={str(p.relative_to(ROOT)):digest(p) for p in files}
    write(HERE/'fingerprint-before.json',{'utc':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'python':sys.version,'platform':platform.platform(),'files':fp,'dependencies':sorted((d.metadata['Name'],d.version) for d in importlib.metadata.distributions())})
    write(HERE/'manifest.json',{'version':'product-readiness-quality/1','frozen_at':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'samples':samples,'rules':['Never treat these fixture identities as human INCLUDE decisions.','Never parse reserved samples in baseline.','No external models, OCR, downloads or full large-PDF run.','No untouched independent holdout currently established.','Existing Gold/history is read-only.','Exact-text diagnostics are not content recall or precision/recall acceptance.']})


def run():
    from pdf_extraction.orchestration.material_parser import parse_material
    manifest=json.loads((HERE/'manifest.json').read_text())
    target=HERE/'before'
    target.mkdir(exist_ok=False)
    for s in manifest['samples']:
        if s['split']=='reserved_do_not_parse':continue
        path=ROOT/s['path']
        if digest(path)!=s['sha256']:raise RuntimeError('source drift: '+s['id'])
        source=dict(source_id=s['engineering_source_id'],snapshot_id=s['engineering_source_id']+'-001',relative_path=path.name,content_hash=s['sha256'],file_format=path.suffix.lstrip('.'),operator_selection_decision='INCLUDE',selection_status='INCLUDE',snapshot_status='STORED',source_status='CURRENT',download_status='SUCCESS',registry_sha256=digest(HERE/'manifest.json'))
        started=time.monotonic()
        try:
            result=parse_material(source,path.parent,target/s['id'],page_indices=s.get('page_indices'))
            report={'id':s['id'],'status':result['status'],'block_count':len(result['blocks']),'block_types':dict(Counter(b['type'] for b in result['blocks'])),'covered_scope':result['covered_scope'],'total_scope_count':len(result['scope']),'warnings':result['warnings'],'unresolved':result['unresolved'],'requirement_status':result['requirement_status'],'artifacts':result['canonical_artifacts']}
        except Exception as exc:
            report={'id':s['id'],'status':'failed','error_type':type(exc).__name__,'error':str(exc)}
        report['elapsed_seconds']=round(time.monotonic()-started,3)
        write(target/(s['id']+'-summary.json'),report)
        print(json.dumps(report,ensure_ascii=False),flush=True)
    fingerprints=json.loads((HERE/'fingerprint-before.json').read_text())['files']
    write(HERE/'source-preservation-after-baseline.json',{'samples':{s['id']:digest(ROOT/s['path'])==s['sha256'] for s in manifest['samples']},'code_config_lock_unchanged':{p:digest(ROOT/p)==h for p,h in fingerprints.items()}})


if __name__=='__main__':
    {'freeze':freeze,'run':run}[sys.argv[1]]()
