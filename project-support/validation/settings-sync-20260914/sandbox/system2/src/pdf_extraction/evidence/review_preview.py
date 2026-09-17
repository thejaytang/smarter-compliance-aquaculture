"""Local, source-bound evidence previews. No remote assets or original scripts execute."""
import base64
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import tempfile
import time
import os
import signal
from lxml import html, etree

VERSION='original-preview/7'


def xpath(locator):
    locator=locator.split('::')[0]
    parts=[]
    for part in locator.split(' > '):
        m=re.fullmatch(r'([a-zA-Z][\w-]*):nth-of-type\((\d+)\)',part)
        if not m: return None
        parts.append(m[1]+'['+m[2]+']')
    return '/'+'/'.join(parts)


def sanitized_region(raw, references, full=False):
    tree=html.document_fromstring(raw)
    located=[]
    for ref in references:
        expr=xpath(ref.get('locator',''))
        if expr:located.extend(tree.xpath(expr))
    warnings=[]
    root=tree
    if not full and located:
        ancestors=[list(reversed([n]+list(n.iterancestors()))) for n in located]
        for level in zip(*ancestors):
            if all(n is level[0] for n in level):root=level[0]
            else:break
    elif not full:
        warnings.append('The saved source location could not be resolved. Showing the complete local snapshot.')
    # Keep published markup and embedded CSS, never execute or fetch source assets.
    styles=[]
    for node in tree.xpath('//style'):
        styles.append(re.sub(r'@import[^;]*;|url\([^)]*\)', '', node.text or '', flags=re.I))
    if tree.xpath('//link[@href] | //script | //iframe | //img[not(starts-with(@src,"data:image/"))]'):
        warnings.append('External styles, images and scripts are unavailable in this offline preview. Verify layout-dependent content against the full original.')
    for node in list(tree.iter()):
        if not isinstance(node.tag,str):continue
        if node.tag.lower() in {'script','iframe','object','embed','link','base','meta','form'}:
            if node.getparent() is not None: node.getparent().remove(node)
            continue
        for key in list(node.attrib):
            val=node.attrib[key]
            if key.lower().startswith('on') or key.lower() in {'srcset','action','formaction','http-equiv','nonce','integrity'}:
                del node.attrib[key]
            elif key in {'href','src','xlink:href'} and not (key=='src' and re.match(r'^data:image/(png|jpeg|gif|webp);base64,',val)):
                del node.attrib[key]
            elif key=='style':node.attrib[key]=re.sub(r'url\([^)]*\)', '',val,flags=re.I)
    for n in located:
        n.attrib['style']=n.attrib.get('style','')+';outline:2px solid #3977bd;outline-offset:3px;'
    body=html.tostring(root,encoding='unicode')
    csp="default-src 'none'; img-src data:; style-src 'unsafe-inline'; font-src 'none'; form-action 'none'; base-uri 'none'"
    output='<!doctype html><html><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="'+csp+'"><style>'+''.join(styles)+'</style></head><body>'+body+'</body></html>'
    return output,warnings


def preview(path, expected_hash, refs, cache, full=False):
    path=Path(path);raw=path.read_bytes()
    if sha256(raw).hexdigest()!=expected_hash:raise ValueError('original_version_changed')
    cache=Path(cache);cache.mkdir(parents=True,exist_ok=True)
    key=sha256(json.dumps([VERSION,expected_hash,refs,full],sort_keys=True).encode()).hexdigest()
    marker=cache/(key+'.json')
    if marker.exists():
        cached=json.loads(marker.read_bytes())
        if cached.get('image') or full or time.time()-marker.stat().st_mtime<30:return cached
    suffix=path.suffix.lower();result={'schema_version':VERSION,'source_sha256':expected_hash,'warnings':[]}
    if suffix in {'.html','.htm'}:
        markup,warnings=sanitized_region(raw,refs,full)
        result.update(kind='html_region',html=markup,warnings=warnings,label='Rendered from the bound local original. Highlighted elements correspond to this unit.')
        chrome=Path('/Applications/Google Chrome.app/Contents/MacOS/Google Chrome')
        if chrome.exists() and not full:
            with tempfile.TemporaryDirectory(dir=cache) as tmp:
                file=Path(tmp)/'region.html';file.write_text(markup)
                png=Path(tmp)/'region.png'
                try:
                    with (Path(tmp)/'browser.log').open('wb') as log:
                        proc=subprocess.Popen([str(chrome),'--headless','--disable-gpu','--disable-crash-reporter','--disable-crashpad-for-testing','--no-first-run','--disable-background-networking','--disable-extensions','--user-data-dir='+str(Path(tmp)/'profile'),'--window-size=1100,900','--screenshot='+str(png),file.as_uri()],stdout=log,stderr=log,start_new_session=True)
                        try:
                            deadline=time.monotonic()+12
                            while not png.exists() and proc.poll() is None and time.monotonic()<deadline:time.sleep(.1)
                            if not png.exists():raise OSError('Original image renderer did not produce a screenshot.')
                            from PIL import Image
                            with Image.open(png) as check:check.verify()
                            result['image']='data:image/png;base64,'+base64.b64encode(png.read_bytes()).decode()
                        finally:
                            if proc.poll() is None:
                                os.killpg(proc.pid,signal.SIGTERM)
                                try:proc.wait(timeout=2)
                                except subprocess.TimeoutExpired:os.killpg(proc.pid,signal.SIGKILL);proc.wait()
                    result['image_label']='Original region snapshot, first viewport. Expand context to inspect the full range.'
                except (OSError,subprocess.SubprocessError) as exc:
                    (cache/(key+'.log')).write_text(str(exc)+'\n'+str(getattr(exc,'stderr',b''))[-1800:])
                    result['warnings'].append('Image rendering is unavailable; inspect the isolated original markup below.')
    elif suffix=='.pdf':
        page_indices=[r['page_index'] for r in refs if type(r.get('page_index')) is int]
        if not page_indices:
            result.update(kind='unlocated',label='The saved PDF range has no verified page location. Open the full original or inspect original pages before confirming this range.')
            return result
        import pypdfium2 as pdfium
        doc=pdfium.PdfDocument(raw);page_index=page_indices[0]
        if any(not 0<=p<len(doc) for p in page_indices):
            doc.close()
            raise ValueError('original_page_out_of_range')
        page=doc[page_index];width,height=page.get_size();image=page.render(scale=1.5).to_pil()
        boxes=[r['bbox'] for r in refs if r.get('page_index')==page_index and r.get('bbox')]
        if boxes and not full:
            boxes=[[b['x0'],b['y0'],b['x1'],b['y1']] if isinstance(b,dict) else b for b in boxes]
            left=min(b[0] for b in boxes);top=min(b[1] for b in boxes);right=max(b[2] for b in boxes);bottom=max(b[3] for b in boxes)
            if max(right,bottom)<=1:left*=width;right*=width;top*=height;bottom*=height
            if right>left and bottom>top:image=image.crop((max(0,int(left*1.5)-20),max(0,int(top*1.5)-20),min(image.width,int(right*1.5)+20),min(image.height,int(bottom*1.5)+20)))
        buffer=BytesIO();image.save(buffer,format='PNG');page.close();doc.close()
        result.update(kind='image_region',image='data:image/png;base64,'+base64.b64encode(buffer.getvalue()).decode(),page=page_index+1,width=width,height=height,label=f'Bound original PDF, page {page_index+1}. Expand the original for other pages and context.')
        pages=sorted({p+1 for p in page_indices})
        if len(pages)>1:result['warnings'].append('This content spans PDF pages '+', '.join(map(str,pages))+'. The crop shows only page '+str(page_index+1)+'. Expand the original and check every referenced page before confirming the range.')
        if any(r.get('split_text_range') or r.get('cell_text_range') for r in refs):result['warnings'].append('This split part maps to a character range within the retained source region. The region can include neighboring text; correct the location if a tighter crop is needed.')
    elif suffix=='.xlsx':
        result=spreadsheet_region(raw,refs)
        result.update(schema_version=VERSION,source_sha256=expected_hash)
    else:raise ValueError('region_preview_format_not_supported')
    temp=marker.with_suffix('.tmp');temp.write_text(json.dumps(result));temp.replace(marker)
    return result


def spreadsheet_region(raw, refs):
    """Resolve ZIP-part locations against this exact original workbook."""
    from zipfile import ZipFile
    import posixpath
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple
    with ZipFile(BytesIO(raw)) as archive:
        rels=etree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets={r.get('Id'):posixpath.normpath('xl/'+r.get('Target')) if not r.get('Target').startswith('/') else r.get('Target').lstrip('/') for r in rels}
        book=etree.fromstring(archive.read('xl/workbook.xml'))
        names={targets[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]:s.get('name') for s in book.findall('.//{*}sheet')}
    positions=[]
    for ref in refs:
        match=re.fullmatch(r'(xl/worksheets/[^#]+)#([A-Z]+[0-9]+)',ref.get('locator',''))
        if match and match[1] in names:positions.append((names[match[1]],*coordinate_to_tuple(match[2])))
    wb=load_workbook(BytesIO(raw),data_only=False);cached=load_workbook(BytesIO(raw),data_only=True)
    name=positions[0][0] if positions else wb.sheetnames[0];ws=wb[name]
    points=[p for p in positions if p[0]==name]
    row=max(1,min(p[1] for p in points)-1) if points else 1
    column=max(1,min(p[2] for p in points)-1) if points else 1
    cells=[]
    for rr in ws.iter_rows(min_row=row,max_row=min(ws.max_row,row+39),min_col=column,max_col=min(ws.max_column,column+11)):
        cells.append([{'address':c.coordinate,'value':str(c.value) if c.value is not None else None,'cached':str(cached[name][c.coordinate].value) if cached[name][c.coordinate].value is not None else None,'hidden_row':bool(ws.row_dimensions[c.row].hidden),'hidden_column':bool(ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(c.column)].hidden)} for c in rr])
    result={'kind':'spreadsheet','sheets':wb.sheetnames,'sheet':name,'state':ws.sheet_state,'row':row,'column':column,'rows':ws.max_row,'columns':ws.max_column,'merged':[str(r) for r in ws.merged_cells.ranges],'cells':cells,'warnings':[],'label':'Bound original worksheet cells, including formulas and saved caches. Use the original file for published layout.'}
    if not positions:result['warnings'].append('Exact cell location is unavailable. Showing the beginning of the original worksheet; report the location for correction.')
    if len({p[0] for p in positions})>1:result['warnings'].append('This unit references more than one worksheet. Check every referenced range before confirming.')
    wb.close();cached.close();return result
