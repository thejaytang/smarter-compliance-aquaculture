import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from system1.formula_cache import FormulaCacheError, WorkbookCalculator
import source_updater as u


class FormulaCacheTests(unittest.TestCase):
    def test_real_template_all_caches_preserve_formulas_history_and_structure(self):
        original = Path(__file__).resolve().parents[2] / 'Requirement_Source_Registry.xlsx'
        before = original.read_bytes()
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / 'registry.xlsx'
            shutil.copy2(original, path)
            wb = load_workbook(path)
            # Normalize only the presentation before testing cache-only preservation.
            # The presentation contract separately verifies pre-migration business data.
            from system1.workbook_presentation import prepare_browser_workbook
            prepare_browser_workbook(wb)
            formulas = {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type == 'f'}
            cells = {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type != 'f' and c.value is not None}
            hidden = {s.title: s.sheet_state for s in wb}
            charts = {s.title: len(s._charts) for s in wb}
            u.save_workbook_atomic(wb, path, u.workbook_mtime(path))
            wb.close()
            after = load_workbook(path)
            values = load_workbook(path, data_only=True)
            self.assertEqual(formulas, {(s.title, c.coordinate): c.value for s in after for row in s for c in row if c.data_type == 'f'})
            self.assertEqual(cells, {(s.title, c.coordinate): c.value for s in after for row in s for c in row if c.data_type != 'f' and c.value is not None})
            self.assertEqual(hidden, {s.title: s.sheet_state for s in after})
            self.assertEqual(charts, {s.title: len(s._charts) for s in after})
            src = after['Source Register']; headers = u.workbook_headers(src, 2)
            rows = [r for r in range(3, src.max_row + 1) if src.cell(r, headers['source_id']).value]
            for r in rows:
                self.assertEqual(values['Source Register'].cell(r, headers['selection_status']).value,
                                 u.selection_from_scores(u.record_from_row(src, r, headers)))
            self.assertEqual(values['Dashboard']['A5'].value, len(rows))
            self.assertEqual(values['Dashboard']['C5'].value,
                             sum(values['Source Register'].cell(r, headers['selection_status']).value == 'INCLUDE' for r in rows))
            ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            with ZipFile(path) as z:
                formula_nodes = [c for name in z.namelist() if name.startswith('xl/worksheets/sheet') and name.endswith('.xml')
                                 for c in ET.fromstring(z.read(name)).iter(ns + 'c') if c.find(ns + 'f') is not None]
                self.assertEqual(len(formula_nodes), len(formulas))
                self.assertTrue(all(c.find(ns + 'v') is not None for c in formula_nodes))
            after.close(); values.close()
        self.assertEqual(original.read_bytes(), before)

    def test_lazy_branches_errors_arithmetic_text_and_dates(self):
        from datetime import date
        wb = Workbook(); s = wb.active
        s['A1'] = '=IF(1=1,"Fish & water",1/0)'
        s['A2'] = '=IFERROR(1/0,"pending")'
        s['A3'] = '=TEXT(1/8,"0%")&" total"'
        s['A4'] = '=INT(TODAY())'
        s['A5'] = '=IF(A8="","empty","filled")'
        s['A6'] = '=MAX(1,2*3)+4/2'
        actual = WorkbookCalculator(wb, date(2026, 9, 7)).calculate()
        self.assertEqual([actual[(s.title, f'A{i}')] for i in range(1, 7)],
                         ['Fish & water', 'pending', '13% total', 46272, 'empty', 8])

    def test_chart_cache_changes_with_source_selection(self):
        original = Path(__file__).resolve().parents[2] / 'Requirement_Source_Registry.xlsx'
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / 'registry.xlsx'; shutil.copy2(original, path)
            wb = load_workbook(path)
            initial = WorkbookCalculator(wb).calculate()[('Dashboard', 'C5')]
            wb['Source Register']['AJ3'] = 'EXCLUDE'
            u.save_workbook_atomic(wb, path, u.workbook_mtime(path)); wb.close()
            ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            with ZipFile(path) as z:
                chart = ET.fromstring(z.read('xl/charts/chart2.xml'))
                self.assertEqual(float(chart.findtext('.//c:numCache/c:pt/c:v', namespaces=ns)), initial - 1)

    def test_unknown_formula_or_cycle_prevents_publication(self):
        for formula in ('=UNSUPPORTED(1)', '=A1', '=1/0', "='[remote.xlsx]Sheet1'!A1"):
            with self.subTest(formula=formula), tempfile.TemporaryDirectory() as name:
                path = Path(name) / 'registry.xlsx'
                wb = Workbook(); wb.save(path); before = path.read_bytes()
                wb.active['A1'] = formula
                with self.assertRaises(FormulaCacheError):
                    u.save_workbook_atomic(wb, path, u.workbook_mtime(path))
                self.assertEqual(before, path.read_bytes())
                self.assertEqual(list(path.parent.glob('.*.tmp.xlsx')), [])

    def test_changed_workbook_during_cache_preparation_is_preserved(self):
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / 'registry.xlsx'
            wb = Workbook(); wb.save(path); wb.active['A1'] = '=1+1'
            original = WorkbookCalculator.calculate
            def concurrent(calculator):
                path.write_bytes(b'new external content')
                return original(calculator)
            with patch.object(WorkbookCalculator, 'calculate', concurrent):
                with self.assertRaisesRegex(u.UpdaterError, 'changed during'):
                    u.save_workbook_atomic(wb, path, u.workbook_mtime(path))
            self.assertEqual(path.read_bytes(), b'new external content')
