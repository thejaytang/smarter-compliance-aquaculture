"""Build an isolated engineering fixture; run with system2/.venv/bin/python.

Only this dedicated fixture root is writable. Its System1 bridge is real, using
an isolated owning database and an unchanged copy of the migration companion.
"""
from pathlib import Path
from hashlib import sha256
from io import BytesIO
import json
import shutil
import sqlite3
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / 'workbench/runtime/human-led-acceptance'


def build(extra_specimens=()):
    TARGET.mkdir(exist_ok=False)
    config_dir = TARGET / 'config'; config_dir.mkdir()
    data_root = TARGET / 'Data'; data_root.mkdir()
    runtime = TARGET / 'governance'; runtime.mkdir()
    config = json.loads((ROOT / 'system1/Code/config/config.json').read_text())
    config.update(workbook='../source-register.xlsx', source_root='../Data',
        governance_db='../governance/governance.sqlite', backup_root='../backups', log_root='../logs',
        manual_intake_root='../intake', manual_intake_archive_root='../intake-archive',
        discovery_candidate_inbox='../inbox.json', random_qa={'enabled':False})
    for folder in ('backups','logs','intake','intake-archive'):(TARGET/folder).mkdir()
    config_path = config_dir / 'config.json'; config_path.write_text(json.dumps(config,indent=2))
    source_db = ROOT / 'system1/Code/runtime/governance.sqlite'
    with sqlite3.connect(source_db.as_uri()+'?mode=ro',uri=True) as db:
        schema = list(db.execute("SELECT sql FROM sqlite_master WHERE sql IS NOT NULL AND name NOT LIKE 'sqlite_%' ORDER BY type DESC"))
        state = db.execute('SELECT * FROM state').fetchone()
        rows = {identity:json.loads(raw) for identity,raw in db.execute('SELECT id,data FROM sources')}
    shutil.copy2(source_db.parent/state[3],runtime/state[3])
    shutil.copy2(ROOT/'system1/Requirement_Source_Registry.xlsx',TARGET/'source-register.xlsx')
    specimens=[]
    for sid in ('PA001','CS010'):
        record=rows[sid]
        path=ROOT/'system1/Data'/record['folder_code'][1]/record['stored_filename'][1]
        specimens.append((sid,path.suffix,path.read_bytes(),record,f"Public saved original: {sid}"))
    html=b'''<!doctype html><html><head><style>p{display:none}</style><script>document.body.innerHTML='UNSAFE';</script></head><body><h1>Engineering fixture: water checks</h1><h2>1 Daily record</h2><p>Record the water temperature each morning.</p><table><tr><th colspan="2">Measurements</th></tr><tr><td>Temperature</td><td>12 C</td></tr></table><p>Keep a note beside each reading.</p><img src="https://example.invalid/image.png" alt="Original image requires source follow-up"><details><summary>Supplementary note</summary><p>Retain empty days in the record.</p></details></body></html>'''
    specimens.append(('TS001','.html',html,None,'ENGINEERING FIXTURE · HTML water checks'))
    book=Workbook();ws=book.active;ws.title='Measurements';ws['A1']='ENGINEERING FIXTURE';ws.merge_cells('A1:C2');ws['A3']='Temperature';ws['B3']=12;ws['C3']='=B3+1';ws['Z105']='Far cell: original retained';ws.row_dimensions[4].hidden=True
    book.create_sheet('Empty sheet');book.create_sheet('Hidden notes').sheet_state='hidden'
    buffer=BytesIO();book.save(buffer);book.close()
    specimens.append(('TS002','.xlsx',buffer.getvalue(),None,'ENGINEERING FIXTURE · Excel measurements'))
    buffer=BytesIO();pdf=canvas.Canvas(buffer,pagesize=(450,600));pdf.drawString(35,555,'ENGINEERING FIXTURE: native page');pdf.drawString(35,525,'Inspect this selectable original sentence.');pdf.showPage()
    scan=Image.new('RGB',(720,900),'white');draw=ImageDraw.Draw(scan);draw.text((40,80),'SCANNED ORIGINAL FIXTURE',fill='black');draw.text((40,130),'Manual transcription: water checked daily.',fill='black')
    pdf.drawImage(ImageReader(scan),0,0,width=450,height=562);pdf.showPage();pdf.showPage();pdf.save()
    specimens.append(('TS003','.pdf',buffer.getvalue(),None,'ENGINEERING FIXTURE · Native, scan and blank PDF'))
    # Additional locally authored fault specimens remain confined to this new
    # fixture. Existing callers keep the same five-source baseline.
    specimens.extend(extra_specimens)
    with sqlite3.connect(runtime/'governance.sqlite') as db:
        for (sql,) in schema:
            if sql.startswith('CREATE TABLE'):db.execute(sql)
        for (sql,) in schema:
            if not sql.startswith('CREATE TABLE'):db.execute(sql)
        db.execute('INSERT INTO state VALUES(?,?,?,?,?)',state)
        records=[]
        for position,(sid,suffix,raw,record,title) in enumerate(specimens,3):
            record=json.loads(json.dumps(record or rows['PA001']))
            snapshot=sid+'-001';filename=snapshot+'_material'+suffix
            folder='A_Public_Authority';directory=data_root/folder;directory.mkdir(exist_ok=True)
            path=directory/filename;path.write_bytes(raw);fingerprint=sha256(raw).hexdigest()
            values=dict(source_id=sid,snapshot_id=snapshot,source_title=title,folder_code=folder,stored_filename=filename,
                file_format=suffix[1:],content_hash=fingerprint,source_status='CURRENT',snapshot_status='STORED',
                download_status='SUCCESS',operator_selection_decision='INCLUDE',selection_status='INCLUDE',
                notes='Isolated engineering validation. No production review decision.',official_url='https://example.invalid/fixture',retrieval_url='https://example.invalid/fixture')
            for key,value in values.items():record[key]=['value',value]
            db.execute('INSERT INTO sources VALUES(?,?,?,?,?)',(sid,position,1,'INCLUDE',json.dumps(record)))
            rel=folder+'/'+filename
            db.execute('INSERT INTO artifacts VALUES(?,?,?)',(rel,fingerprint,len(raw)))
            db.execute('INSERT INTO source_versions VALUES(?,?,?,?,?)',(sid,snapshot,rel,fingerprint,1))
            records.append({'source_id':sid,'path':rel,'sha256':fingerprint,'kind':'public_original' if sid in {'PA001','CS010'} else 'synthetic'})
    app=TARGET/'workbench';app.mkdir();(app/'runtime').mkdir()
    for name in ('ui','src','.venv'):(app/name).symlink_to(ROOT/'workbench'/name,target_is_directory=True)
    (TARGET/'system2').symlink_to(ROOT/'system2',target_is_directory=True)
    (TARGET/'manifest.json').write_text(json.dumps({'scope':'isolated_engineering_only','sources':records,'config':str(config_path)},indent=2))
    print(TARGET)


if __name__=='__main__':build()
