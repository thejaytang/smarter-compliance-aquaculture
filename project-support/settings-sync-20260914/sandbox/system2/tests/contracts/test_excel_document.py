"""Source fidelity and corruption tests for independent spreadsheet extraction."""
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from pdf_extraction.contracts.excel import ExcelDocument
from pdf_extraction.formats.excel import parse_excel
from pdf_extraction.verification.excel import verify_excel
from pdf_extraction.orchestration.source_batch import run_item
from tests.contracts.test_source_intake import source, store


def workbook():
    w = Workbook(); ws = w.active; ws.title = 'Krav ø'
    ws.append(['Requirement', 'Limit', 'Formula'])
    ws.append(['Operator shall record ≥ 5 fish.', 5, '=B2*2'])
    ws['B2'].number_format = '0.00" kg"'
    ws['A4'] = 'Merged heading'; ws.merge_cells('A4:C4')
    ws['C7'] = '=SUM(B2,2)'; ws['E9'] = True; ws['F9'] = '#DIV/0!'
    ws['A2'].comment = Comment('Source note', 'Test')
    ws['A2'].hyperlink = 'https://example.invalid/evidence'
    ws.row_dimensions[2].hidden = True; ws.column_dimensions['B'].hidden = True
    ws.freeze_panes = 'A2'
    ws.add_table(Table(displayName='Requirements',ref='A1:C2'))
    hidden = w.create_sheet('Hidden annex'); hidden.sheet_state = 'veryHidden'; hidden['Z100'] = 'Do not omit'
    stream = BytesIO(); w.save(stream)
    return stream.getvalue()


def mutate_zip(raw, target, transform):
    out = BytesIO()
    with ZipFile(BytesIO(raw)) as src, ZipFile(out,'w',ZIP_DEFLATED) as dst:
        for n in src.namelist():
            dst.writestr(n, transform(src.read(n)) if n == target else src.read(n))
    return out.getvalue()


def test_structure_and_sparse_hidden_content():
    raw = workbook(); doc = parse_excel(raw,source(raw,'xlsx'))
    report = verify_excel(raw, doc)
    assert report['status'] == 'passed', report
    assert doc.sheets[1].state == 'veryHidden' and doc.sheets[1].cells[0].coordinate == 'Z100'
    assert doc.sheets[0].merged_ranges == ['A4:C4']
    assert any(r.get('hidden') == '1' for r in doc.sheets[0].rows)
    assert any('hidden="1"' in x for x in doc.sheets[0].structure_xml)
    cells = {c.coordinate:c for c in doc.sheets[0].cells}
    assert cells['C2'].formula == 'B2*2' and cells['C2'].cache_status == 'missing'
    assert cells['B2'].value == '5' and cells['B2'].attributes['s']
    assert cells['A2'].text == 'Operator shall record ≥ 5 fish.'
    assert any('tableColumns' in (p.xml or '') for p in doc.parts)
    assert any('Source note' in (p.xml or '') for p in doc.parts)
    assert any(i.startswith('external_relationship_not_fetched:') for i in doc.issues)


@pytest.mark.parametrize('change', ['drop_cell','text','formula','merge','hidden','rows','xml','parts','issue','review','order'])
def test_independent_verifier_detects_tampering(change):
    raw=workbook(); value=parse_excel(raw,source(raw,'xlsx')).model_dump()
    sheet=value['sheets'][0]
    if change=='drop_cell': sheet['cells'].pop()
    if change=='text': sheet['cells'][0]['text']='Invented'
    if change=='formula': next(c for c in sheet['cells'] if c['formula'])['formula']='1'
    if change=='merge': sheet['merged_ranges']=[]
    if change=='hidden': value['sheets'][1]['state']='visible'
    if change=='rows': sheet['rows'][1].pop('hidden')
    if change=='xml': sheet['structure_xml']=[]
    if change=='parts': value['parts'].pop()
    if change=='issue': value['issues']=[]
    if change=='review': sheet['cells'][0]['confidence']=1.0
    if change=='order': sheet['cells'].reverse()
    assert verify_excel(raw,ExcelDocument(**value))['status']=='failed'


def test_formula_cache_and_shared_formula_source_are_preserved():
    raw=workbook()
    raw=mutate_zip(raw,'xl/worksheets/sheet1.xml',lambda b:b.replace(b'<f>B2*2</f><v></v>',b'<f t="shared" si="0" ref="C2:C3">B2*2</f><v>10</v>'))
    doc=parse_excel(raw,source(raw,'xlsx')); cell=next(c for c in doc.sheets[0].cells if c.coordinate=='C2')
    assert cell.value=='10' and cell.cache_status=='present' and cell.formula_attributes['t']=='shared'
    assert verify_excel(raw,doc)['status']=='passed'


def test_empty_workbook_fails_instead_of_empty_success():
    w=Workbook(); out=BytesIO(); w.save(out); raw=out.getvalue()
    with pytest.raises(ValueError,match='empty_workbook'): parse_excel(raw,source(raw,'xlsx'))


def test_hash_mismatch():
    raw=workbook(); s=source(raw,'xlsx').model_copy(update={'content_hash':'0'*64})
    with pytest.raises(ValueError,match='hash_mismatch'): parse_excel(raw,s)


def test_duplicate_coordinate_fails():
    raw=workbook(); raw=mutate_zip(raw,'xl/worksheets/sheet1.xml',lambda b:b.replace(b'r="B1"',b'r="A1"'))
    with pytest.raises(ValueError,match='duplicate'): parse_excel(raw,source(raw,'xlsx'))


def test_pipeline_artifacts_and_failure_propagation(tmp_path,monkeypatch):
    raw=workbook(); s=source(raw,'xlsx'); store(tmp_path/'sources',s,raw)
    result=run_item(s,tmp_path/'sources',tmp_path/'run',{})
    assert result.status=='review_required' and result.canonical_schema_version=='excel-document/1'
    assert result.verification_sha256==sha256((tmp_path/'run/verification.json').read_bytes()).hexdigest()
    import pdf_extraction.verification.excel as verification
    monkeypatch.setattr(verification,'verify_excel',lambda *_:{'status':'failed','errors':['missing_cell']})
    result=run_item(s,tmp_path/'sources',tmp_path/'failure',{})
    assert result.status=='failed' and 'missing_cell' in result.reason_codes
    assert (tmp_path/'sources'/s.relative_path).read_bytes()==raw


def test_schema_matches_published_contract():
    assert json.loads(Path('config/schemas/excel-document.schema.json').read_text())==ExcelDocument.model_json_schema()


def test_rich_shared_string_and_empty_string_formula_cache():
    raw=workbook()
    raw=mutate_zip(raw,'xl/worksheets/sheet1.xml',lambda b:b.replace(
        b'<c r="A1" t="inlineStr"><is><t>Requirement</t></is></c>', b'<c r="A1" t="s"><v>0</v></c>').replace(
        b'<c r="C2"><f>B2*2</f><v></v></c>', b'<c r="C2" t="str"><f>IF(B2=5,&quot;&quot;,1)</f><v></v></c>'))
    out=BytesIO()
    with ZipFile(BytesIO(raw)) as src, ZipFile(out,'w',ZIP_DEFLATED) as dst:
        for n in src.namelist():
            data=src.read(n)
            if n=='xl/_rels/workbook.xml.rels':
                data=data.replace(b'</Relationships>',b'<Relationship Id="rStrings" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
            dst.writestr(n,data)
        dst.writestr('xl/sharedStrings.xml','<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><rPr><b/></rPr><t>Req</t></r><r><t>uirement</t></r><rPh sb="0" eb="1"><t>phonetic only</t></rPh></si></sst>')
    raw=out.getvalue(); doc=parse_excel(raw,source(raw,'xlsx'))
    cells={c.coordinate:c for c in doc.sheets[0].cells}
    assert cells['A1'].text=='Requirement'
    assert cells['C2'].cache_status=='present' and cells['C2'].value is None
    assert verify_excel(raw,doc)['status']=='passed'


def test_local_diagnostic_source_does_not_invent_registry_approval():
    from pdf_extraction.contracts.excel import LocalExcelSource
    raw=workbook(); s=LocalExcelSource(relative_path='reference.xlsx',content_hash=sha256(raw).hexdigest())
    doc=parse_excel(raw,s)
    assert doc.source.production_eligible is False
    assert 'operator_selection_decision' not in doc.source.model_dump()
    assert any(p.path.endswith('.vml') and p.xml for p in doc.parts)
    assert verify_excel(raw,doc)['status']=='passed'
