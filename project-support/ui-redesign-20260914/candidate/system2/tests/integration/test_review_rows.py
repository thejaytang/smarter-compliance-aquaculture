from copy import deepcopy
import json
import sqlite3
import time

import pytest
from pdf_extraction.review import row_store
from pdf_extraction.review.browser_view import browser_state
from pdf_extraction.domains.requirements.review_groups import group_units
from pdf_extraction.evidence.review_preview import sanitized_region
from tests.integration.test_two_stage_workflow import system, current, unit, request


def guarded(system, uid, action='correct', **kwargs):
    store,did=system
    detail=browser_state(store,'unit',did,uid)
    return request(document_id=did,unit_id=uid,revision=detail['revision'],policy_revision=detail['policy_revision'],
        source_sha256='a'*64,guard=detail['guard'],action=action,**kwargs)


def test_independent_edits_and_same_unit_conflict(system):
    store,did=system
    store.install(did,[unit('other')],[])
    first=guarded(system,'clause',fields={'body':'One shall be correct.'},note='Checked original')
    independent=guarded(system,'other',fields={'body':'Two shall be correct.'},note='Checked original')
    stale=guarded(system,'clause',fields={'body':'Obsolete correction'},note='Old browser')
    result=store.decision(first)
    assert store.decision(first)==result
    store.decision(independent)
    with pytest.raises(ValueError,match='stale_unit'):store.decision(stale)


def test_shared_dependency_invalidates_guard(system):
    store,did=system
    req=guarded(system,'clause',action='classify',classification='requirement')
    store.decision(guarded(system,'coverage',fields={'body':'Changed coverage text'},note='Source correction'))
    with pytest.raises(ValueError,match='stale_unit'):store.decision(req)


def test_migration_preserves_history_and_units(system):
    store,did=system
    doc=current(system)
    with store.transaction() as db:
        db.execute('DELETE FROM storage_version')
        db.execute('UPDATE documents SET data=? WHERE id=?',(json.dumps({k:v for k,v in doc.items() if not k.startswith('_')}),did))
    receipt=row_store.migrate(store)
    assert receipt['status']=='migrated'
    assert current(system)['units']==doc['units']
    with sqlite3.connect(receipt['backup']) as db:
        assert json.loads(db.execute('SELECT data FROM documents').fetchone()[0])['units']==doc['units']


def test_paginated_read_30k_without_full_hydration(system,monkeypatch):
    store,did=system
    seed=deepcopy(current(system)['units'][1])
    with store.transaction() as db:
        doc=store._load(db,did)
        doc['units']=[deepcopy(seed) for _ in range(30001)]
        for n,u in enumerate(doc['units']):u['id']=str(n);u['original']['fields']['body']='Farm shall keep records '+str(n)
        store._save(db,doc)
    monkeypatch.setattr(store,'_load',lambda *a,**k:(_ for _ in ()).throw(AssertionError('Full document load')))
    start=time.monotonic()
    page=browser_state(store,'tasks',did,offset=50,include_reviewed=True)
    detail=browser_state(store,'unit',did,'70')
    elapsed=time.monotonic()-start
    assert page['total']==30001 and len(page['items'])==50
    assert len(json.dumps(page))<40000 and 'original' not in page['items'][0]
    assert detail['unit']['id']=='70' and elapsed<1


def test_original_preview_strips_active_and_remote_content():
    raw=b'<html><body><script>evil()</script><p onclick="evil()">Original text<img src="https://remote/test"></p></body></html>'
    markup,warnings=sanitized_region(raw,[{'locator':'html:nth-of-type(1) > body:nth-of-type(1) > p:nth-of-type(1)::text(1)'}])
    assert 'Original text' in markup and 'evil()' not in markup and 'https://remote' not in markup
    assert "default-src 'none'" in markup and warnings


def test_grouping_retains_evidence_and_local_dependencies():
    units=[]
    for uid,kind,text,loc in [('h','source_heading','Chapter','h2'),('a','source_paragraph','One','p'),('b','source_paragraph','Two','p')]:
        u=unit(uid,kind);u['original']={'fields':{'title' if kind=='source_heading' else 'body':text},'references':[{'locator':loc}],'structure':[]};units.append(u)
    grouped=group_units(units)
    group=next(u for u in grouped if u['kind']=='source_section')
    assert group['original']['fields']['body']=='One\nTwo'
    assert {m['id'] for m in group['members']}=={'a','b'}
    assert 'h' in group['dependencies'] and 'b' not in grouped[0]['dependencies']


def test_table_sections_keep_duplicate_row_numbers():
    items=[]
    for i,part in enumerate(('thead','tbody')):
        u=unit(str(i),'source_table_cell')
        u['original']['fields']={'body':part+' wording'}
        u['original']['references']=[{'locator':f'html:nth-of-type(1) > body:nth-of-type(1) > table:nth-of-type(1) > {part}:nth-of-type(1) > tr:nth-of-type(1) > td:nth-of-type(1)::text(0)'}]
        items.append(u)
    table=next(u for u in group_units(items) if u['kind']=='source_table')
    assert len(table['table_rows'])==2
    assert table['original']['fields']['body']=='thead wording\ntbody wording'


def test_excel_part_locator_resolves_actual_sheet_and_merged_cells():
    from io import BytesIO
    from openpyxl import Workbook
    from pdf_extraction.evidence.review_preview import spreadsheet_region
    wb=Workbook();wb.active.title='Cover';ws=wb.create_sheet('Requirements');ws['D90']='Exact original';ws.merge_cells('D90:E90');ws.row_dimensions[90].hidden=True
    buf=BytesIO();wb.save(buf)
    result=spreadsheet_region(buf.getvalue(),[{'locator':'xl/worksheets/sheet2.xml#D90'}])
    assert result['sheet']=='Requirements' and result['row']==89 and result['column']==3
    cell=next(c for row in result['cells'] for c in row if c['address']=='D90')
    assert cell['value']=='Exact original' and cell['hidden_row']
    assert 'D90:E90' in result['merged']


def test_normative_note_is_not_excluded():
    from pdf_extraction.domains.requirements.classification import propose
    assert propose({'kind':'source_note','fields':{'notes':'The farm must retain records.'}})['classification']=='requirement'


def test_read_search_and_human_reason(system):
    store,did=system
    data=browser_state(store,'tasks',did,query='shall',include_reviewed=True)
    assert all('original' not in row and 'reason' in row for row in data['items'])


def test_pdf_crop_uses_coordinate_names_not_json_key_order(tmp_path):
    from pypdf import PdfWriter
    from pdf_extraction.evidence.review_preview import preview
    from hashlib import sha256
    from PIL import Image
    from io import BytesIO
    import base64
    writer=PdfWriter();writer.add_blank_page(width=600,height=800)
    p=tmp_path/'original.pdf'
    with p.open('wb') as f:writer.write(f)
    # Canonical JSON sorts keys alphabetically: x0,x1,y0,y1.
    result=preview(p,sha256(p.read_bytes()).hexdigest(),[{'page_index':0,'bbox':{'x0':100,'x1':200,'y0':300,'y1':400}}],tmp_path/'cache')
    image=Image.open(BytesIO(base64.b64decode(result['image'].split(',')[1])))
    assert image.size==(190,190)


def test_issues_can_be_resolved_independently(system):
    store,did=system
    with store.transaction() as db:
        doc=store._load(db,did);u=next(u for u in doc['units'] if u['id']=='clause');u['blockers']=['first','second'];store._recompute(db,doc,store.policy(db));store._save(db,doc)
    store.decision(guarded(system,'clause',action='resolve_content',resolved_issues=['first'],note='Verified first issue against the original.'))
    u=browser_state(store,'unit',did,'clause')['unit']
    assert u['blockers']==['second'] and u['resolved_issues']==['first']
    assert u['content_status']=='pending'


def test_location_correction_keeps_original_and_requires_recheck(system):
    store,did=system;before=browser_state(store,'unit',did,'clause')['unit']['original']
    store.decision(guarded(system,'clause',action='structure',structure=[{'role':'source_region','locator':'page 1: 10,20,100,200'}],note='Corrected original region only.'))
    u=browser_state(store,'unit',did,'clause')['unit']
    assert u['original']==before and u['content_status']=='pending'
    assert store.resolved(u)['references'][0]['bbox']==[10,20,100,200]
    assert 'human_location' in u['blockers']


def test_two_pending_levels_do_not_duplicate_units(system):
    store,did=system;store.install(did,[unit('unscored',score=None)],[])
    first=browser_state(store,'tasks',did,stage='content')
    second=browser_state(store,'tasks',did,stage='requirement')
    assert 'unscored' in {u['id'] for u in first['items']}
    assert 'unscored' not in {u['id'] for u in second['items']}
    store.decision(guarded(system,'unscored',action='accept_content'))
    first=browser_state(store,'tasks',did,stage='content')
    second=browser_state(store,'tasks',did,stage='requirement')
    assert 'unscored' not in {u['id'] for u in first['items']}
    assert 'unscored' in {u['id'] for u in second['items']}
    assert second['stage_counts']['requirement']==1
    all_content=browser_state(store,'tasks',did,stage='content',include_reviewed=True)
    assert 'unscored' in {u['id'] for u in all_content['items']}


def test_system1_original_hold_suspends_incremental_delivery(system):
    store,did=system;doc=current(system)
    assert doc['published']
    source=dict(doc['source'],_original_issues=[{'task_id':'source-issue','system2_document_id':did,'source_sha256':doc['source']['content_hash']}])
    store.reconcile_sources({'PA001':source})
    assert not current(system)['published']
    assert any(e['kind']=='suspend' for e in store.feed()['events'])
    store.reconcile_sources({'PA001':doc['source']})
    assert current(system)['published']


def test_task_summary_identifies_overlap_evidence_without_full_content(system):
    s,did=system
    with s.transaction() as db:
        doc=s._load(db,did);u=next(u for u in doc['units'] if u['id']=='clause');u['evidence_only']=True;u['blockers']=['overlap_evidence_only'];s._recompute(db,doc,s.policy(db));s._save(db,doc)
    from pdf_extraction.review.browser_view import browser_state
    result=browser_state(s,'tasks',did)
    item=next(i for i in result['items'] if i['id']=='clause')
    assert item['evidence_only']==1 and 'another parsing window' in item['reason']
    assert 'original' not in item and 'structure' not in item
