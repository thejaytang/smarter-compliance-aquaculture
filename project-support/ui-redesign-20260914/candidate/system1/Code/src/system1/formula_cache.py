"""Calculate the bounded formula vocabulary used by the System1 workbook.

This is not a general Excel engine. Unknown syntax, functions and references stop
the atomic save. Formulas remain intact; only their OOXML result caches change.
"""
from __future__ import annotations

from datetime import date, datetime
import math
from pathlib import Path
import re
from xml.sax.saxutils import escape
from zipfile import ZipFile

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel


class FormulaCacheError(ValueError):
    pass


class FormulaValueError(FormulaCacheError):
    """A calculation error that Excel IFERROR may handle."""


FUNCTIONS = {'IF', 'IFERROR', 'AND', 'OR', 'COUNTIF', 'COUNTIFS', 'COUNTA',
             'SUM', 'MAX', 'MAXIFS', 'TEXT', 'INT', 'TODAY'}
PRECEDENCE = {'=': 10, '<>': 10, '<': 10, '>': 10, '<=': 10, '>=': 10,
              '&': 20, '+': 30, '-': 30, '*': 40, '/': 40}


class FormulaParser:
    def __init__(self, formula):
        self.tokens = [t for t in Tokenizer(formula).items if t.type != 'WHITE-SPACE']
        self.pos = 0

    def take(self):
        if self.pos == len(self.tokens):
            raise FormulaCacheError('incomplete_formula')
        token = self.tokens[self.pos]
        self.pos += 1
        return token

    def expression(self, minimum=0):
        token = self.take()
        if token.type == 'OPERAND':
            left = ('operand', token.subtype, token.value)
        elif token.type == 'OPERATOR-PREFIX' and token.value in ('+', '-'):
            left = ('unary', token.value, self.expression(50))
        elif token.type == 'PAREN' and token.subtype == 'OPEN':
            left = self.expression()
            end = self.take()
            if (end.type, end.subtype) != ('PAREN', 'CLOSE'):
                raise FormulaCacheError('unclosed_parenthesis')
        elif token.type == 'FUNC' and token.subtype == 'OPEN':
            name = token.value[:-1].removeprefix('_xlfn.').upper()
            if name not in FUNCTIONS:
                raise FormulaCacheError('unsupported_function:' + name)
            args = []
            if self.pos < len(self.tokens) and self.tokens[self.pos].subtype != 'CLOSE':
                while True:
                    args.append(self.expression())
                    if self.pos == len(self.tokens) or self.tokens[self.pos].type != 'SEP':
                        break
                    self.take()
            end = self.take()
            if (end.type, end.subtype) != ('FUNC', 'CLOSE'):
                raise FormulaCacheError('unclosed_function')
            left = ('call', name, args)
        else:
            raise FormulaCacheError('unsupported_formula_token:' + token.value)
        while self.pos < len(self.tokens):
            token = self.tokens[self.pos]
            if token.type != 'OPERATOR-INFIX' or PRECEDENCE.get(token.value, -1) < minimum:
                break
            self.take()
            left = ('binary', token.value, left, self.expression(PRECEDENCE[token.value] + 1))
        return left

    def parse(self):
        result = self.expression()
        if self.pos != len(self.tokens):
            raise FormulaCacheError('unconsumed_formula_token')
        return result


def number(value):
    if value is None or value == '':
        return 0
    if isinstance(value, (int, float)):
        return value
    raise FormulaValueError('non_numeric_operand')


def text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def compare(left, right, operator):
    if left is None:
        left = '' if isinstance(right, str) else 0
    if right is None:
        right = '' if isinstance(left, str) else 0
    # Excel sorts numeric values before text, then logical values.
    def key(value):
        if isinstance(value, bool):
            return (2, value)
        if isinstance(value, str):
            return (1, value.casefold())
        return (0, value)
    a, b = key(left), key(right)
    return {'=': a == b, '<>': a != b, '<': a < b, '>': a > b,
            '<=': a <= b, '>=': a >= b}[operator]


def matches(value, criteria):
    if not isinstance(criteria, str):
        return compare(value, criteria, '=')
    match = re.fullmatch(r'(<=|>=|<>|=|<|>)?(.*)', criteria)
    op, target = match.group(1) or '=', match.group(2)
    if target and re.fullmatch(r'-?\d+(\.\d+)?', target):
        return compare(value, float(target), op)
    if op in ('=', '<>'):
        pattern = ''
        i = 0
        while i < len(target):
            char = target[i]
            if char == '~' and i + 1 < len(target):
                i += 1
                pattern += re.escape(target[i])
            else:
                pattern += '.*' if char == '*' else '.' if char == '?' else re.escape(char)
            i += 1
        matched = re.fullmatch(pattern, text(value), re.IGNORECASE | re.DOTALL) is not None
        return matched if op == '=' else not matched
    return compare(value, target, op)


class WorkbookCalculator:
    def __init__(self, workbook, today=None):
        self.workbook = workbook
        self.today = today or date.today()
        self.values = {}
        self.active = set()
        self.ranges = {}

    def cell(self, sheet, coordinate):
        key = (sheet, coordinate.replace('$', ''))
        if key in self.values:
            return self.values[key]
        if key in self.active:
            raise FormulaCacheError('circular_reference:' + str(key))
        cell = self.workbook[sheet][key[1]]
        self.active.add(key)
        try:
            if cell.data_type == 'f':
                value = self.evaluate(FormulaParser(cell.value).parse(), sheet)
                # A direct reference to a blank cell evaluates to zero in Excel.
                if value is None:
                    value = 0
            elif cell.data_type == 'e':
                raise FormulaValueError('source_error_cell:' + str(key))
            else:
                value = cell.value
                if isinstance(value, (datetime, date)):
                    value = to_excel(value, self.workbook.epoch)
            if isinstance(value, list):
                raise FormulaCacheError('array_result_not_supported')
            self.values[key] = value
            return value
        finally:
            self.active.remove(key)

    def reference(self, value, sheet):
        key = (sheet, value)
        if key in self.ranges:
            return self.ranges[key]
        table_ref = re.fullmatch(r'([A-Za-z_][\w.]*)\[([^\[\]]+)\]', value)
        if table_ref:
            name, header = table_ref.groups()
            tables = [(s, t) for s in self.workbook for t in s.tables.values() if t.name == name]
            if len(tables) != 1:
                raise FormulaCacheError('missing_or_duplicate_table:' + name)
            ws, table = tables[0]
            c1, r1, c2, r2 = range_boundaries(table.ref)
            columns = [c for c in range(c1, c2 + 1) if ws.cell(r1, c).value == header]
            if len(columns) != 1:
                raise FormulaCacheError('missing_or_duplicate_table_column:' + header)
            result = [self.cell(ws.title, ws.cell(r, columns[0]).coordinate)
                      for r in range(r1 + (table.headerRowCount or 0), r2 + 1 - (table.totalsRowCount or 0))]
        else:
            if '!' in value:
                name, value = value.rsplit('!', 1)
                sheet = name[1:-1].replace("''", "'") if name.startswith("'") and name.endswith("'") else name
            if sheet not in self.workbook.sheetnames or not re.fullmatch(r'\$?[A-Z]+\$?\d+(:\$?[A-Z]+\$?\d+)?', value):
                raise FormulaCacheError('unsupported_reference:' + value)
            if ':' not in value:
                return self.cell(sheet, value)
            c1, r1, c2, r2 = range_boundaries(value)
            if (r2 - r1 + 1) * (c2 - c1 + 1) > 100000:
                raise FormulaCacheError('range_too_large')
            ws = self.workbook[sheet]
            result = [self.cell(sheet, ws.cell(r, c).coordinate)
                      for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        self.ranges[key] = result
        return result

    def evaluate(self, node, sheet):
        kind, *args = node
        if kind == 'operand':
            subtype, value = args
            if subtype == 'NUMBER':
                return float(value)
            if subtype == 'TEXT':
                return value[1:-1].replace('""', '"')
            if subtype == 'LOGICAL':
                return value == 'TRUE'
            if subtype == 'RANGE':
                return self.reference(value, sheet)
            raise FormulaCacheError('unsupported_operand:' + subtype)
        if kind == 'unary':
            return number(self.evaluate(args[1], sheet)) * (-1 if args[0] == '-' else 1)
        if kind == 'binary':
            op, a, b = args
            left, right = self.evaluate(a, sheet), self.evaluate(b, sheet)
            if op in ('=', '<>', '<', '>', '<=', '>='):
                return compare(left, right, op)
            if op == '&':
                return text(left) + text(right)
            left, right = number(left), number(right)
            if op == '+': return left + right
            if op == '-': return left - right
            if op == '*': return left * right
            if right == 0: raise FormulaValueError('division_by_zero')
            return left / right
        name, expressions = args
        if name == 'IF':
            if len(expressions) != 3: raise FormulaCacheError('IF_arity')
            return self.evaluate(expressions[1 if self.evaluate(expressions[0], sheet) else 2], sheet)
        if name == 'IFERROR':
            if len(expressions) != 2: raise FormulaCacheError('IFERROR_arity')
            try: return self.evaluate(expressions[0], sheet)
            except FormulaValueError: return self.evaluate(expressions[1], sheet)
        values = [self.evaluate(e, sheet) for e in expressions]
        flat = [x for v in values for x in (v if isinstance(v, list) else [v])]
        if name == 'AND': return all(flat)
        if name == 'OR': return any(flat)
        if name == 'COUNTA': return sum(x is not None for x in flat)
        if name in ('SUM', 'MAX'):
            numeric = [x for x in flat if isinstance(x, (int, float)) and not isinstance(x, bool)]
            return sum(numeric) if name == 'SUM' else max(numeric, default=0)
        if name in ('COUNTIF', 'COUNTIFS', 'MAXIFS'):
            criteria_values = values[1:] if name == 'MAXIFS' else values
            if not criteria_values or len(criteria_values) % 2 or (name == 'COUNTIF' and len(values) != 2):
                raise FormulaCacheError('criteria_arity')
            pairs = list(zip(criteria_values[::2], criteria_values[1::2]))
            if any(not isinstance(r, list) for r, _ in pairs) or len({len(r) for r, _ in pairs}) != 1:
                raise FormulaCacheError('criteria_range_shape')
            indexes = [i for i in range(len(pairs[0][0])) if all(matches(r[i], criterion) for r, criterion in pairs)]
            if name != 'MAXIFS': return len(indexes)
            if not isinstance(values[0], list) or len(values[0]) != len(pairs[0][0]):
                raise FormulaCacheError('MAXIFS_range_shape')
            return max((v for i in indexes if isinstance(v := values[0][i], (int, float))), default=0)
        if name == 'TODAY' and not values: return to_excel(self.today, self.workbook.epoch)
        if name == 'INT' and len(values) == 1: return math.floor(number(values[0]))
        if name == 'TEXT' and len(values) == 2 and values[1] == '0%':
            value = number(values[0]) * 100
            rounded = math.copysign(math.floor(abs(value) + .5), value)
            return str(int(rounded)) + '%'
        raise FormulaCacheError('unsupported_function_arguments:' + name)

    def calculate(self):
        results = {}
        for sheet in self.workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f':
                        try:
                            results[(sheet.title, cell.coordinate)] = self.cell(sheet.title, cell.coordinate)
                        except Exception as exc:
                            raise FormulaCacheError(f'{sheet.title}!{cell.coordinate}: {exc}') from exc
        return results


def write_formula_caches(path: Path, workbook, results, calculator=None):
    """Patch only <v>/cell type in the freshly saved file; leave formula XML intact."""
    from xml.etree import ElementTree as ET
    from posixpath import normpath
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with ZipFile(path) as archive:
        entries = [(i, archive.read(i.filename)) for i in archive.infolist()]
    parts = {i.filename: data for i, data in entries}
    relations = {r.attrib['Id']: normpath('xl/' + r.attrib['Target']) if not r.attrib['Target'].startswith('/')
                 else r.attrib['Target'].lstrip('/') for r in ET.fromstring(parts['xl/_rels/workbook.xml.rels'])}
    patched = {}
    for sheet in ET.fromstring(parts['xl/workbook.xml']).findall('s:sheets/s:sheet', ns):
        name = sheet.attrib['name']
        target = relations[sheet.attrib['{' + rel_ns + '}id']]
        expected = {coordinate: value for (title, coordinate), value in results.items() if title == name}
        seen = set()
        def replace(match):
            attributes, body = match.group(1), match.group(2)
            coordinate = re.search(r'\br="([^"]+)"', attributes).group(1)
            if coordinate not in expected:
                return match.group(0)
            if '<f' not in body:
                raise FormulaCacheError('formula_missing_from_saved_cell:' + coordinate)
            value = expected[coordinate]
            kind = 'b' if isinstance(value, bool) else 'str' if isinstance(value, str) else 'n'
            if kind == 'n' and (not isinstance(value, (int, float)) or not math.isfinite(value)):
                raise FormulaCacheError('invalid_numeric_cache:' + coordinate)
            serialized = ('1' if value else '0') if kind == 'b' else text(value)
            attributes = re.sub(r'\s+t="[^"]*"', '', attributes) + f' t="{kind}"'
            body = re.sub(r'<v(?:\s[^>]*)?>.*?</v>|<v\s*/>', '', body, flags=re.S)
            seen.add(coordinate)
            return '<c' + attributes + '>' + body + '<v>' + escape(serialized) + '</v></c>'
        patched[target] = re.sub(r'<c(\s[^>]*?)(?<!/)>(.*?)</c>', replace, parts[target].decode('utf-8'), flags=re.S).encode('utf-8')
        if seen != set(expected):
            raise FormulaCacheError('formula_cache_coverage_mismatch:' + name)
    # Charts have separate result caches, which openpyxl otherwise retains stale.
    calculator = calculator or WorkbookCalculator(workbook)
    from html import unescape
    def chart_reference(match):
        prefix, kind, body = match.group(1) or '', match.group(2), match.group(3)
        formula = re.search(fr'<{prefix}f>(.*?)</{prefix}f>', body, re.S)
        if not formula or '!' not in unescape(formula.group(1)):
            raise FormulaCacheError('unsupported_chart_reference')
        values = calculator.reference(unescape(formula.group(1)), workbook.sheetnames[0])
        values = values if isinstance(values, list) else [values]
        cache = kind + 'Cache'
        old = re.search(fr'<{prefix}{cache}\b[^>]*>(.*?)</{prefix}{cache}>', body, re.S)
        format_code = re.search(fr'<{prefix}formatCode>.*?</{prefix}formatCode>', old.group(1), re.S) if old else None
        content = format_code.group(0) if format_code else ''
        content += f'<{prefix}ptCount val="{len(values)}"/>'
        for index, value in enumerate(values):
            if value is None:
                continue
            if kind == 'num' and (not isinstance(value, (int, float)) or not math.isfinite(value)):
                raise FormulaCacheError('non_numeric_chart_cache')
            content += f'<{prefix}pt idx="{index}"><{prefix}v>{escape(text(value))}</{prefix}v></{prefix}pt>'
        body = re.sub(fr'<{prefix}{cache}\b[^>]*>.*?</{prefix}{cache}>|<{prefix}{cache}\s*/>', '', body, flags=re.S)
        return f'<{prefix}{kind}Ref>{body}<{prefix}{cache}>{content}</{prefix}{cache}></{prefix}{kind}Ref>'
    for name, raw in parts.items():
        if re.fullmatch(r'xl/charts/chart\d+\.xml', name):
            patched[name] = re.sub(r'<((?:[A-Za-z_]\w*:)?)(num|str)Ref>(.*?)</\1\2Ref>',
                                   chart_reference, raw.decode('utf-8'), flags=re.S).encode('utf-8')
    with ZipFile(path, 'w') as archive:
        for info, data in entries:
            archive.writestr(info, patched.get(info.filename, data))
