from collections import Counter
from hashlib import sha256
from io import BytesIO
from pathlib import Path, PurePosixPath
from openpyxl import load_workbook
from ..contracts.source import Snapshot

class IntakeError(ValueError):
    pass

def read_registry(path: Path) -> tuple[list[dict], str]:
    # Read one immutable byte image. Never save/recalculate the business workbook.
    raw = path.read_bytes()
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        rows = workbook['Source Register'].iter_rows(values_only=True)
        next(rows)
        headers = next(rows)
        required = set(Snapshot.model_fields) - {'relative_path', 'registry_sha256'}
        required |= {'folder_code', 'stored_filename'}
        if not required <= set(headers):
            raise IntakeError('registry_headers_missing')
        if len([h for h in headers if h]) != len({h for h in headers if h}):
            raise IntakeError('duplicate_registry_headers')
        records = [dict(zip(headers, row)) for row in rows]
        return [r for r in records if r.get('source_id')], sha256(raw).hexdigest()
    finally:
        workbook.close()

def snapshot_path(source: Snapshot, root: Path) -> Path:
    rel = PurePosixPath(source.relative_path)
    if rel.is_absolute() or '..' in rel.parts or '\\' in source.relative_path:
        raise IntakeError('unsafe_snapshot_path')
    target = (root / rel).resolve()
    if not target.is_relative_to(root.resolve()):
        raise IntakeError('snapshot_outside_source_root')
    return target

def read_snapshot(source: Snapshot, root: Path) -> bytes:
    raw = snapshot_path(source, root).read_bytes()
    if sha256(raw).hexdigest() != source.content_hash:
        raise IntakeError('snapshot_hash_mismatch')
    return raw

def build_manifest(records: list[dict], registry_hash: str, root: Path,
                   source_ids: set[str]) -> dict:
    if not source_ids:
        raise IntakeError('explicit_source_ids_required')
    counts = Counter(str(r.get('source_id')) for r in records)
    items, rejected = [], []
    for source_id in sorted(source_ids):
        if counts[source_id] != 1:
            rejected.append({'source_id': source_id, 'reason': 'missing_or_duplicate_source_id'})
            continue
        row = next(r for r in records if r.get('source_id') == source_id)
        gates = {'operator_selection_decision':'INCLUDE', 'selection_status':'INCLUDE',
                 'snapshot_status':'STORED', 'source_status':'CURRENT', 'download_status':'SUCCESS'}
        bad = [k for k,v in gates.items() if row.get(k) != v]
        if bad:
            rejected.append({'source_id':source_id, 'reason':'ineligible_or_uncached_selection', 'fields':bad})
            continue
        try:
            folder, filename = str(row['folder_code']), str(row['stored_filename'])
            if any('/' in s or '\\' in s or s in {'', '.', '..'} for s in (folder, filename)):
                raise IntakeError('unsafe_snapshot_path')
            if not str(row['snapshot_id']).startswith(source_id + '-') or not filename.startswith(row['snapshot_id'] + '_'):
                raise IntakeError('snapshot_identity_mismatch')
            item = Snapshot(**{k:row[k] for k in Snapshot.model_fields if k not in {'relative_path','registry_sha256'}},
                            relative_path=f'{folder}/{filename}', registry_sha256=registry_hash)
            read_snapshot(item, root)
            items.append(item.model_dump())
        except (ValueError, OSError) as exc:
            rejected.append({'source_id':source_id, 'reason':str(exc)})
    return {'schema_version':'source-manifest/1', 'registry_sha256':registry_hash,
            'items':items, 'rejected':rejected}
