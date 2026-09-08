"""Failure-driven non-PDF regressions, including real XLSX containers and CLI IO."""
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import pytest
from openpyxl import Workbook
from pdf_extraction.formats.html import parse_html
from pdf_extraction.formats.router import route_format, FormatError
from pdf_extraction.intake.registry import build_manifest
from pdf_extraction.orchestration.source_batch import write_json
from tests.contracts.test_source_intake import CONFIG, HTML, source, row, store
from tests.support.paths import PROJECT_ROOT


@pytest.fixture
def subprocess_source_env(monkeypatch):
    # A relative PYTHONPATH changes meaning when the child changes cwd.
    monkeypatch.setenv('PYTHONPATH', str(PROJECT_ROOT / 'src'))


def html_with(body):
    return ('<html><body><div id="documentMeta"><h1>Test</h1></div><div id="documentBody">'+body+'</div></body></html>').encode()


def test_chapter_intro_is_flagged_as_unrepresented():
    raw=html_with('<div class="kapittel"><h2>Chapter</h2><p>Important introductory obligation.</p><div class="paragraf"><p>Clause.</p></div></div>')
    result=parse_html(raw,source(raw),CONFIG)
    assert any('unrepresented_body_text' in i for i in result.issues)


def test_parent_chapter_does_not_borrow_child_heading():
    raw=html_with('<div class="kapittel" id="outer"><div class="kapittel" id="inner"><h2>Child only</h2><div class="paragraf"><p>Clause.</p></div></div></div>')
    result=parse_html(raw,source(raw),CONFIG)
    chapters=[n for n in result.nodes if n.kind=='chapter']
    assert chapters[0].data['title']=='' and chapters[1].data['title']=='Child only'


def test_custom_metadata_selector_retains_fields():
    raw=HTML.replace(b'id="documentMeta"',b'id="customMeta"')
    result=parse_html(raw,source(raw),dict(CONFIG,metadata='#customMeta'))
    assert result.nodes[0].data['metadata']=={'Dato':'2026'}


def test_nested_paragraphs_do_not_duplicate_child_text():
    raw=html_with('<div class="paragraf" id="outer"><p>Parent.</p><div class="paragraf" id="inner"><p>Child.</p></div></div>')
    result=parse_html(raw,source(raw),CONFIG)
    paras=[n for n in result.nodes if n.kind=='paragraph']
    assert [n.data['text'] for n in paras]==['Parent.','Child.']
    assert paras[1].parent_id==paras[0].id


def test_nested_list_text_not_duplicated():
    raw=html_with('<div class="paragraf"><p>List:</p><table class="listeItem"><tr><td>a.</td><td>Parent <table class="listeItem" data-level="2"><tr><td>i.</td><td>Child</td></tr></table></td></tr></table></div>')
    result=parse_html(raw,source(raw),CONFIG)
    items=[n for n in result.nodes if n.kind=='paragraph'][0].data['list_items']
    assert [x['text'] for x in items]==['Parent','Child']


def test_bad_record_is_rejected_without_aborting_manifest(tmp_path):
    r=row(source()); r['snapshot_id']=None
    result=build_manifest([r],'a'*64,tmp_path,{'PA001'})
    assert not result['items'] and result['rejected']


def test_xlsx_cannot_be_registered_as_xls():
    stream=BytesIO(); Workbook().save(stream)
    with pytest.raises(FormatError,match='registered_format_mismatch'):
        route_format('example.xlsx','xls',stream.getvalue())


def test_json_serialization_failure_leaves_no_partial_file(tmp_path):
    p=tmp_path/'result.json'
    with pytest.raises(TypeError): write_json(p,{'unsupported':object()})
    assert not p.exists()


def test_json_does_not_overwrite_existing(tmp_path):
    p=tmp_path/'result.json'; p.write_text('original')
    with pytest.raises(FileExistsError): write_json(p,{'new':True})
    assert p.read_text()=='original'


def test_excel_only_cli_works_outside_project_directory(tmp_path, subprocess_source_env):
    raw_stream=BytesIO(); w=Workbook(); w.active['A1']='=1+1'; w.save(raw_stream)
    raw=raw_stream.getvalue(); s=source(raw,'xlsx'); store(tmp_path/'sources',s,raw)
    r=row(s); w=Workbook(); ws=w.active; ws.title='Source Register'; ws.append(['Display']); ws.append(list(r)); ws.append(list(r.values()))
    registry=tmp_path/'registry.xlsx'; w.save(registry)
    out=tmp_path/'run'
    proc=subprocess.run([sys.executable,'-m','pdf_extraction.orchestration.source_batch','--registry',str(registry),'--source-root',str(tmp_path/'sources'),'--source-id','PA001','--output',str(out)],cwd=tmp_path,capture_output=True,text=True)
    assert proc.returncode==0
    batch=json.loads((out/'batch.json').read_text())
    assert batch['results'][0]['status']=='review_required'
    assert batch['results'][0]['canonical_schema_version']=='excel-document/1'
    assert (tmp_path/'sources'/s.relative_path).read_bytes()==raw


def test_cli_missing_registry_writes_failure_receipt(tmp_path, subprocess_source_env):
    out=tmp_path/'run'
    proc=subprocess.run([sys.executable,'-m','pdf_extraction.orchestration.source_batch','--registry',str(tmp_path/'missing.xlsx'),'--source-root',str(tmp_path/'sources'),'--source-id','PA001','--output',str(out)],cwd=tmp_path,capture_output=True,text=True)
    assert proc.returncode==1
    batch=json.loads((out/'batch.json').read_text())
    assert batch['status']=='failed' and batch['reason_codes'] and batch['results']==[]


def test_publish_failure_cleans_temporary_file(tmp_path,monkeypatch):
    import pdf_extraction.orchestration.source_batch as batch
    def fail_link(*args):
        raise OSError('simulated storage failure')
    monkeypatch.setattr(batch.os,'link',fail_link)
    with pytest.raises(OSError): write_json(tmp_path/'result.json',{'value':1})
    assert list(tmp_path.iterdir())==[]
